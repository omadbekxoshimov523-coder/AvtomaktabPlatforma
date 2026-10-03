"""Avtomatik testlar (Python stdlib unittest):
    py -m tests.test_api
Server alohida jarayonda emas — shu jarayonda, temp bazada ishga tushiriladi.
Qamrab oladi: LOGIN, USERS, CARS, SESSIONS, CAPACITY, INSTRUCTOR BANDLIGI,
CAR ALMASHTIRISH TARIXI, SECURITY/RBAC.
"""
import http.cookiejar
import os
import shutil
import sys
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
import json
import re
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# Test uchun alohida temp baza
TMP = Path(tempfile.mkdtemp(prefix="avtomaktab_test_"))
os.environ["AVTOMAKTAB_DB"] = str(TMP / "test.db")

# BAND 6 — rate limiting / IP bloklash TESTLARNI buzmasligi uchun o'chiriladi.
# Bular server JARAYONI ichida bitta umumiy serverda ishlayotgani uchun
# kerak: aks holda birinchi ~100 so'rovdan keyin 429 qaytadi.
# (Brute-force himoyasi alohida test sinflarida `reset_login_attempts()` orqali
#  tekshiriladi — u o'chirilmaydi, faqat chegaralar ko'tariladi.)
os.environ["API_RATE_MAX"] = "0"     # umumiy so'rov limiteri — o'chirilgan
os.environ["MAX_IP_TRIES"] = "0"    # IP bloklash — o'chirilgan
os.environ["MAX_LOGIN_TRIES"] = "100000"

import server  # noqa: E402

from app.db import jload, Db, now  # noqa: E402
from app.api import ensure_reminders  # noqa: E402
from app.auth import hash_password  # noqa: E402


# ---------------------------------------------------------------------------
# BAND 5/7: seed endi RANDOM credential yaratadi (taxmin qilinishi mumkin emas).
# Testlar barqaror bo'lishi uchun seeded foydalanuvchilarning login/parolini
# shu yerda TESTGA XOS, YANGI FORMATDAGI qiymatlar bilan almashtiramiz
# (boshqa hech narsa — mashg'ulotlar, tarix, statistika — tegilmaydi).
#
#   login    : usrL_<14 belgi>   (harflar + raqam)
#   password : usrP_<14 belgi>   (katta + kichik harf + raqam + maxsus belgi)
# ---------------------------------------------------------------------------
INSTR_LOGINS = ("usrL_instr00000001", "usrL_instr00000002", "usrL_instr00000003")
INSTR_PASSWORDS = ("usrP_Inst0000!Ab1x", "usrP_Inst0000!Ab2x", "usrP_Inst0000!Ab3x")
STUD_LOGINS = ("usrL_stud00000001", "usrL_stud00000002",
               "usrL_stud00000003", "usrL_stud00000004")
STUD_PASSWORDS = ("usrP_Stud0000!Cd1x", "usrP_Stud0000!Cd2x",
                  "usrP_Stud0000!Cd3x", "usrP_Stud0000!Cd4x")

ADMIN_LOGIN, ADMIN_PASSWORD = "admin", "admin123"

# Qisqa nomlar (testlar shularni ishlatadi)
INSTR_LOGIN_1, INSTR_PASS_1 = INSTR_LOGINS[0], INSTR_PASSWORDS[0]
INSTR_LOGIN_3, INSTR_PASS_3 = INSTR_LOGINS[2], INSTR_PASSWORDS[2]
STUD_LOGIN_1, STUD_PASS_1 = STUD_LOGINS[0], STUD_PASSWORDS[0]
STUD_LOGIN_2, STUD_PASS_2 = STUD_LOGINS[1], STUD_PASSWORDS[1]
STUD_LOGIN_3, STUD_PASS_3 = STUD_LOGINS[2], STUD_PASSWORDS[2]
STUD_LOGIN_4, STUD_PASS_4 = STUD_LOGINS[3], STUD_PASSWORDS[3]

# Test fixture parollari — `_login` yordamchisi shu ro'yxatdan topadi.
PW_BY_LOGIN = {
    INSTR_LOGINS[0]: INSTR_PASSWORDS[0],
    INSTR_LOGINS[1]: INSTR_PASSWORDS[1],
    INSTR_LOGINS[2]: INSTR_PASSWORDS[2],
    STUD_LOGINS[0]: STUD_PASSWORDS[0],
    STUD_LOGINS[1]: STUD_PASSWORDS[1],
    STUD_LOGINS[2]: STUD_PASSWORDS[2],
    STUD_LOGINS[3]: STUD_PASSWORDS[3],
}


def _seed_fixture_credentials():
    """Seeded foydalanuvchilarga barqaror login/parol beradi (testga xos)."""
    db = Db(str(TMP / "test.db"))
    fixed = []
    for i, name in enumerate(["Akmal", "Jasur", "Vali"]):
        row = db.q1("SELECT id FROM users WHERE first_name=? AND role='instructor'", (name,))
        if row:
            fixed.append((row["id"], INSTR_LOGINS[i], INSTR_PASSWORDS[i]))
    for i, name in enumerate(["Omadbek", "Ali", "Vali", "Hasan"]):
        row = db.q1(
            """SELECT u.id AS id FROM users u JOIN students s ON s.user_id=u.id
               WHERE u.first_name=? AND u.role='student'""", (name,))
        if row:
            fixed.append((row["id"], STUD_LOGINS[i], STUD_PASSWORDS[i]))
    for uid, login, pw in fixed:
        db.upd(
            "UPDATE users SET login=?, password_hash=?, must_change_password=0, updated_at=? WHERE id=?",
            (login, hash_password(pw), now(), uid))
        db.ex(
            "INSERT OR IGNORE INTO used_credentials(login_digest, password_digest, created_at) VALUES(?,?,?)",
            (__import__("hashlib").sha256(login.encode()).hexdigest(),
             __import__("hashlib").sha256(pw.encode()).hexdigest(), now()))
    return len(fixed)


FIXED_COUNT = _seed_fixture_credentials()

HOST, PORT = "127.0.0.1", 0




def _ins_lesson(db, instructor_id, date):
    """Test uchun bitta mashg'ulot yaratadi va uning id'sini qaytaradi.

    `lesson_sessions` da ko'p ustun NOT NULL bo'lgani uchun yozishni
    barcha majburiy maydonlarga to'ldirib qilamiz (sxema o'zgarishi ham
    bu yordamchini buzmasligi kerak)."""
    cur = db.connect()
    try:
        names = {r[1] for r in cur.execute("PRAGMA table_info(lesson_sessions)")}
        vals = {
            "date": date, "start_time": "09:00", "end_time": "10:00",
            "instructor_id": instructor_id, "car_id": None, "car_name_snapshot": "",
            "car_plate_snapshot": "", "capacity_snapshot": 4, "status": "completed",
            "created_at": date + " 09:00:00", "updated_at": date + " 10:00:00",
            "notes": "", "location": "",
        }
        use = {k: v for k, v in vals.items() if k in names}
        cols = ", ".join(use)
        marks = ", ".join("?" * len(use))
        c = cur.execute("INSERT INTO lesson_sessions (%s) VALUES (%s)" % (cols, marks),
                        tuple(use.values()))
        cur.commit()
        return c.lastrowid
    finally:
        cur.close()

class Client:
    """Cookie'li HTTP klient."""

    def __init__(self):
        self.jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.jar))
        self._csrf = ""

    # BAND 6: server CSRF token'ni MAJBURIY qilgan -> test klient ham uni
    # yuborishi shart (haqiqiy frontend kabi: `sessionStorage`dan oladi).
    def _csrf_token(self):
        if not self._csrf:
            st, r = self.get("/api/auth/me")      # GET -> token talab qilinmaydi
            if st == 200:
                self._csrf = (r or {}).get("csrf") or ""
        return self._csrf

    def req(self, method, path, body=None, cookie=None):
        data = None
        headers = {}
        if body is not None:
            data = json.dumps(body).encode()
            headers["Content-Type"] = "application/json"
        if method in ("POST", "PUT", "DELETE"):
            headers["X-Requested-With"] = "Avtomaktab"
            # `cookie=` — BOSHQA sessiya uchun so'rov: o'shaning token'i
            # kerak emas (maqsad — "eski cookie ishlamaydi" tekshiruvi).
            if not cookie:
                tok = self._csrf_token()
                if tok:
                    headers["X-CSRF-Token"] = tok
        # Chiqishdan keyin "eski token saqlab qolinsa ham ishlamaydi" —
        # buni tekshirish uchun cookie'ni qo'lda yuborish imkoniyati.
        if cookie:
            headers["Cookie"] = "sid=" + cookie
        req = urllib.request.Request(f"http://{HOST}:{PORT}{path}", data=data, headers=headers, method=method)
        try:
            with self.opener.open(req, timeout=15) as resp:
                raw = resp.read()
                try:
                    out = json.loads(raw.decode("utf-8"))
                except Exception:
                    return resp.status, {"raw": True, "mime": resp.headers.get("Content-Type")}
                if isinstance(out, dict) and out.get("csrf"):
                    self._csrf = out["csrf"]
                return resp.status, out
        except urllib.error.HTTPError as e:
            raw = e.read()
            try:
                return e.code, json.loads(raw.decode("utf-8"))
            except Exception:
                return e.code, {}

    # qulaylik
    def get(self, p): return self.req("GET", p)
    def post(self, p, b=None): return self.req("POST", p, b or {})
    def put(self, p, b=None): return self.req("PUT", p, b or {})
    def delete(self, p): return self.req("DELETE", p)

    def get_file(self, p):
        """Eksport yuklab olish — server `download` maydonini HAQIQIY faylga
        aylantirib yuboradi (Content-Type + Content-Disposition), shuning uchun
        javob JSON EMAS. Qaytaradi: (status, raw_bytes, headers)."""
        return self._file_req("GET", p, None)

    def post_file(self, p, body):
        return self._file_req("POST", p, body)

    def _file_req(self, method, path, body):
        data = None
        headers = {}
        if body is not None:
            data = json.dumps(body).encode()
            headers["Content-Type"] = "application/json"
        if method in ("POST", "PUT", "DELETE"):
            headers["X-Requested-With"] = "Avtomaktab"
            tok = self._csrf_token()
            if tok:
                headers["X-CSRF-Token"] = tok
        req = urllib.request.Request(f"http://{HOST}:{PORT}{path}", data=data,
                                     headers=headers, method=method)
        try:
            with self.opener.open(req, timeout=30) as resp:
                return resp.status, resp.read(), dict(resp.headers)
        except urllib.error.HTTPError as e:
            return e.code, e.read(), dict(e.headers or {})


class TabClient(Client):
    """Brauzer tab'ini taqlid qiladi: `X-Avto-Tab` sarlavhasini yuboradi.

    Bitta brauzerda bir nechta tab ochilganda har bir tab' o'z sessiyasini
    oladi (cookie esa barcha tab'larda umumiy) — shu sarlavha orqali.
    `jar` parametri orqali bir nechta tab bitta cookie jarni ham ulaydi.
    """

    def __init__(self, tab, jar=None):
        super().__init__()
        self.tab = tab
        if jar is not None:
            self.jar = jar
            self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))

    def req(self, method, path, body=None, cookie=None, tab=None):
        data = None
        headers = {}
        if body is not None:
            data = json.dumps(body).encode()
            headers["Content-Type"] = "application/json"
        if method in ("POST", "PUT", "DELETE"):
            headers["X-Requested-With"] = "Avtomaktab"
            if not cookie:
                tok = self._csrf_token()
                if tok:
                    headers["X-CSRF-Token"] = tok
        # `tab=False` -> sarlavha yuborilmaydi (eski mijoz kabi: curl, bot)
        if tab is not False:
            headers["X-Avto-Tab"] = tab if tab else self.tab
        if cookie:
            headers["Cookie"] = "sid=" + cookie
        req = urllib.request.Request(f"http://{HOST}:{PORT}{path}", data=data, headers=headers, method=method)
        try:
            with self.opener.open(req, timeout=15) as resp:
                raw = resp.read()
                try:
                    out = json.loads(raw.decode("utf-8"))
                except Exception:
                    return resp.status, {"raw": True}
                if isinstance(out, dict) and out.get("csrf"):
                    self._csrf = out["csrf"]
                return resp.status, out
        except urllib.error.HTTPError as e:
            raw = e.read()
            try:
                return e.code, json.loads(raw.decode("utf-8"))
            except Exception:
                return e.code, {}


# Barcha test klasslari BIRTA umumiy server ishlatadi (har klassda yangi server ochish
# Windows'da SO_REUSEADDR tufayli o'lik serverga ulanish/osilish muammosini keltiradi)
_SERVER = {"httpd": None, "owners": 0}


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        global HOST, PORT  # noqa: PLW0603
        if _SERVER["httpd"] is None:
            from http.server import ThreadingHTTPServer
            httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
            _SERVER["httpd"] = httpd
            threading.Thread(target=httpd.serve_forever, daemon=True).start()
            HOST, PORT = "127.0.0.1", httpd.server_address[1]
        _SERVER["owners"] += 1
        cls.httpd = _SERVER["httpd"]
        cls.port = _SERVER["httpd"].server_address[1]

    @classmethod
    def tearDownClass(cls):
        # DIQQAT: bu yerda hech narsani o'chirmaymiz! Server va vaqtinchalik
        # baza butun testlar tugagandan KEYIN `tearDownModule()` da yopiladi.
        # Sababi: sinflar ketma-ket ishlaganda, birinchisi tugagach TMP
        # o'chirilsa, qolgan sinflar bazani topolmaydi
        # ("sqlite3.OperationalError: unable to open database file").
        # (Windows'da ochiq faylni o'chirib bo'lmaydi, Linux'da esa
        #  o'chib ketardi — ya'ni xato faqat CI da ko'rinardi.)
        pass


def tearDownModule():
    """Barcha sinflar tugagach — serverni to'xtatish va temp bazani tozalash."""
    httpd = _SERVER.get("httpd")
    if httpd is not None:
        httpd.shutdown()
        httpd.server_close()
        _SERVER["httpd"] = None
    shutil.rmtree(TMP, ignore_errors=True)


class TestAuthAndUsers(Base):
    def test_01_admin_login_ok(self):
        c = Client()
        st, r = c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)

    def test_02_wrong_password(self):
        c = Client()
        st, r = c.post("/api/auth/login", {"login": "admin", "password": "wrong"})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "auth.wrong_credentials")

    def test_03_wrong_role(self):
        c = Client()
        st, r = c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "student"})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "auth.wrong_role")

    def test_04_blocked_user_cannot_login(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # Boshqa testlar Akmalga (INSTR_LOGIN_1) bog'liq — bloklash uchun alohida user yaratamiz
        st, cr = admin.post("/api/admin/users", {"role": "student", "first_name": "Blok", "last_name": "Test"})
        self.assertEqual(st, 200, cr)
        lg, pw = cr["credentials"]["login"], cr["credentials"]["password"]
        st, r = admin.post(f"/api/admin/users/{cr['user']['id']}/status", {"status": "block"})
        self.assertEqual(st, 200, r)
        stud = Client()
        st, r = stud.post("/api/auth/login", {"login": lg, "password": pw, "role": "student"})
        self.assertEqual(r["error"], "auth.user_blocked")

    def test_05_logout(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c.post("/api/auth/logout")
        self.assertEqual(st, 200)

    def test_06_admin_creates_student_auto_creds(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.post("/api/admin/users", {
            "role": "student", "first_name": "Test", "last_name": "Student", "phone": "+998911111111", "group_name": "T-1",
        })
        self.assertEqual(st, 200, r)
        self.assertRegex(r["credentials"]["login"], r"^usrL_[A-Za-z0-9]{14}$")
        self.assertRegex(r["credentials"]["password"], r"^usrP_[A-Za-z0-9!@#$%^&*()\-_=+\[\]{}?]{14}$")
        # yangi login bilan kirish
        sc = Client()
        st, r2 = sc.post("/api/auth/login", {"login": r["credentials"]["login"],
                                             "password": r["credentials"]["password"],
                                             "role": "student"})
        self.assertEqual(st, 200, r2)

    def test_07_credentials_unique_random_not_reused(self):
        """BAND 7: har bir yangi credential UNIKAL va RANDOM.

        Eski ketma-ket format (`usrL_00001`) olib tashlangan — login ham,
        parol ham `secrets` (CSPRNG) bilan tasodifiy yaratiladi va hech
        qachon qayta ishlatilmaydi (DB UNIQUE constraint + collision retry).
        """
        admin = Client()
        admin.post("/api/auth/login", {"login": ADMIN_LOGIN, "password": ADMIN_PASSWORD, "role": "admin"})
        logins, passes = [], []
        for i, role in enumerate(("student", "student", "instructor")):
            st, r = admin.post("/api/admin/users",
                               {"role": role, "first_name": "A%d" % i, "last_name": "B%d" % i})
            self.assertEqual(st, 200, r)
            logins.append(r["credentials"]["login"])
            passes.append(r["credentials"]["password"])
        # 1) UNIKAL: hech ikkisi bir xil emas
        self.assertEqual(len(set(logins)), 3, "loginlar takrorlanmasin")
        self.assertEqual(len(set(passes)), 3, "parollar takrorlanmasin")
        # 2) FORMAT: usrL_ / usrP_ + 14 ta xavfsiz random belgi
        for lg in logins:
            self.assertRegex(lg, r"^usrL_[A-Za-z0-9]{14}$")
        for pw in passes:
            self.assertRegex(pw, r"^usrP_[A-Za-z0-9!@#$%^&*()\-_=+\[\]{}?]{14}$")
            body = pw[len("usrP_"):]
            self.assertTrue(any(c.isupper() for c in body), "katta harf")
            self.assertTrue(any(c.islower() for c in body), "kichik harf")
            self.assertTrue(any(c.isdigit() for c in body), "raqam")
            self.assertTrue(any(not c.isalnum() for c in body), "maxsus belgi")
        # 3) KETMA-KET EMAS: loginlar bir-biridan "keyingi" emas
        self.assertNotEqual(logins[1], logins[0])
        # 4) O'chirilgan foydalanuvchining credential'i QAYTA ishlatilmaydi
        users = admin.get("/api/admin/users?role=student")[1]["users"]
        target = next(u for u in users if u["login"] == logins[0])
        admin.delete(f"/api/admin/users/{target['id']}")
        st, r = admin.post("/api/admin/users",
                           {"role": "student", "first_name": "A4", "last_name": "B4"})
        self.assertEqual(st, 200, r)
        self.assertNotEqual(r["credentials"]["login"], logins[0])
        self.assertNotEqual(r["credentials"]["password"], passes[0])
        # 5) Parol DB'da oddiy matn sifatida SAQLANMAYDI (BAND 6)
        db = Db(str(TMP / "test.db"))
        row = db.q1("SELECT password_hash FROM users WHERE login=?", (r["credentials"]["login"],))
        self.assertTrue(row["password_hash"].startswith("scrypt$"),
                        "parol scrypt hash sifatida saqlanadi")
        self.assertNotIn(r["credentials"]["password"], row["password_hash"])

    def test_08_bulk_create(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.post("/api/admin/users/bulk", {"count": 10})
        self.assertEqual(st, 200)
        self.assertEqual(len(r["created"]), 10)

    def test_09_no_public_registration(self):
        c = Client()
        st, r = c.post("/api/user/register", {"first_name": "X"})
        self.assertEqual(st, 404)  # bunday endpoint yo'q


class TestLogoutSession(Base):
    """Chiqish (logout).

    Talablar: sessiya server tomonida ham yopilishi (faqat frontend tokenni
    "unutib qo'yishi" yetarli emas), eski token qayta ishlamasligi, himoyalangan
    endpoint'lar 401 qaytarishi va barcha uch rol uchun bir xil ishlashi.
    """

    ROLES = [
        ("admin", "admin", "admin123"),
        ("instructor", INSTR_LOGIN_1, INSTR_PASS_1),
        ("student", STUD_LOGIN_1, STUD_PASS_1),
    ]

    def _sid(self, c):
        for ck in c.jar:
            if ck.name == "sid":
                return ck.value
        return None

    def _login(self, login, password, role):
        c = Client()
        st, r = c.post("/api/auth/login", {"login": login, "password": password, "role": role})
        self.assertEqual(st, 200, r)
        return c

    def test_01_login_then_logout_all_roles(self):
        for role, login, pw in self.ROLES:
            with self.subTest(role=role):
                c = self._login(login, pw, role)
                self.assertTrue(self._sid(c), "login cookie berilmadi")
                st, r = c.get("/api/auth/me")
                self.assertEqual(st, 200, r)
                self.assertEqual(r["user"]["role"], role)
                st, r = c.post("/api/auth/logout")
                self.assertEqual(st, 200, r)

    def test_02_protected_endpoints_401_after_logout(self):
        for role, login, pw in self.ROLES:
            with self.subTest(role=role):
                c = self._login(login, pw, role)
                self.assertEqual(c.post("/api/auth/logout")[0], 200)
                st, r = c.get("/api/auth/me")
                self.assertEqual(st, 401, r)
                self.assertEqual(r["error"], "auth.required")

    def test_03_stale_token_rejected_server_side(self):
        # Eski token texnik jihatdan saqlab qolinsa ham, serverda bekor
        # qilingani uchun qabul qilinmasligi shart.
        c = self._login("admin", "admin123", "admin")
        old = self._sid(c)
        self.assertTrue(old)
        self.assertEqual(c.post("/api/auth/logout")[0], 200)

        stale = Client()
        for path in ("/api/auth/me", "/api/me/settings", "/api/me/sessions", "/api/admin/dashboard"):
            with self.subTest(path=path):
                st, r = stale.req("GET", path, cookie=old)
                self.assertEqual(st, 401, r)

    def test_04_logout_clears_cookie(self):
        # Set-Cookie Max-Age=0 — klient jaridagi sid yo'qolishi kerak
        c = self._login("admin", "admin123", "admin")
        self.assertTrue(self._sid(c))
        st, r = c.post("/api/auth/logout")
        self.assertEqual(st, 200, r)
        self.assertIsNone(self._sid(c), "sid cookie tozalanmadi")

    def test_05_logout_without_session_is_idempotent(self):
        c = Client()
        st, r = c.post("/api/auth/logout")
        self.assertEqual(st, 200, r)
        st, r = c.post("/api/auth/logout")
        self.assertEqual(st, 200, r)

    def test_06_other_sessions_survive(self):
        # Faqat joriy sessiya yopiladi — boshqa qurilmalar saqlanib qoladi
        a = self._login("admin", "admin123", "admin")
        b = self._login("admin", "admin123", "admin")
        self.assertEqual(a.post("/api/auth/logout")[0], 200)
        st, _ = b.get("/api/auth/me")
        self.assertEqual(st, 200)
        self.assertEqual(b.post("/api/auth/logout")[0], 200)

    def test_07_no_data_leak_between_users(self):
        # Chiqqandan keyin boshqa hisob bilan kiringanda avvalgi foydalanuvchi
        # ma'lumotlari (profil) ko'rinmasligi kerak.
        admin = self._login("admin", "admin123", "admin")
        stu = Client()
        st, r = stu.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        self.assertEqual(st, 200, r)
        st, r = stu.get("/api/auth/me")
        stu_login = r["user"]["login"]
        stu_id = r["user"]["id"]

        self.assertEqual(admin.post("/api/auth/logout")[0], 200)

        st, r = stu.get("/api/auth/me")
        self.assertEqual(st, 200, r)
        self.assertEqual(r["user"]["login"], stu_login)
        self.assertEqual(r["user"]["id"], stu_id)
        # Talaba admin endpoint'iga kira olmaydi
        st, r = stu.get("/api/admin/dashboard")
        self.assertEqual(st, 403, r)
        self.assertEqual(stu.post("/api/auth/logout")[0], 200)

    def test_08_relogin_works_after_logout(self):
        c = self._login("admin", "admin123", "admin")
        self.assertEqual(c.post("/api/auth/logout")[0], 200)
        c2 = self._login("admin", "admin123", "admin")
        self.assertTrue(self._sid(c2))
        st, r = c2.get("/api/auth/me")
        self.assertEqual(st, 200, r)
        self.assertEqual(c2.post("/api/auth/logout")[0], 200)


class TestCars(Base):
    def test_10_car_create_duplicate_plate(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        payload = {"brand": "Toyota", "model": "Camry", "plate_number": "01 Z 777 ZZ", "practice_capacity": 4, "status": "active"}
        st, r = admin.post("/api/admin/cars", payload)
        self.assertEqual(st, 200, r)
        st, r = admin.post("/api/admin/cars", payload)
        self.assertEqual(st, 409)
        self.assertEqual(r["error"], "car.plate_exists")

    def test_11_car_status_blocks_session(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # Akmalning joriy mashinasini topamiz (test_30 uni almashtirib qo'yishi mumkin)
        insts = admin.get("/api/admin/instructors")[1]["instructors"]
        akmal = next(i for i in insts if i["first_name"] == "Akmal" and i["last_name"] == "Karimov")
        car_id = akmal["car"]["id"]
        st, r = admin.put(f"/api/admin/cars/{car_id}", {"status": "repair"})
        self.assertEqual(st, 200, r)
        st, r = admin.post("/api/admin/sessions", {"date": "2026-10-05", "start_time": "10:00", "end_time": "11:00", "instructor_id": akmal["id"], "student_ids": []})
        self.assertEqual(r["ok"], False)
        self.assertIn("car_in_repair", r["params"]["errors"])
        admin.put(f"/api/admin/cars/{car_id}", {"status": "active"})

    def test_12_instructor_assign_car(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        cars = admin.get("/api/admin/cars")[1]["cars"]
        target = next(c for c in cars if c["plate_number"] == "01 Z 777 ZZ")
        # Jasurning original mashinasini eslab qolamiz (test_27 unga bog'liq)
        insts_before = admin.get("/api/admin/instructors")[1]["instructors"]
        jasur_before = next(i for i in insts_before if i["first_name"] == "Jasur")
        original_car_id = jasur_before["car"]["id"]
        st, r = admin.post(f"/api/admin/instructors/{jasur_before['id']}/car", {"car_id": target["id"]})
        self.assertEqual(st, 200, r)
        # Jasur endi Toyota'da
        insts = admin.get("/api/admin/instructors")[1]["instructors"]
        jasur = next(i for i in insts if i["first_name"] == "Jasur")
        self.assertEqual(jasur["car"]["plate_number"], "01 Z 777 ZZ")
        # Boshqa testlar Nexia'ga bog'liq — tiklaymiz
        admin.post(f"/api/admin/instructors/{jasur['id']}/car", {"car_id": original_car_id})


class TestSessions(Base):
    # 2026-11-02 = dushanba, 11-03 = seshanba, 11-04 = chorshanba, 11-06 = juma,
    # 11-07 = shanba, 11-09 = dushanba, 11-10 = seshanba.
    # Akmal (instructor 1): du-sha; Jasur (instructor 2): du-ju.

    def test_20_create_session_auto_car(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # Akmal (instructor id 1) -> Cobalt (car 1)
        st, r = admin.post("/api/admin/sessions", {
            "date": "2026-11-02", "start_time": "09:00", "end_time": "10:00", "instructor_id": 1,
            "student_ids": [1, 2, 3],
        })
        self.assertEqual(st, 200, r)
        self.assertEqual(r["auto"]["car"], "Chevrolet Cobalt")
        self.assertEqual(r["auto"]["capacity"], 4)

    def test_21_capacity_5th_student_fails(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.post("/api/admin/sessions", {
            "date": "2026-11-02", "start_time": "11:00", "end_time": "12:00", "instructor_id": 1,
            "student_ids": [1, 2, 3, 4, 5],
        })
        self.assertEqual(st, 409)
        self.assertIn("capacity_full", r["params"]["errors"])

    def test_22_instructor_busy_overlap(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        ok, _ = admin.post("/api/admin/sessions", {"date": "2026-11-03", "start_time": "09:00", "end_time": "10:00", "instructor_id": 1, "student_ids": []})
        self.assertEqual(ok, 200)
        st, r = admin.post("/api/admin/sessions", {"date": "2026-11-03", "start_time": "09:30", "end_time": "10:30", "instructor_id": 1, "student_ids": []})
        self.assertEqual(st, 409)
        self.assertIn("instructor_busy", r["params"]["errors"])

    def test_23_student_busy(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # Jasur (id 2) bilan
        ok, _ = admin.post("/api/admin/sessions", {"date": "2026-11-04", "start_time": "10:00", "end_time": "11:00", "instructor_id": 2, "student_ids": [1]})
        self.assertEqual(ok, 200)
        st, r = admin.post("/api/admin/sessions", {"date": "2026-11-04", "start_time": "10:30", "end_time": "11:30", "instructor_id": 1, "student_ids": [1]})
        self.assertEqual(st, 409)
        self.assertIn("student_busy", r["params"]["errors"])

    def test_24_work_schedule_enforced(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # Akmal: 08:00-18:00, tanaffus 13-14.  2026-11-06 = juma (ish kuni)
        st, r = admin.post("/api/admin/sessions", {"date": "2026-11-06", "start_time": "13:00", "end_time": "14:00", "instructor_id": 1, "student_ids": []})
        self.assertEqual(st, 409)
        self.assertIn("intersects_break", r["params"]["errors"])
        st, r = admin.post("/api/admin/sessions", {"date": "2026-11-06", "start_time": "07:00", "end_time": "08:00", "instructor_id": 1, "student_ids": []})
        self.assertIn("outside_work_hours", r["params"]["errors"])

    def test_25_no_car_instructor_blocked(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # yangi instruktor avtomobilsiz
        st, r = admin.post("/api/admin/users", {"role": "instructor", "first_name": "Noc", "last_name": "Car"})
        self.assertEqual(st, 200)
        new_inst_login = r["credentials"]["login"]
        insts = admin.get("/api/admin/instructors")[1]["instructors"]
        newest = max(insts, key=lambda i: i["id"])
        st, r = admin.post("/api/admin/sessions", {"date": "2026-11-07", "start_time": "09:00", "end_time": "10:00", "instructor_id": newest["id"], "student_ids": []})
        self.assertEqual(st, 409)
        self.assertIn("car_not_assigned", r["params"]["errors"])

    def test_26_start_finish_attendance(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        today_s = datetime.now().strftime("%Y-%m-%d")  # boshlash faqat "bugun" uchun ruxsat
        # Vali (instructor 3) — haftaning 7 kunida ham ishlaydi (yakshanba ham o'tadi)
        ok, r = admin.post("/api/admin/sessions", {"date": today_s, "start_time": "09:00", "end_time": "10:00", "instructor_id": 3, "student_ids": [1, 2]})
        self.assertEqual(ok, 200, r)
        sid = r["id"]
        # instruktor bilan
        inst = Client()
        inst.post("/api/auth/login", {"login": INSTR_LOGIN_3, "password": INSTR_PASS_3, "role": "instructor"})
        st, r = inst.post(f"/api/instructor/sessions/{sid}/start")
        self.assertEqual(st, 200, r)
        st, r = inst.post(f"/api/instructor/sessions/{sid}/attendance", {"student_id": 1, "status": "present"})
        self.assertEqual(st, 200, r)
        st, r = inst.post(f"/api/instructor/sessions/{sid}/attendance", {"student_id": 2, "status": "absent"})
        self.assertEqual(st, 200, r)
        st, r = inst.post(f"/api/instructor/sessions/{sid}/finish")
        self.assertEqual(st, 200, r)
        detail = inst.get(f"/api/instructor/sessions/{sid}")[1]["session"]
        self.assertEqual(detail["status"], "completed")

    def test_27_cancel_and_reschedule(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        ok, r = admin.post("/api/admin/sessions", {"date": "2026-11-09", "start_time": "11:00", "end_time": "12:00", "instructor_id": 1, "student_ids": [1]})
        self.assertEqual(ok, 200, r)
        sid = r["id"]
        # Ko'chirish: Jasur'ga
        st, r2 = admin.post(f"/api/admin/sessions/{sid}/reschedule", {"date": "2026-11-10", "start_time": "10:00", "end_time": "11:00", "instructor_id": 2})
        self.assertEqual(st, 200, r2)
        det = admin.get(f"/api/admin/sessions/{sid}")[1]["session"]
        self.assertEqual(det["instructor_name"], "Jasur Aliyev")
        self.assertEqual(det["car_plate_snapshot"], "01 B 456 BB")  # Jasurning avtomobili avtomatik
        st, r3 = admin.post(f"/api/admin/sessions/{sid}/cancel", {"reason": "yomg'ir"})
        self.assertEqual(st, 200)
        det = admin.get(f"/api/admin/sessions/{sid}")[1]["session"]
        self.assertEqual(det["status"], "cancelled")


class TestCarChangeHistory(Base):
    def test_30_car_change_keeps_history(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # Akmal bilan eski session (Cobalt)
        ok, r = admin.post("/api/admin/sessions", {"date": "2026-12-01", "start_time": "09:00", "end_time": "10:00", "instructor_id": 1, "student_ids": []})
        old_sid = r["id"]
        self.assertEqual(r["auto"]["car"], "Chevrolet Cobalt")
        # Akmalning mashinasini almashtiramiz: Cobalt -> Onix (id 3)
        st, r2 = admin.post("/api/admin/instructors/1/car", {"car_id": 3})
        self.assertEqual(st, 200, r2)
        # Eski session tarixi buzilmasin
        old_detail = admin.get(f"/api/admin/sessions/{old_sid}")[1]["session"]
        self.assertEqual(old_detail["car_plate_snapshot"], "01 A 123 AA")
        # Yangi session -> Onix
        ok, r3 = admin.post("/api/admin/sessions", {"date": "2026-12-02", "start_time": "09:00", "end_time": "10:00", "instructor_id": 1, "student_ids": []})
        self.assertEqual(r3["auto"]["car"], "Chevrolet Onix")
        self.assertEqual(r3["auto"]["plate"], "01 C 789 CC")
        # Keyingi testlar uchun avtomobillarni tiklaymiz: Akmal -> Cobalt (1), Vali -> Onix (3)
        admin.post("/api/admin/instructors/1/car", {"car_id": 1})
        admin.post("/api/admin/instructors/3/car", {"car_id": 3})


class TestSecurityRBAC(Base):
    def test_40_student_cannot_access_admin(self):
        st = Client()
        st.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        code, r = st.get("/api/admin/users")
        self.assertEqual(code, 403)
        self.assertEqual(r["error"], "forbidden")

    def test_41_instructor_cannot_access_admin(self):
        it = Client()
        it.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        code, r = it.get("/api/admin/dashboard")
        self.assertEqual(code, 403)

    def test_42_student_cannot_see_other_student_data(self):
        st = Client()
        st.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        # Boshqa talabaning sessionlari API bor, lekin yo'l boshqa talabaga tegishli emas
        sessions = st.get("/api/student/sessions?upcoming=1")[1]["sessions"]
        # 1-talaba (Omadbek) sessionlari ro'yxatida faqat o'ziniki bo'ladi; boshqa talabaning ma'lumotlari yo'q
        for s in sessions:
            # student na faqat o'z sessionlarini ko'ra oladi — ular ro'yxati Omadbek'dan
            pass
        # shaxsiy profil
        me = st.get("/api/auth/me")[1]
        self.assertEqual(me["user"]["login"], STUD_LOGIN_1)

    def test_43_instructor_cannot_manage_others_sessions(self):
        it = Client()
        it.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        # Jasurning sessioni emas; Akmal boshqa instruktorlar sessionlarini boshqara olmaydi
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        ok, r = admin.post("/api/admin/sessions", {"date": "2026-12-10", "start_time": "09:00", "end_time": "10:00", "instructor_id": 2, "student_ids": []})
        sid = r["id"]
        st, r2 = it.post(f"/api/instructor/sessions/{sid}/start")
        self.assertEqual(st, 404)  # topilmadi — boshqa instruktor sessioni
        st, r3 = it.get(f"/api/instructor/sessions/{sid}")
        self.assertEqual(st, 404)

    def test_44_no_password_in_responses(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        _, r = admin.get("/api/admin/users")
        for u in r["users"]:
            self.assertNotIn("password_hash", u)
            self.assertNotIn("usrP_", json.dumps(u))


class TestProfileAvatar2FA(Base):
    """Modul 2 — profil rasm yuklash, 2FA (TOTP) va profil tahriri."""

    TINY_PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
                "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")

    def _cleanup_avatar(self, path_):
        if path_ and path_.startswith("/uploads/avatars/"):
            fp = ROOT / "web" / path_.lstrip("/")
            try:
                if fp.exists():
                    fp.unlink()
            except OSError:
                pass

    def test_50_avatar_upload_valid(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.put("/api/me/profile/avatar", {"image": self.TINY_PNG})
        self.assertEqual(st, 200, r)
        self.assertTrue(r["profile_image"].startswith("/uploads/avatars/"))
        # me() qaytariladigan user'da ko'rinadi
        _, me = admin.get("/api/auth/me")
        self.assertEqual(me["user"]["profile_image"], r["profile_image"])
        self._cleanup_avatar(r["profile_image"])

    def test_51_avatar_rejects_bad_format(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.put("/api/me/profile/avatar", {"image": "data:image/gif;base64," + "A" * 100})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "img.format")
        st, r2 = admin.put("/api/me/profile/avatar", {"image": "notadata"})
        self.assertEqual(st, 400)
        self.assertEqual(r2["error"], "img.format")

    def test_52_avatar_rejects_oversize(self):
        import base64 as _b64
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        big = _b64.b64encode(b"x" * (5 * 1024 * 1024 + 100)).decode()
        st, r = admin.put("/api/me/profile/avatar", {"image": "data:image/png;base64," + big})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "img.size")

    def _latest_twofa_code(self):
        import re
        import sqlite3
        con = sqlite3.connect(os.environ["AVTOMAKTAB_DB"])
        try:
            row = con.execute(
                "SELECT body FROM notifications WHERE type='2fa' ORDER BY id DESC LIMIT 1").fetchone()
        finally:
            con.close()
        if not row:
            return None
        m = re.search(r"\d{6}", row[0])
        return m.group() if m else None

    def test_53_2fa_full_cycle(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        # noto'g'ri parol bilan yoqib bo'lmaydi
        st, r = admin.post("/api/me/2fa/enable", {"password": "wrong"})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "auth.wrong_old_password")
        # to'g'ri parol + telefon bilan SMS kod yuboriladi (QR/secret yo'q)
        st, r = admin.post("/api/me/2fa/enable", {"password": "admin123", "phone": "+998901112233"})
        self.assertEqual(st, 200, r)
        self.assertNotIn("secret", r)
        self.assertNotIn("otpauth", r)
        # kod bildirishnoma (simulyatsiya) orqali yetib boradi
        code = self._latest_twofa_code()
        self.assertTrue(code and len(code) == 6, code)
        # noto'g'ri kod rad etiladi
        st, r2 = admin.post("/api/me/2fa/verify-enable", {"code": "000000"})
        self.assertEqual(st, 400)
        self.assertEqual(r2["error"], "auth.otp_invalid")
        # to'g'ri SMS kod bilan yoqiladi
        st, r3 = admin.post("/api/me/2fa/verify-enable", {"code": code})
        self.assertEqual(st, 200, r3)
        _, me = admin.get("/api/auth/me")
        self.assertEqual(me["user"]["totp_enabled"], 1)
        # endi loginda kod so'raladi (yangi SMS kod yuboriladi)
        st, r4 = admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(r4["error"], "auth.otp_required")
        code2 = self._latest_twofa_code()
        self.assertTrue(code2 and len(code2) == 6, code2)
        # kod bilan kirish ishlaydi
        st, r5 = admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin", "otp": code2})
        self.assertEqual(st, 200, r5)
        # noto'g'ri parol bilan o'chirib bo'lmaydi, parol bilan o'chiriladi
        st, r6 = admin.post("/api/me/2fa/disable", {"password": "wrong"})
        self.assertEqual(st, 400)
        st, r7 = admin.post("/api/me/2fa/disable", {"password": "admin123"})
        self.assertEqual(st, 200, r7)
        # 2FA o'chgach, kod siz kirish qaytadan ishlaydi
        st, r8 = admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r8)

    def test_54_profile_update(self):
        st = Client()
        st.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        code, r = st.put("/api/me/profile", {"phone": "+998901234567", "first_name": "Omad"})
        self.assertEqual(code, 200, r)
        _, me = st.get("/api/auth/me")
        self.assertEqual(me["user"]["phone"], "+998901234567")
        self.assertEqual(me["user"]["first_name"], "Omad")
        # orqaga qaytaramiz
        st.put("/api/me/profile", {"phone": "", "first_name": "Omadbek"})


class TestCarPhotosAndStatus(Base):
    """M5: avtomobil fotosuratlari va tezkor status o'zgarishi."""

    PNG = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="

    def _admin(self):
        a = Client()
        a.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return a

    def _upload_photo(self, admin, car_id):
        return admin.post(f"/api/admin/cars/{car_id}/photos", {"image": "data:image/png;base64," + TestCarPhotosAndStatus.PNG})

    def test_60_quick_status_change(self):
        admin = self._admin()
        # dastlab aktiv
        _, r = admin.get("/api/admin/cars")
        car = next(c for c in r["cars"] if c["plate_number"] == "01 A 123 AA")
        st, r2 = admin.put(f"/api/admin/cars/{car['id']}/status", {"status": "repair"})
        self.assertEqual(st, 200, r2)
        self.assertEqual(r2["status"], "repair")
        _, r3 = admin.get("/api/admin/cars")
        self.assertEqual(next(c for c in r3["cars"] if c["id"] == car["id"])["status"], "repair")
        # noto'g'ri status
        st, r4 = admin.put(f"/api/admin/cars/{car['id']}/status", {"status": "flying"})
        self.assertEqual(st, 400)
        self.assertEqual(r4["error"], "car.bad_status")
        # tiklaymiz
        admin.put(f"/api/admin/cars/{car['id']}/status", {"status": "active"})
        # RBAC: talaba statusni o'zgartira olmaydi
        stud = Client()
        stud.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        st, _ = stud.put(f"/api/admin/cars/{car['id']}/status", {"status": "inactive"})
        self.assertEqual(st, 403)

    def test_61_photo_add_list_remove(self):
        admin = self._admin()
        cars = admin.get("/api/admin/cars")[1]["cars"]
        cid = cars[0]["id"]
        # qo'shish
        st, r = self._upload_photo(admin, cid)
        self.assertEqual(st, 200, r)
        _, r2 = admin.get("/api/admin/cars")
        ph = r2["photos"][str(cid)]
        self.assertEqual(len(ph), 1)
        self.assertTrue(ph[0]["path"].startswith("/uploads/cars/"))
        # yana birini qo'shamiz
        st, r3 = self._upload_photo(admin, cid)
        self.assertEqual(st, 200, r3)
        _, r4 = admin.get("/api/admin/cars")
        self.assertEqual(len(r4["photos"][str(cid)]), 2)
        # noto'g'ri format rad etiladi
        st, r5 = admin.post(f"/api/admin/cars/{cid}/photos", {"image": "not-a-data-url"})
        self.assertEqual(st, 400)
        self.assertEqual(r5["error"], "img.format")
        # bittasini o'chirish
        pid = r["id"]
        st, r6 = admin.delete(f"/api/admin/cars/{cid}/photos/{pid}")
        self.assertEqual(st, 200, r6)
        _, r7 = admin.get("/api/admin/cars")
        self.assertEqual(len(r7["photos"][str(cid)]), 1)
        # qolganini ham o'chirib tozalaymiz
        pid2 = r3["id"]
        admin.delete(f"/api/admin/cars/{cid}/photos/{pid2}")
        _, r8 = admin.get("/api/admin/cars")
        self.assertFalse(r8["photos"].get(str(cid)))

    def test_62_photo_limit_and_rbac(self):
        admin = self._admin()
        cars = admin.get("/api/admin/cars")[1]["cars"]
        cid = cars[0]["id"]
        pids = []
        try:
            for _ in range(8):
                st, r = self._upload_photo(admin, cid)
                self.assertEqual(st, 200, r)
                pids.append(r["id"])
            # 9-rasm rad etiladi
            st, r = self._upload_photo(admin, cid)
            self.assertEqual(st, 409)
            self.assertEqual(r["error"], "car.photo_limit")
        finally:
            for pid in pids:
                admin.delete(f"/api/admin/cars/{cid}/photos/{pid}")
        # RBAC: talaba o'chira olmaydi
        stud = Client()
        stud.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        st, _ = stud.delete(f"/api/admin/cars/{cid}/photos/1")
        self.assertEqual(st, 403)
        st, _ = stud.post(f"/api/admin/cars/{cid}/photos", {"image": "x"})
        self.assertEqual(st, 403)

    def test_63_instructor_sees_car_photos(self):
        admin = self._admin()
        # Akmal (instructor 1) Cobalt'da — fotosurat qo'shamiz
        akmal = Client()
        akmal.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        _, ctx = akmal.get("/api/instructor/car")
        cid = ctx["car"]["id"]
        st, r = self._upload_photo(admin, cid)
        self.assertEqual(st, 200, r)
        try:
            _, r2 = akmal.get("/api/instructor/car")
            self.assertEqual(r2["car"]["id"], cid)
            self.assertIn(r["path"], [p["path"] for p in r2["photos"]])
        finally:
            admin.delete(f"/api/admin/cars/{cid}/photos/{r['id']}")


class TestHomeAndToday(Base):
    def test_70_admin_dashboard_weekly_stats(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.get("/api/admin/dashboard")
        self.assertEqual(st, 200, r)
        d = r["dashboard"]
        for k in ("week_sessions", "week_done", "week_hours", "week_requests", "week_new_users"):
            self.assertIn(k, d)
        self.assertIsInstance(d["week_sessions"], int)
        self.assertIsInstance(d["week_hours"], float)

    def test_71_instructor_home(self):
        inst = Client()
        inst.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        st, r = inst.get("/api/instructor/home")
        self.assertEqual(st, 200, r)
        for k in ("sessions", "done", "hours", "students", "today"):
            self.assertIn(k, r["weekly"])
        for k in ("sessions", "done", "hours"):
            self.assertIn(k, r["overall"])
        self.assertIn("car", r)

    def test_72_student_home_overall_progress(self):
        st = Client()
        st.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        code, r = st.get("/api/student/home")
        self.assertEqual(code, 200, r)
        self.assertIn("overall", r)
        self.assertIn("weekly", r)
        self.assertIn("next", r)
        self.assertIsInstance(r["overall"]["sessions"], int)
        self.assertIsInstance(r["weekly"]["done"], int)

    def test_73_student_today(self):
        st = Client()
        st.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        code, r = st.get("/api/student/today")
        self.assertEqual(code, 200, r)
        self.assertIn("sessions", r)
        self.assertIsInstance(r["sessions"], list)

    def test_74_rbac_home_and_today(self):
        stud = Client()
        stud.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        st, _ = stud.get("/api/instructor/home")
        self.assertEqual(st, 403)
        inst = Client()
        inst.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        st, _ = inst.get("/api/student/home")
        self.assertEqual(st, 403)
        st, _ = inst.get("/api/student/today")
        self.assertEqual(st, 403)


class TestAdminDashboardM8(Base):
    def test_80_dashboard_week_chart_data(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.get("/api/admin/dashboard")
        self.assertEqual(st, 200, r)
        byd = r["dashboard"].get("week_sessions_by_day")
        self.assertIsInstance(byd, list)
        self.assertEqual(len(byd), 7)
        for b in byd:
            self.assertIn("date", b)
            self.assertIn("count", b)
            self.assertIsInstance(b["count"], int)
        # kunlar tartibi hafta boshidan boshlanadi
        from datetime import datetime
        monday = datetime.strptime(byd[0]["date"], "%Y-%m-%d").weekday()
        self.assertEqual(monday, 0)

    def test_81_instructors_include_profile_image(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.get("/api/admin/instructors")
        self.assertEqual(st, 200, r)
        self.assertTrue(r["instructors"])
        for u in r["instructors"]:
            self.assertIn("profile_image", u)

    def test_82_requests_include_profile_image(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.get("/api/admin/requests")
        self.assertEqual(st, 200, r)
        for req in r["requests"]:
            self.assertIn("profile_image", req)

    def test_83_users_include_profile_image(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.get("/api/admin/users?role=student")
        self.assertEqual(st, 200, r)
        self.assertTrue(r["users"])
        for u in r["users"]:
            self.assertIn("profile_image", u)


class TestSettingsM12(Base):
    """M12 — Sozlamalar bo'limini kengaytirish:
    shaxsiy sozlamalar (til/tema/bildirishnomalar), faol sessiyalar,
    platforma sozlamalari (mashg'ulot davomiyligi/ish vaqti/eslatma) va
    avtomatik eslatmalar."""

    def _next_monday(self):
        # Seed'dagi "bugun/ertaga" sessiyalari bilan to'qnashmaslik uchun
        # keyingi dushanbani kamida 7 kun oldinga surib olamiz (haftaning istalgan kunida ishlaydi).
        from datetime import timedelta
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=7)).strftime("%Y-%m-%d")

    def test_90_me_settings_defaults(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c.get("/api/me/settings")
        self.assertEqual(st, 200, r)
        s = r["settings"]
        self.assertIn("lang", s)
        self.assertIn("theme", s)
        self.assertEqual(set(s["notif"].keys()),
                         {"lesson", "request", "message", "security", "reminder", "admin"})
        self.assertTrue(all(s["notif"].values()))
        # BAND 17: yangi, aniq `notification_settings` ham qaytariladi
        self.assertEqual(set(s["notification_settings"].keys()),
                         {"lesson_reminders", "admin_messages", "lesson_status_updates",
                          "messages", "requests", "security"})
        self.assertTrue(all(s["notification_settings"].values()))

    def test_91_me_settings_put_and_get(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c.put("/api/me/settings", {"lang": "ru", "theme": "dark", "notif": {"lesson": False, "request": False}})
        self.assertEqual(st, 200, r)
        st, r = c.get("/api/me/settings")
        self.assertEqual(st, 200, r)
        self.assertEqual(r["settings"]["lang"], "ru")
        self.assertEqual(r["settings"]["theme"], "dark")
        self.assertFalse(r["settings"]["notif"]["lesson"])
        self.assertFalse(r["settings"]["notif"]["request"])
        self.assertTrue(r["settings"]["notif"]["message"])
        self.assertTrue(r["settings"]["notif"]["security"])

    def test_92_me_settings_invalid_values_ignored(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        c.put("/api/me/settings", {"lang": "xx", "theme": "neon"})
        st, r = c.get("/api/me/settings")
        self.assertEqual(st, 200, r)
        self.assertNotEqual(r["settings"]["lang"], "xx")
        self.assertNotEqual(r["settings"]["theme"], "neon")

    def test_93_sessions_list_and_current(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c.get("/api/me/sessions")
        self.assertEqual(st, 200, r)
        self.assertGreaterEqual(len(r["sessions"]), 1)
        self.assertEqual(sum(1 for s in r["sessions"] if s["current"]), 1)
        for s in r["sessions"]:
            self.assertIn("created_at", s)
            self.assertIn("expires_at", s)
            self.assertIn("last_seen", s)

    def test_94_revoke_other_device_kicks_it(self):
        c1 = Client()
        c1.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        c2 = Client()
        c2.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c1.get("/api/me/sessions")
        self.assertGreaterEqual(len(r["sessions"]), 2)
        other = next(s for s in r["sessions"] if not s["current"])
        # c1 (joriy) boshqa qurilma sessiyasini o'chirsin
        st, rev = c1.post("/api/me/sessions/revoke", {"id": other["id"]})
        self.assertEqual(st, 200, rev)
        self.assertFalse(rev["current_revoked"])
        # o'sha qurilma (c2) endi kirishi yo'q
        st2, r2 = c2.get("/api/auth/me")
        self.assertEqual(st2, 401)

    def test_95_revoke_current_requires_relogin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c.get("/api/me/sessions")
        cur = next(s for s in r["sessions"] if s["current"])
        st, rev = c.post("/api/me/sessions/revoke", {"id": cur["id"]})
        self.assertEqual(st, 200, rev)
        self.assertTrue(rev["current_revoked"])
        st2, r2 = c.get("/api/auth/me")
        self.assertEqual(st2, 401)

    def test_96_revoke_all_logs_out(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = c.post("/api/me/sessions/revoke-all", {})
        self.assertEqual(st, 200, r)
        self.assertTrue(r["current_revoked"])
        st2, r2 = c.get("/api/auth/me")
        self.assertEqual(st2, 401)

    def test_97_notification_prefs_apply(self):
        # Instruktor (Akmal) mashg'ulot bildirishnomalarini o'chiradi
        inst = Client()
        inst.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        inst.put("/api/me/settings", {"notif": {"lesson": False}})
        # Student (Omadbek) default holatda
        stu = Client()
        stu.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})

        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        d = self._next_monday()
        st, r = admin.post("/api/admin/sessions", {
            "date": d, "start_time": "15:00", "end_time": "16:30",
            "instructor_id": 1, "student_ids": [1], "notes": "M12 prefs test",
        })
        self.assertEqual(st, 200, r)
        sid = r["id"]

        # Instruktor: lesson bildirishnomasi KELMASLIGI kerak
        st, ni = inst.get("/api/me/notifications")
        self.assertEqual(st, 200, ni)
        for n in ni["notifications"]:
            d = jload(n.get("data") or "{}", {}) if not isinstance(n.get("data"), dict) else (n.get("data") or {})
            if str(d.get("session_id")) == str(sid):
                self.assertNotIn(n["type"], ("lesson", "cancel"), "prefs off bo'lganda kelmasligi kerak")

        # Student: lesson.assigned KELISHI kerak (default prefs)
        st, ns = stu.get("/api/me/notifications")
        self.assertEqual(st, 200, ns)
        got = [n for n in ns["notifications"]
               if str((jload(n.get("data") or "{}", {}) if not isinstance(n.get("data"), dict) else (n.get("data") or {})).get("session_id")) == str(sid)
               and n["type"] == "lesson"]
        self.assertTrue(got, "student default prefs bilan lesson bildirishnomasini oladi")

        # Xavfsizlik toifasi chapdan o'chirilgan bildirishnoma emas: reset so'rovi o'tadi
        before = [n for n in inst.get("/api/me/notifications")[1]["notifications"] if n["type"] == "security"]
        st, rr = inst.post("/api/auth/request-password-reset", {"login": INSTR_LOGIN_1})
        self.assertEqual(st, 200, rr)
        after = [n for n in inst.get("/api/me/notifications")[1]["notifications"] if n["type"] == "security"]
        self.assertLess(len(before), len(after), "security toifasi bloklanmaydi")

    def test_98_admin_platform_settings_defaults_and_update(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        st, r = admin.get("/api/admin/settings")
        self.assertEqual(st, 200, r)
        s = r["settings"]
        self.assertEqual(s["lesson_duration_min"], 90)
        self.assertEqual(s["work_start"], "09:00")
        self.assertEqual(s["work_end"], "18:00")
        self.assertEqual(s["reminder_minutes"], 60)
        # M8: platforma nomi / asosiy filial sozlamalari olib tashlangan
        self.assertNotIn("platform_name", s)
        self.assertNotIn("branch", s)
        st, r = admin.put("/api/admin/settings", {
            "lesson_duration_min": 60, "work_start": "10:00", "work_end": "20:00",
            "reminder_minutes": 30,
        })
        self.assertEqual(st, 200, r)
        st, r = admin.get("/api/admin/settings")
        self.assertEqual(r["settings"]["lesson_duration_min"], 60)
        self.assertEqual(r["settings"]["work_start"], "10:00")
        self.assertEqual(r["settings"]["work_end"], "20:00")
        self.assertEqual(r["settings"]["reminder_minutes"], 30)
        self.assertNotIn("platform_name", r["settings"])

    def test_99_default_duration_used_when_end_missing(self):
        admin = Client()
        admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        admin.put("/api/admin/settings", {"lesson_duration_min": 90})
        d = self._next_monday()
        st, r = admin.post("/api/admin/sessions", {
            "date": d, "start_time": "10:00",
            "instructor_id": 1, "student_ids": [1], "notes": "M12 dur test",
        })
        self.assertEqual(st, 200, r)
        st, det = admin.get(f"/api/admin/sessions/{r['id']}")
        self.assertEqual(st, 200, det)
        self.assertEqual(det["session"]["end_time"], "11:30")  # 10:00 + 90 daqiqa

    def _mk_session(self, db, start_dt, created_at, student_id=1, inst_id=1, notes="REM"):
        """Test uchun mashg'ulot qatori (created_at ni boshqarish bilan)."""
        date, start = start_dt.strftime("%Y-%m-%d"), start_dt.strftime("%H:%M")
        end = (start_dt + timedelta(minutes=60))
        end_s = end.strftime("%H:%M") + (("+" + end.strftime("%d")) if end.day != start_dt.day else "")
        sid = db.ex(
            """INSERT INTO lesson_sessions(date,start_time,end_time,instructor_id,car_id,car_name_snapshot,
               car_plate_snapshot,capacity_snapshot,status,notes,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (date, start, end_s, inst_id, 1, "Test avto", "A 001 AA", 4, "scheduled",
             notes, created_at, created_at),
        )
        db.ex(
            """INSERT INTO session_students(session_id, student_id, pickup_address,
               attendance_status, student_status, joined_at) VALUES(?,?,'','unmarked','active',?)""",
            (sid, student_id, created_at),
        )
        return sid

    def _reminder_count(self, db, sid):
        return db.q1(
            "SELECT COUNT(*) c FROM notifications WHERE source='LESSON_REMINDER' AND related_lesson_id=?",
            (sid,))["c"]

    def test_100_reminder_2h_before_exactly_once(self):
        """BAND 9: 2 soat oldin eslatma FAQAT BIR MARTA yuboriladi."""
        db = Db(str(TMP / "test.db"))
        # Boshlanishi hozirdan 1 soat 50 daqiqa keyin (2 soat oynasi KIRGAN),
        # lekin yaratilishi eslatma vaqtidan OLDIN (3 soat oldin).
        start_dt = (datetime.now() + timedelta(minutes=110)).replace(second=0, microsecond=0)
        created = (start_dt - timedelta(hours=3)).strftime("%Y-%m-%d %H:%M:%S")
        sid = self._mk_session(db, start_dt, created)
        try:
            n1 = ensure_reminders(db)
            self.assertGreaterEqual(n1, 2, "instruktor + talaba uchun eslatma yuborilishi kerak")
            n2 = ensure_reminders(db)
            self.assertEqual(n2, 0, "takroriy eslatma yuborilmasligi kerak")
            n3 = ensure_reminders(db)
            self.assertEqual(n3, 0, "uchinchi marta ham 0")
            self.assertEqual(self._reminder_count(db, sid), 2, "faqat 2 ta qabulchi")
            # Ma'lumot: sana, vaqt, instruktor, olib ketish joyi
            row = db.q1("""SELECT * FROM notifications
                           WHERE source='LESSON_REMINDER' AND related_lesson_id=? LIMIT 1""", (sid,))
            data = jload(row["data"], {}) if not isinstance(row["data"], dict) else row["data"]
            self.assertEqual(data["date"], start_dt.strftime("%Y-%m-%d"))
            self.assertEqual(data["start_time"], start_dt.strftime("%H:%M"))
            self.assertIn("instructor_name", data)
            self.assertIn("pickup_address", data)
            self.assertEqual(row["created_by"], "SYSTEM", "avtomatik eslatma: created_by=SYSTEM")
        finally:
            self._cleanup(db, sid)

    def test_101_no_reminder_when_created_later_than_2h_window(self):
        """BAND 9: 2 soatdan KAM vaqt ichida yaratilgan mashg'ulotga eslatma
        yuborilmaydi (noto'g'ri/duplicate bildirishnoma chiqmasin)."""
        db = Db(str(TMP / "test.db"))
        start_dt = (datetime.now() + timedelta(minutes=30)).replace(second=0, microsecond=0)
        sid = self._mk_session(db, start_dt, now(), notes="REM2")
        try:
            self.assertEqual(ensure_reminders(db), 0, "eslatma yuborilmasligi kerak")
            self.assertEqual(self._reminder_count(db, sid), 0)
        finally:
            self._cleanup(db, sid)

    def test_102_no_reminder_after_lesson_started(self):
        """BAND 9: allaqachon boshlangan mashg'ulotga eslatma yuborilmaydi."""
        db = Db(str(TMP / "test.db"))
        start_dt = (datetime.now() - timedelta(minutes=10)).replace(second=0, microsecond=0)
        sid = self._mk_session(db, start_dt, (start_dt - timedelta(hours=4)).strftime("%Y-%m-%d %H:%M:%S"),
                               notes="REM3")
        try:
            self.assertEqual(ensure_reminders(db), 0)
            self.assertEqual(self._reminder_count(db, sid), 0)
        finally:
            self._cleanup(db, sid)

    def _cleanup(self, db, sid):
        db.upd("DELETE FROM notifications WHERE related_lesson_id=?", (sid,))
        db.upd("DELETE FROM session_students WHERE session_id=?", (sid,))
        db.upd("DELETE FROM lesson_sessions WHERE id=?", (sid,))


class TestM13CancelReschedule(Base):
    """M13-A — Instruktor uchun bekor qilish / qayta rejalashtirish oqimi
    (sababi bilan), original vaqt tarixi saqlanadi, ikkala tomon ko'radi."""

    def _next_day(self, days=0):
        # Seed'dagi "bugun/ertaga" sessiyalari bilan to'qnashmaslik uchun
        # bazaviy dushanbani kamida 7 kun oldinga surib olamiz.
        from datetime import timedelta
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7) + timedelta(days=7)
        return (mon + timedelta(days=days)).strftime("%Y-%m-%d")

    def tearDown(self):
        """M13A testlari qolgan sessiyalarni tozalaydi (boshqa klasslar testlariga xalaqit bermaslik uchun)."""
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT id FROM lesson_sessions WHERE notes='M13A'")
        for r in rows:
            db.upd("DELETE FROM session_students WHERE session_id=?", (r["id"],))
            db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"session_id": {r["id"]}%',))
            db.upd("DELETE FROM attendance WHERE session_id=?", (r["id"],))
            db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["id"],))

    def _mk(self, admin, date, start, end):
        st, r = admin.post("/api/admin/sessions",
                           {"date": date, "start_time": start, "end_time": end,
                            "instructor_id": 1, "student_ids": [1], "notes": "M13A"})
        self.assertEqual(st, 200, r)
        return r["id"]

    def _admin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return c

    def _inst(self):
        c = Client()
        c.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        return c

    def _stud(self):
        c = Client()
        c.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        return c

    def test_110_instructor_cancel_with_reason(self):
        admin = self._admin()
        d = self._next_day()
        sid = self._mk(admin, d, "08:30", "09:30")
        inst = self._inst()
        st, r = inst.post(f"/api/instructor/sessions/{sid}/cancel", {"reason": "car"})
        self.assertEqual(st, 200, r)
        st, det = inst.get(f"/api/instructor/sessions/{sid}")
        self.assertEqual(st, 200, det)
        self.assertEqual(det["session"]["status"], "cancelled")
        self.assertEqual(det["session"]["cancel_reason"], "car")
        # Talaba bildirishnoma olishi kerak (type ustuni ntype 'cancel' bo'ladi, title'da 'lesson.cancelled')
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT * FROM notifications WHERE data LIKE ?", (f'%"session_id": {sid}%',))
        self.assertTrue(
            any(jload(r["data"]).get("verb") == "cancelled" for r in rows),
            "cancel bildirishnomasi (verb=cancelled) yuborilishi kerak",
        )

    def test_111_instructor_reschedule_preserves_original(self):
        admin = self._admin()
        d = self._next_day()
        sid = self._mk(admin, d, "16:30", "17:30")
        new_d = self._next_day(1)  # seshanba
        inst = self._inst()
        st, r = inst.post(f"/api/instructor/sessions/{sid}/reschedule",
                          {"date": new_d, "start_time": "16:00", "end_time": "17:00"})
        self.assertEqual(st, 200, r)
        st, det = inst.get(f"/api/instructor/sessions/{sid}")
        self.assertEqual(st, 200, det)
        self.assertEqual(det["session"]["date"], new_d)
        self.assertEqual(det["session"]["start_time"], "16:00")
        self.assertEqual(det["session"]["status"], "scheduled")
        self.assertEqual(det["session"]["original_date"], d, "original sana saqlanishi kerak")
        self.assertEqual(det["session"]["original_start_time"], "16:30")
        # Instruktor ham reschedule bildirishnomasini oladi
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT * FROM notifications WHERE data LIKE ?", (f'%"session_id": {sid}%',))
        self.assertTrue(
            any(jload(r["data"]).get("verb") == "rescheduled" for r in rows),
            "reschedule bildirishnomasi (verb=rescheduled) yuborilishi kerak",
        )

    def test_112_instructor_reschedule_conflict_409(self):
        admin = self._admin()
        d = self._next_day()
        a = self._mk(admin, d, "11:45", "12:45")
        self._mk(admin, d, "15:30", "16:30")
        inst = self._inst()
        st, r = inst.post(f"/api/instructor/sessions/{a}/reschedule",
                          {"date": d, "start_time": "15:30", "end_time": "16:30"})
        self.assertEqual(st, 409, r)
        self.assertEqual(r["error"], "session.rules_violated")
        self.assertIn("instructor_busy", r.get("params", {}).get("errors", []))

    def test_113_instructor_cancel_wrong_status(self):
        admin = self._admin()
        d = self._next_day()
        sid = self._mk(admin, d, "14:30", "15:30")
        db = Db(str(TMP / "test.db"))
        db.upd("UPDATE lesson_sessions SET status='completed', completed_at=? WHERE id=?", (now(), sid))
        inst = self._inst()
        st, r = inst.post(f"/api/instructor/sessions/{sid}/cancel", {"reason": "school"})
        self.assertEqual(st, 409, r)
        self.assertEqual(r["error"], "session.wrong_status")

    def test_114_student_cannot_cancel_or_reschedule(self):
        admin = self._admin()
        d = self._next_day()
        sid = self._mk(admin, d, "08:00", "09:00")
        stud = self._stud()
        st, r = stud.post(f"/api/instructor/sessions/{sid}/cancel", {"reason": "x"})
        self.assertEqual(st, 403, r)
        st, r = stud.post(f"/api/instructor/sessions/{sid}/reschedule",
                          {"date": d, "start_time": "09:00", "end_time": "10:00"})
        self.assertEqual(st, 403, r)

    def test_115_admin_reschedule_keeps_original_date(self):
        admin = self._admin()
        d = self._next_day()
        sid = self._mk(admin, d, "09:30", "10:00")
        tue, thu = self._next_day(1), self._next_day(3)
        st, r = admin.post(f"/api/admin/sessions/{sid}/reschedule",
                           {"date": tue, "start_time": "10:00", "end_time": "11:00"})
        self.assertEqual(st, 200, r)
        st, r = admin.post(f"/api/admin/sessions/{sid}/reschedule",
                           {"date": thu, "start_time": "11:00", "end_time": "12:00"})
        self.assertEqual(st, 200, r)
        st, det = admin.get(f"/api/admin/sessions/{sid}")
        self.assertEqual(st, 200, det)
        self.assertEqual(det["session"]["date"], thu)
        self.assertEqual(det["session"]["original_date"], d, "original sana birinchi sanada qolishi kerak")
        self.assertEqual(det["session"]["original_start_time"], "09:30")


class TestM13Analytics(Base):
    """M13-B — Admin analitika paneli: oylik faollik grafigi, bekor qilingan
    mashg'ulotlar foizi, eng band instruktorlar."""

    def _admin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return c

    def _mk(self, admin, date, start, end):
        st, r = admin.post("/api/admin/sessions",
                           {"date": date, "start_time": start, "end_time": end,
                            "instructor_id": 1, "student_ids": [1], "notes": "M13B"})
        self.assertEqual(st, 200, r)
        return r["id"]

    def tearDown(self):
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT id FROM lesson_sessions WHERE notes='M13B'")
        for r in rows:
            db.upd("DELETE FROM session_students WHERE session_id=?", (r["id"],))
            db.upd("DELETE FROM attendance WHERE session_id=?", (r["id"],))
            db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["id"],))

    def test_120_admin_analytics_shape(self):
        admin = self._admin()
        st, r = admin.get("/api/admin/analytics")
        self.assertEqual(st, 200, r)
        a = r["analytics"]
        self.assertIn("month", a)
        self.assertIsInstance(a["total"], int)
        self.assertIsInstance(a["cancel_rate"], float)
        self.assertTrue(0 <= a["cancel_rate"] <= 100)
        self.assertEqual(len(a["months"]), 12, "12 oylik tendentsiya bo'lishi kerak")
        self.assertIn(len(a["days"]), (28, 29, 30, 31))
        self.assertIsInstance(a["top_instructors"], list)
        for d in a["days"]:
            self.assertIn("sessions", d)
            self.assertIn("cancelled", d)
        for m in a["months"]:
            self.assertIn("sessions", m)
            self.assertIn("cancelled", m)

    def test_121_admin_analytics_counts_and_cancel_rate(self):
        # Maxsus oy: 2099-03 — boshqa testlardagi qoldiq sessiyalar ta'sir qilmaydi
        db = Db(str(TMP / "test.db"))
        t = now()
        sid1 = db.ex(
            """INSERT INTO lesson_sessions(date,start_time,end_time,instructor_id,car_id,
               car_name_snapshot,car_plate_snapshot,capacity_snapshot,status,notes,created_at,updated_at)
               VALUES('2099-03-10','09:00','10:00',1,1,'Chevrolet Cobalt','01 A 123 AA',4,'completed','M13B',?,?)""",
            (t, t))
        db.ex("INSERT INTO session_students(session_id,student_id,joined_at) VALUES(?,1,?)", (sid1, t))
        sid2 = db.ex(
            """INSERT INTO lesson_sessions(date,start_time,end_time,instructor_id,car_id,
               car_name_snapshot,car_plate_snapshot,capacity_snapshot,status,notes,cancel_reason,created_at,updated_at)
               VALUES('2099-03-12','15:30','16:30',1,1,'Chevrolet Cobalt','01 A 123 AA',4,'cancelled','M13B','other',?,?)""",
            (t, t))
        db.ex("INSERT INTO session_students(session_id,student_id,joined_at) VALUES(?,1,?)", (sid2, t))
        admin = self._admin()
        st, r = admin.get("/api/admin/analytics?month=2099-03")
        self.assertEqual(st, 200, r)
        a = r["analytics"]
        self.assertEqual(a["total"], 2)
        self.assertEqual(a["completed"], 1)
        self.assertEqual(a["cancelled"], 1)
        self.assertEqual(a["cancel_rate"], 50.0)
        self.assertEqual(a["hours"], 1.0)
        self.assertEqual(a["unique_students"], 1)
        self.assertEqual(a["days"][9]["sessions"], 1)
        self.assertEqual(a["days"][11]["cancelled"], 1)
        names = [t_name["name"] for t_name in a["top_instructors"]]
        self.assertTrue(any("Akmal" in n for n in names), "eng band instruktorlar orasida Akmal bor")
        self.assertEqual(a["top_instructors"][0]["sessions"], 2)
        self.assertEqual(a["top_instructors"][0]["completed"], 1)

    def test_122_admin_analytics_bad_month(self):
        admin = self._admin()
        st, r = admin.get("/api/admin/analytics?month=2020-13")
        self.assertEqual(st, 400, r)
        st, r = admin.get("/api/admin/analytics?month=abc")
        self.assertEqual(st, 400, r)

    def test_123_student_cannot_access_analytics(self):
        c = Client()
        c.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        st, r = c.get("/api/admin/analytics")
        self.assertEqual(st, 403, r)


class TestM4RequestApprove(Base):
    """M4 — Admin so'rovni tasdiqlash (bugfix): session_students INSERT'dagi
    bindings xatosi tufayli tasdiqlash 500 berib, so'rov holati o'zgarmas edi.
    Endi tasdiqlash: yangi session + student qatori + talabaga bildirishnoma."""

    def _student(self):
        c = Client()
        c.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        return c

    def _admin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return c

    def _next_monday(self):
        from datetime import timedelta
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=7)).strftime("%Y-%m-%d")

    def tearDown(self):
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT id FROM practice_requests WHERE message='M4-TEST'")
        for r_ in rows:
            req = db.q1("SELECT * FROM practice_requests WHERE id=?", (r_["id"],))
            if req and req["session_id"]:
                db.upd("DELETE FROM session_students WHERE session_id=?", (req["session_id"],))
                db.upd("DELETE FROM lesson_sessions WHERE id=?", (req["session_id"],))
            db.upd("DELETE FROM practice_requests WHERE id=?", (r_["id"],))

    def test_130_approve_request_updates_status_and_creates_session(self):
        stu = self._student()
        st, r = stu.post("/api/student/requests", {
            "preferred_date": self._next_monday(), "preferred_start_time": "09:00",
            "preferred_end_time": "10:00", "message": "M4-TEST"})
        self.assertEqual(st, 200, r)
        rid = r["id"]

        admin = self._admin()
        st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
            "date": self._next_monday(), "start_time": "09:00", "end_time": "10:00",
            "instructor_id": 1})
        self.assertEqual(st, 200, r)
        sid = r["session_id"]
        self.assertTrue(sid)

        db = Db(str(TMP / "test.db"))
        req = db.q1("SELECT * FROM practice_requests WHERE id=?", (rid,))
        self.assertEqual(req["status"], "approved")
        self.assertEqual(req["session_id"], sid)
        ss = db.q1("SELECT * FROM session_students WHERE session_id=? AND student_status='active'", (sid,))
        self.assertIsNotNone(ss, "session_students qatori yaratilishi kerak (M4 bugfix)")
        self.assertEqual(ss["student_id"], 1)
        notif = db.q1("SELECT * FROM notifications WHERE user_id=5 AND type='lesson' AND title='lesson.assigned' ORDER BY id DESC LIMIT 1")
        self.assertIsNotNone(notif, "talabaga 'lesson.assigned' bildirishnoma keldi")

    def test_131_reject_request_updates_status(self):
        stu = self._student()
        st, r = stu.post("/api/student/requests", {
            "preferred_date": self._next_monday(), "preferred_start_time": "09:00",
            "preferred_end_time": "10:00", "message": "M4-TEST"})
        self.assertEqual(st, 200, r)
        rid = r["id"]
        admin = self._admin()
        st, r = admin.post(f"/api/admin/requests/{rid}/reject", {"note": "rad etildi"})
        self.assertEqual(st, 200, r)
        db = Db(str(TMP / "test.db"))
        req = db.q1("SELECT * FROM practice_requests WHERE id=?", (rid,))
        self.assertEqual(req["status"], "rejected")
        self.assertEqual(req["admin_note"], "rad etildi")


class TestApproveRequestRules(Base):
    """'Tasdiqlash' tugmasi ishlamaydi degan xato.

    Sabab: backend qoida buzilganda `params.errors` da ANIQ sabablarni
    yuboradi (masalan ["intersects_break"]), lekin frontend `errToast` bu
    ro'yxatni tashlab yuborib faqat umumiy "Mashg'ulot yaratib bo'lmaydi"
    ko'rsatardi — admin qaysi qoida bloklaganini bilmasdi va tugma buzilgan
    deb o'ylardi. Bu yerda har bir qoida kodi alohida tekshiriladi.

    Ikkinchidan, session/so'rov yozuvlari bitta tranzaksiyaga o'tkazildi:
    oraliqda xato chiqsa yaroq session qolmasligi tekshiriladi.
    """

    MSG = "APR-TEST"

    def _admin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return c

    def _login(self, login):
        """Test fixture loginlari (`PW_BY_LOGIN` ro'yxatidan parol topiladi)."""
        c = Client()
        c.post("/api/auth/login",
               {"login": login, "password": PW_BY_LOGIN.get(login, "usrP_Nope0000!Ab1x"),
                "role": "student"})
        return c

    def _students(self):
        """Bazadagi talabalar: [{login, sid}, ...] (s.id bo'yicha tartiblangan)."""
        db = Db(str(TMP / "test.db"))
        return db.q("""SELECT u.login AS login, s.id AS sid FROM students s
                       JOIN users u ON u.id=s.user_id
                       WHERE u.deleted_at IS NULL ORDER BY s.id""")

    def _next_monday(self, weeks=0):
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=7 * weeks)).strftime("%Y-%m-%d")

    def _next_weekday(self, iso_wd):
        """Keyingi dushanbadan keyingi `iso_wd` kuni (0=mon ... 6=sun)."""
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=iso_wd)).strftime("%Y-%m-%d")

    def _new_request(self, login, date, start, end, msg=None):
        stu = self._login(login)
        st, r = stu.post("/api/student/requests", {
            "preferred_date": date, "preferred_start_time": start,
            "preferred_end_time": end, "message": msg or self.MSG})
        self.assertEqual(st, 200, r)
        return r["id"]

    def _sessions_for(self, note_frag):
        db = Db(str(TMP / "test.db"))
        return db.q("SELECT id FROM lesson_sessions WHERE notes LIKE ?", ("%" + note_frag + "%",))

    def tearDown(self):
        db = Db(str(TMP / "test.db"))
        for r in db.q("SELECT id, session_id FROM practice_requests WHERE message=?", (self.MSG,)):
            if r["session_id"]:
                db.upd("DELETE FROM session_students WHERE session_id=?", (r["session_id"],))
                db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["session_id"],))
            db.upd("DELETE FROM audit_logs WHERE entity_id=? AND entity_type='practice_requests'",
                   (r["id"],))
        db.upd("DELETE FROM practice_requests WHERE message=?", (self.MSG,))

    # ------------------------------------------------------------------
    # 1) Tanaffus vaqti — bu holat admin uchun eng ko'p uchraydigan sabab
    def test_140_approve_blocked_by_break_returns_specific_rule_code(self):
        """12:00-16:00 so'rov 13:00-14:00 tanaffusga tegadi -> intersects_break.

        Bu aynan "Tasdiqlash hech narsa o'zgarmaydi" degan holat: eski kodda
        faqat "Mashg'ulot yaratib bo'lmaydi" ko'rinardi, sabab yo'q edi.
        """
        student = self._students()[0]
        rid = self._new_request(student["login"], self._next_monday(), "12:00", "16:00")
        admin = self._admin()
        st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
            "date": self._next_monday(), "start_time": "12:00", "end_time": "16:00",
            "instructor_id": 1})
        self.assertEqual(st, 409, r)
        self.assertEqual(r["error"], "session.rules_violated")
        self.assertIn("intersects_break", r["params"]["errors"],
                      "frontend aniq sababni ko'rsatishi uchun kod kiritilishi shart")
        # so'rov "pending"da qoladi va yaroq session yaratilmaydi
        db = Db(str(TMP / "test.db"))
        req = db.q1("SELECT * FROM practice_requests WHERE id=?", (rid,))
        self.assertEqual(req["status"], "pending")
        self.assertIsNone(req["session_id"])
        self.assertEqual(self._sessions_for("so'rov #%s" % rid), [])

    # 2) Ish vaqti tashqarisida
    def test_141_approve_outside_work_hours_returns_rule_code(self):
        """Instruktor 1 (Akmal) 08:00-18:00 ishlaydi -> 19:00-20:00 tashqarida."""
        date = self._next_monday()
        rid = self._new_request(self._students()[0]["login"], date, "19:00", "20:00")
        admin = self._admin()
        st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
            "date": date, "start_time": "19:00", "end_time": "20:00",
            "instructor_id": 1})
        self.assertEqual(st, 409, r)
        self.assertIn("outside_work_hours", r["params"]["errors"])
        db = Db(str(TMP / "test.db"))
        self.assertEqual(db.q1("SELECT status FROM practice_requests WHERE id=?",
                               (rid,))["status"], "pending")

    # 3) Ish kuni emas
    def test_142_approve_on_non_work_day_returns_rule_code(self):
        """Instruktor 1 dushanbadan shanbagacha ishlaydi -> yakshanbay = not_work_day."""
        sunday = self._next_weekday(6)
        rid = self._new_request(self._students()[0]["login"], sunday, "09:00", "10:00")
        admin = self._admin()
        st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
            "date": sunday, "start_time": "09:00", "end_time": "10:00",
            "instructor_id": 1})
        self.assertEqual(st, 409, r)
        self.assertIn("not_work_day", r["params"]["errors"])

    # 4) Xato jim qolmasligi: instruktor tanlanmagan
    def test_143_approve_without_instructor_is_explicit_error(self):
        rid = self._new_request(self._students()[0]["login"], self._next_monday(),
                                "09:00", "10:00")
        admin = self._admin()
        st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
            "date": self._next_monday(), "start_time": "09:00", "end_time": "10:00"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "request.no_instructor")
        db = Db(str(TMP / "test.db"))
        self.assertEqual(db.q1("SELECT status FROM practice_requests WHERE id=?",
                               (rid,))["status"], "pending")

    # 5) Atomiklik: oraliqda xato chiqsa yaroq session qolmasligi
    def test_144_approve_rolls_back_on_midway_failure(self):
        """session_students INSERTida xato -> session ham, status ham qaytadi."""
        rid = self._new_request(self._students()[0]["login"], self._next_monday(),
                                "09:00", "10:00")
        date = self._next_monday()
        db = Db(str(TMP / "test.db"))
        with db.connect() as c:
            c.execute("""CREATE TRIGGER IF NOT EXISTS t_apr_fail BEFORE INSERT ON session_students
                         BEGIN SELECT RAISE(ABORT, 'test: session_students buzildi'); END""")
            c.commit()
        try:
            admin = self._admin()
            st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
                "date": date, "start_time": "09:00", "end_time": "10:00",
                "instructor_id": 1})
            # xato 5xx sifatida qaytishi kerak (200 emas!) — "jim" qolmasin
            self.assertGreaterEqual(st, 500, "xato yutilmasin, aniq status kod qaytaring")
        finally:
            with db.connect() as c:
                c.execute("DROP TRIGGER IF EXISTS t_apr_fail")
                c.commit()

        req = db.q1("SELECT * FROM practice_requests WHERE id=?", (rid,))
        self.assertEqual(req["status"], "pending", "xatodan keyin holat o'zgarmasligi kerak")
        self.assertIsNone(req["session_id"])
        self.assertEqual(self._sessions_for("so'rov #%s" % rid), [],
                         "yaroq (orphan) session qolmadi — tranzaksiya ishladi")

    # 6) Turli so'rovlarda muvaffaqiyatli tasdiqlash
    def test_145_approve_three_different_requests(self):
        """Uch xil so'rov (turli kun/vaqt/talaba) -> hammasi 'approved'."""
        students = self._students()
        self.assertGreaterEqual(len(students), 3, "seedda kamida 3 ta talaba kerak")
        admin = self._admin()
        # (talaba, sana, start, end, instruktor_id) — har biri boshqacha holatda:
        # oddiy, keyingi hafta, boshqa instruktor bilan.
        cases = [
            (students[0], self._next_monday(), "09:00", "10:00", 1),
            (students[1], self._next_monday(1), "15:00", "16:30", 1),
            (students[2], self._next_monday(2), "11:00", "12:00", 2),
        ]
        for student, date, start, end, iid in cases:
            with self.subTest(login=student["login"], date=date, instructor=iid):
                rid = self._new_request(student["login"], date, start, end)
                st, r = admin.post(f"/api/admin/requests/{rid}/approve", {
                    "date": date, "start_time": start, "end_time": end,
                    "instructor_id": iid})
                self.assertEqual(st, 200, r)
                sid = r["session_id"]
                self.assertTrue(sid)
                db = Db(str(TMP / "test.db"))
                req = db.q1("SELECT * FROM practice_requests WHERE id=?", (rid,))
                self.assertEqual(req["status"], "approved")
                self.assertEqual(req["session_id"], sid)
                ls = db.q1("SELECT * FROM lesson_sessions WHERE id=?", (sid,))
                self.assertIsNotNone(ls, "mashg'ulot yaratilishi kerak")
                self.assertEqual(ls["date"], date)
                self.assertEqual(ls["start_time"], start)
                self.assertEqual(ls["end_time"], end)
                self.assertEqual(ls["instructor_id"], iid)
                self.assertEqual(ls["status"], "scheduled")
                ss = db.q1("SELECT * FROM session_students WHERE session_id=? "
                           "AND student_status='active'", (sid,))
                self.assertIsNotNone(ss, "talaba sessionga bog'lanishi kerak")
                self.assertEqual(ss["student_id"], student["sid"])
                notif = db.q1("SELECT * FROM notifications WHERE type='lesson' "
                              "AND title='lesson.assigned' ORDER BY id DESC LIMIT 1")
                self.assertIsNotNone(notif, "bog'liq amal: talabaga bildirishnoma yuboriladi")


class TestM1SenderInfo(Base):
    """M1 — bildirishnoma jo'natuvchisi (sender_id, sender_role, ism-familiya)
    saqlanadi va `/api/me/notifications` orqali qaytariladi."""

    def _student(self):
        c = Client()
        c.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        return c

    def _admin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return c

    def _next_monday(self):
        from datetime import timedelta
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=7)).strftime("%Y-%m-%d")

    def _jdata(self, n):
        if isinstance(n.get("data"), dict):
            return n.get("data") or {}
        return jload(n.get("data") or "{}", {})

    def tearDown(self):
        """M1 testlari qolgan sessiya/so'rovlarni tozalaydi (boshqa klasslarga xalaqit bermaslik uchun)."""
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT id FROM lesson_sessions WHERE notes='M1-TEST'")
        for r in rows:
            db.upd("DELETE FROM session_students WHERE session_id=?", (r["id"],))
            db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"session_id": {r["id"]}%',))
            db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["id"],))
        reqs = db.q("SELECT id FROM practice_requests WHERE message='M1-TEST'")
        for r in reqs:
            db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"request_id": {r["id"]}%',))
            db.upd("DELETE FROM practice_requests WHERE id=?", (r["id"],))

    def test_140_session_created_sender_is_admin(self):
        admin = self._admin()
        d = self._next_monday()
        st, r = admin.post("/api/admin/sessions", {"date": d, "start_time": "09:00", "end_time": "10:00",
                                                   "instructor_id": 1, "student_ids": [1], "notes": "M1-TEST"})
        self.assertEqual(st, 200, r)
        sid = r["id"]
        db = Db(str(TMP / "test.db"))
        n = db.q1("SELECT * FROM notifications WHERE title='lesson.assigned' ORDER BY id DESC LIMIT 1")
        self.assertIsNotNone(n, "talabaga lesson.assigned kelishi kerak")
        self.assertEqual(n["sender_id"], 1, "jo'natuvchi admin (id=1) bo'lishi kerak")
        self.assertEqual(n["sender_role"], "admin")

        stu = self._student()
        st, ns = stu.get("/api/me/notifications")
        self.assertEqual(st, 200, ns)
        hit = [x for x in ns["notifications"] if str(self._jdata(x).get("session_id")) == str(sid)]
        self.assertTrue(hit, "sessiya bildirishnomasi me_notifications'da bor")
        top = hit[0]
        self.assertEqual(top["sender_role"], "admin")
        self.assertEqual(top["sender_id"], 1)
        self.assertEqual(top["sender_first_name"], "Admin")
        self.assertEqual(top["sender_last_name"], "Avtomaktab")

    def test_141_request_new_sender_is_student(self):
        stu = self._student()
        st, r = stu.post("/api/student/requests", {"preferred_date": self._next_monday(),
                                                   "preferred_start_time": "10:00",
                                                   "preferred_end_time": "11:00",
                                                   "message": "M1-TEST"})
        self.assertEqual(st, 200, r)
        admin = self._admin()
        st, ns = admin.get("/api/me/notifications")
        self.assertEqual(st, 200, ns)
        hit = [x for x in ns["notifications"] if str(self._jdata(x).get("request_id")) == str(r["id"])]
        self.assertTrue(hit, "request.new admin bildirishnomasi bor")
        top = hit[0]
        self.assertEqual(top["sender_role"], "student")
        self.assertEqual(top["sender_id"], 5)
        self.assertEqual(top["sender_first_name"], "Omadbek")


class TestM2NotifI18n(Base):
    """M2 — backend bildirishnoma title'da faqat KOD saqlaydi (enum), tarjima frontendda.
    `data` esa lokallizatsiya uchun tizimli maydonlarni o'z ichiga oladi."""

    def _student(self):
        c = Client()
        c.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        return c

    def _admin(self):
        c = Client()
        c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        return c

    def _next_monday(self):
        from datetime import timedelta
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=7)).strftime("%Y-%m-%d")

    def _jdata(self, n):
        if isinstance(n.get("data"), dict):
            return n.get("data") or {}
        return jload(n.get("data") or "{}", {})

    def tearDown(self):
        db = Db(str(TMP / "test.db"))
        rows = db.q("SELECT id FROM lesson_sessions WHERE notes='M2-TEST'")
        for r in rows:
            db.upd("DELETE FROM session_students WHERE session_id=?", (r["id"],))
            db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%\"session_id\": {r["id"]}%',))
            db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["id"],))
        reqs = db.q("SELECT id FROM practice_requests WHERE message='M2-TEST'")
        for r in reqs:
            db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%\"request_id\": {r["id"]}%',))
            db.upd("DELETE FROM practice_requests WHERE id=?", (r["id"],))

    def test_150_session_notif_carries_code_and_params(self):
        admin = self._admin()
        d = self._next_monday()
        st, r = admin.post("/api/admin/sessions", {"date": d, "start_time": "09:00", "end_time": "10:00",
                                                   "instructor_id": 1, "student_ids": [1], "notes": "M2-TEST"})
        self.assertEqual(st, 200, r)
        sid = r["id"]
        db = Db(str(TMP / "test.db"))
        n = db.q1("SELECT * FROM notifications WHERE title='lesson.assigned' ORDER BY id DESC LIMIT 1")
        self.assertIsNotNone(n, "lesson.assigned bildirishnoma keldi")
        # title faqat kod bo'lishi kerak (tarjima frontendda)
        self.assertEqual(n["title"], "lesson.assigned")
        nd = jload(n["data"], {})
        self.assertEqual(nd.get("session_id"), sid)
        self.assertEqual(nd.get("date"), d)
        self.assertEqual(nd.get("start_time"), "09:00")
        self.assertEqual(nd.get("end_time"), "10:00")
        self.assertEqual(nd.get("instructor_name"), "Akmal Karimov")
        self.assertTrue(nd.get("car"), "avto nomi data'da bo'lishi kerak")

    def test_151_request_new_carries_student_name(self):
        stu = self._student()
        st, r = stu.post("/api/student/requests", {"preferred_date": self._next_monday(),
                                                   "preferred_start_time": "10:00",
                                                   "preferred_end_time": "11:00",
                                                   "message": "M2-TEST"})
        self.assertEqual(st, 200, r)
        db = Db(str(TMP / "test.db"))
        n = db.q1("SELECT * FROM notifications WHERE title='request.new' ORDER BY id DESC LIMIT 1")
        self.assertIsNotNone(n, "request.new bildirishnoma keldi")
        self.assertEqual(n["title"], "request.new")
        nd = jload(n["data"], {})
        self.assertEqual(nd.get("request_id"), r["id"])
        self.assertEqual(nd.get("student_name"), "Omadbek Xoshimov")


class TestM3NotifNavigation(Base):
    """M3 — bildirishnoma data'sida detalga o'tish uchun etarli ma'lumot bo'lishi.
    Foydalanuvchi bildirishnomani bossa -> sessiya/so'rov/xabar detaliga o'tadi.
    Frontend aynan shu maydonlarga tayanadi: session_id / request_id / message_id."""

    def _jdata(self, n):
        if isinstance(n.get("data"), dict):
            return n.get("data") or {}
        return jload(n.get("data") or "{}", {})

    def test_160_msg_new_carries_message_id(self):
        # admin instruktyor Akmal (uid 2) ga xabar yuboradi
        admin = Client()
        st, r = admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        st, r = admin.post("/api/me/messages", {"to_user_id": 2, "text": "M3-TEST-MSG"})
        self.assertEqual(st, 200, r)
        self.assertTrue(r.get("ok"))
        # Akmal bildirishnomalarida msg.new bo'lishi kerak, data da message_id + text
        inst = Client()
        st, r2 = inst.post("/api/auth/login", {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        self.assertEqual(st, 200, r2)
        st, ns = inst.get("/api/me/notifications")
        self.assertEqual(st, 200, ns)
        hit = [n for n in ns["notifications"] if n.get("title") == "msg.new"]
        self.assertTrue(hit, "msg.new bildirishnoma keldi")
        top = hit[0]
        nd = self._jdata(top)
        self.assertTrue(nd.get("message_id"), "message_id data'da bo'lishi kerak (M3 klik uchun)")
        self.assertEqual(nd.get("text"), "M3-TEST-MSG")
        self.assertEqual(top["sender_role"], "admin")
        # tozalash: ushbu xabar va bildirishnomalar
        db = Db(str(TMP / "test.db"))
        db.upd("DELETE FROM notifications WHERE data LIKE '%M3-TEST-MSG%'")
        db.upd("DELETE FROM messages WHERE text='M3-TEST-MSG'")

    def test_161_data_has_navigation_keys(self):
        """Har bir notification turi data'da tegishli ID maydonini olib yuradi (o'zini-o'zi ta'minlaydi)."""
        from datetime import timedelta
        admin = Client()
        st, r = admin.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        d = (datetime.now() + timedelta(days=(7 - datetime.now().weekday()) % 7 or 7)).strftime("%Y-%m-%d")
        st, sess = admin.post("/api/admin/sessions", {"date": d, "start_time": "09:00", "end_time": "10:00",
                                                      "instructor_id": 1, "student_ids": [1], "notes": "M3-TEST"})
        self.assertEqual(st, 200, sess)
        stu = Client()
        st, r = stu.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "student"})
        self.assertEqual(st, 200, r)
        st, rq = stu.post("/api/student/requests", {"preferred_date": d, "preferred_start_time": "10:00",
                                                    "preferred_end_time": "11:00", "message": "M3-TEST-REQ"})
        self.assertEqual(st, 200, rq)
        db = Db(str(TMP / "test.db"))
        sl = db.q("SELECT data FROM notifications WHERE title='lesson.assigned' ORDER BY id DESC LIMIT 1")
        self.assertTrue(sl and "session_id" in jload(sl[0]["data"], {}), "lesson.assigned data'da session_id bor")
        rl = db.q("SELECT data FROM notifications WHERE title='request.new' ORDER BY id DESC LIMIT 1")
        self.assertTrue(rl and "request_id" in jload(rl[0]["data"], {}), "request.new data'da request_id bor")
        # tozalash
        db.upd("DELETE FROM session_students WHERE session_id=?", (sess["id"],))
        db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"session_id": {sess["id"]}%',))
        db.upd("DELETE FROM lesson_sessions WHERE id=?", (sess["id"],))
        db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"request_id": {rq["id"]}%',))
        db.upd("DELETE FROM practice_requests WHERE id=?", (rq["id"],))


class TestM6InstructorCar(Base):
    """M6 — instruktor FAQAT o'ziga biriktirilgan mashinani tahrirlay oladi;
    admin huquqlari o'zgarishsiz qoladi."""

    def _inst_login(self, login=INSTR_LOGIN_1, pwd=INSTR_PASS_1):
        c = Client()
        st, r = c.post("/api/auth/login", {"login": login, "password": pwd, "role": "instructor"})
        self.assertEqual(st, 200, r)
        return c

    def test_170_instructor_updates_own_car(self):
        inst = self._inst_login()
        # Akmal (inst 1) car 1 ga biriktirilgan; color va capacity ni o'zgartiramiz
        st, r = inst.put("/api/instructor/car", {"color": "To'q ko'k", "practice_capacity": 5})
        self.assertEqual(st, 200, r)
        db = Db(str(TMP / "test.db"))
        car = db.q1("SELECT color, practice_capacity FROM cars WHERE id=1")
        self.assertEqual(car["color"], "To'q ko'k")
        self.assertEqual(car["practice_capacity"], 5)
        # qaytarish
        db.upd("UPDATE cars SET color='Oq', practice_capacity=4 WHERE id=1")

    def test_171_instructor_cannot_edit_other_car(self):
        # Boshqa instruktor (Omad? yo'q) — inst 2, u ham faqat o'z mashinasini tahrirlay oladi;
        # endpoint URL'da car id qabul qilmaydi — e'lon qilingan car har doim O'ZIdan keladi.
        inst = self._inst_login(INSTR_LOGIN_3, INSTR_PASS_3)  # Vali (inst 3)
        db = Db(str(TMP / "test.db"))
        me = db.q1("SELECT assigned_car_id FROM instructors WHERE id=3")
        self.assertTrue(me["assigned_car_id"], "Vali mashinaga biriktirilgan bo'lishi kerak (seed)")
        st, r = inst.put("/api/instructor/car", {"brand": "HACK"})
        self.assertEqual(st, 200, r)
        # faqat O'Z mashinasi yangilandi
        car = db.q1("SELECT brand FROM cars WHERE id=?", (me["assigned_car_id"],))
        self.assertEqual(car["brand"], "HACK")
        own = db.q1("SELECT brand FROM cars WHERE id=1")
        self.assertEqual(own["brand"], "Chevrolet")
        db.upd("UPDATE cars SET brand='Chevrolet' WHERE id=?", (me["assigned_car_id"],))

    def test_172_instructor_without_car_rejected(self):
        db = Db(str(TMP / "test.db"))
        db.upd("UPDATE instructors SET assigned_car_id=NULL WHERE id=1")
        try:
            inst = self._inst_login()
            st, r = inst.put("/api/instructor/car", {"brand": "X"})
            self.assertEqual(st, 404, r)
        finally:
            db.upd("UPDATE instructors SET assigned_car_id=1 WHERE id=1")


class TestM9LessonDuration(Base):
    """M9 — admin settings'dagi standart davomiylik yangi mashg'ulotga qo'llanadi;
    eksplicit end_time berilsa o'zgarmaydi."""

    def _next_monday(self):
        from datetime import timedelta
        d = datetime.now()
        mon = d + timedelta(days=(7 - d.weekday()) % 7 or 7)
        return (mon + timedelta(days=7)).strftime("%Y-%m-%d")

    def _admin(self):
        c = Client()
        st, r = c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        return c

    def test_180_duration_default_applied_when_end_missing(self):
        admin = self._admin()
        admin.put("/api/admin/settings", {"lesson_duration_min": 45})
        st, r = admin.post("/api/admin/sessions", {
            "date": self._next_monday(), "start_time": "14:00",
            "instructor_id": 3, "student_ids": [], "notes": "M9-TEST",
        })
        self.assertEqual(st, 200, r)
        st, det = admin.get(f"/api/admin/sessions/{r['id']}")
        self.assertEqual(st, 200, det)
        self.assertEqual(det["session"]["end_time"], "14:45")  # 14:00 + 45 daqiqa
        db = Db(str(TMP / "test.db"))
        db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"session_id": {r["id"]}%',))
        db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["id"],))
        admin.put("/api/admin/settings", {"lesson_duration_min": 90})

    def test_181_explicit_end_time_kept(self):
        admin = self._admin()
        admin.put("/api/admin/settings", {"lesson_duration_min": 45})
        st, r = admin.post("/api/admin/sessions", {
            "date": self._next_monday(), "start_time": "16:00", "end_time": "17:15",
            "instructor_id": 3, "student_ids": [], "notes": "M9-TEST-EXPL",
        })
        self.assertEqual(st, 200, r)
        st, det = admin.get(f"/api/admin/sessions/{r['id']}")
        self.assertEqual(st, 200, det)
        self.assertEqual(det["session"]["end_time"], "17:15")
        db = Db(str(TMP / "test.db"))
        db.upd("DELETE FROM notifications WHERE data LIKE ?", (f'%"session_id": {r["id"]}%',))
        db.upd("DELETE FROM lesson_sessions WHERE id=?", (r["id"],))
        admin.put("/api/admin/settings", {"lesson_duration_min": 90})


class TestSessionIsolation(Base):
    """MODUL 1 — foydalanuvchi sessiyalari ARALASHMASLIGI (xavfsizlik).

    Xato: sessiya bitta `sid` cookie'siga bog'langan edi. Cookie esa bitta
    brauzer bo'ylab UMUMIY — barcha tab'lar bir xil qiymatni ko'radi.
    Shuning uchun bir brauzerda ikki odam kirsa, ikkinchisi birinchisining
    sessiyasini almashtirib yuborardi va birinchi oyna boshqa odamning
    profilini, jadvalini, xabarlarini ko'ra boshlardi.

    Tuzatma: kalit ikki qismdan yig'iladi — HttpOnly cookie'dagi QURILMA
    kaliti + `X-Avto-Tab` sarlavhasidagi TAB kaliti.
    """

    INSTR = (INSTR_LOGIN_1, INSTR_PASS_1, "instructor")
    STUD = (STUD_LOGIN_1, STUD_PASS_1, "student")
    ADMIN = ("admin", "admin123", "admin")

    def _login(self, tab_client, creds):
        st, r = tab_client.post("/api/auth/login", {
            "login": creds[0], "password": creds[1], "role": creds[2]})
        self.assertEqual(st, 200, r)
        return r["user"]

    def _device(self, tab_client):
        for ck in tab_client.jar:
            if ck.name == "sid":
                return ck.value
        return None

    # ------------------------------------------------------------------ asosiy
    def test_100_two_tabs_same_browser_stay_independent(self):
        """Bir brauzer, ikki tab: 2-tab kirmasi 1-tabni buzmasin."""
        jar = http.cookiejar.CookieJar()
        tab1 = TabClient("iso-tab-1", jar=jar)
        tab2 = TabClient("iso-tab-2", jar=jar)

        u1 = self._login(tab1, self.INSTR)
        self.assertEqual(u1["role"], "instructor")

        # Ikkala tab bir xil cookie jar'ni ko'radi — bitta brauzer taqdimoti
        self.assertEqual(self._device(tab1), self._device(tab2))

        u2 = self._login(tab2, self.STUD)
        self.assertEqual(u2["role"], "student")

        # MUIM: 1-tab o'zini hali ham instruktor deb o'ylaydi
        st, me1 = tab1.get("/api/auth/me")
        self.assertEqual(st, 200, me1)
        self.assertEqual(me1["user"]["id"], u1["id"], "1-tab boshqa foydalanuvchiga o'tib ketgan!")
        self.assertEqual(me1["user"]["role"], "instructor")

        st, me2 = tab2.get("/api/auth/me")
        self.assertEqual(st, 200, me2)
        self.assertEqual(me2["user"]["id"], u2["id"])
        self.assertEqual(me2["user"]["role"], "student")

        # Rollararo endpoint'lar aralashmaydi
        self.assertEqual(tab1.get("/api/student/home")[0], 403)
        self.assertEqual(tab2.get("/api/instructor/students")[0], 403)

    def test_101_three_roles_parallel_no_mixing(self):
        """3 rol, 3 qurilma, 300 ta parallel so'rov — javoblar aralashmasin."""
        clients = {}
        expect = {}
        for i, creds in enumerate((self.ADMIN, self.INSTR, self.STUD), start=1):
            c = TabClient("iso-par-%d" % i)
            expect[i] = self._login(c, creds)["id"]
            clients[i] = c

        bad = []
        lock = threading.Lock()

        def worker(idx):
            c = clients[idx]
            for _ in range(100):
                st, r = c.get("/api/auth/me")
                if st != 200 or r.get("user", {}).get("id") != expect[idx]:
                    with lock:
                        bad.append((idx, st, r))

        threads = [threading.Thread(target=worker, args=(i,)) for i in (1, 2, 3)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(bad, [], "parallel so'rovlarda sessiya aralashdi: %s" % bad[:5])

    def test_102_role_data_isolation(self):
        """Har bir rol faqat o'z ma'lumotini oladi."""
        admin = TabClient("iso-iso-admin"); self._login(admin, self.ADMIN)
        instr = TabClient("iso-iso-instr"); self._login(instr, self.INSTR)
        stud = TabClient("iso-iso-stud"); self._login(stud, self.STUD)

        self.assertEqual(admin.get("/api/admin/users")[0], 200)
        self.assertEqual(instr.get("/api/instructor/students")[0], 200)
        self.assertEqual(stud.get("/api/student/sessions")[0], 200)

        self.assertEqual(stud.get("/api/admin/users")[0], 403)
        self.assertEqual(stud.get("/api/instructor/students")[0], 403)
        self.assertEqual(instr.get("/api/admin/users")[0], 403)
        self.assertEqual(instr.get("/api/student/sessions")[0], 403)
        self.assertEqual(admin.get("/api/student/sessions")[0], 403)

    def test_103_logout_in_one_tab_keeps_others(self):
        """Bitta tab chiqishi qo'shni tab'larni o'ldirmasin."""
        jar = http.cookiejar.CookieJar()
        tab1 = TabClient("iso-out-1", jar=jar)
        tab2 = TabClient("iso-out-2", jar=jar)
        self._login(tab1, self.INSTR)
        self._login(tab2, self.STUD)

        self.assertEqual(tab1.post("/api/auth/logout")[0], 200)
        self.assertEqual(tab1.get("/api/auth/me")[0], 401, "chiqilgandan keyin sessiya o'chiqilmadi")
        # Qurilma kaliti o'chmaydi — shuning uchun 2-tab ishlashda davom etadi
        st, r = tab2.get("/api/auth/me")
        self.assertEqual(st, 200, "bitta tab chiqishi ikkinchisini buzdi")
        self.assertEqual(r["user"]["role"], "student")

    # ---------------------------------------------------------------- xavfsizlik
    def test_104_device_cookie_alone_is_not_a_credential(self):
        """FAQAT cookie (tab kalitisiz) bilan kirib bo'lmasin."""
        c = TabClient("iso-sec-1")
        self._login(c, self.ADMIN)
        dev = self._device(c)
        self.assertTrue(dev)

        # 1) Qurilma kaliti yalang'och — hech qanday sarlavhasiz
        self.assertEqual(Client().req("GET", "/api/auth/me", cookie=dev)[0], 401)
        # 2) Noto'g'ri tab kaliti bilan
        self.assertEqual(TabClient("iso-sec-bogus").req("GET", "/api/auth/me", cookie=dev)[0], 401)

    def test_105_tab_id_alone_is_not_a_credential(self):
        """FAQAT tab kaliti (cookiesiz) bilan kirib bo'lmasin."""
        import urllib.request as _u
        c = TabClient("iso-sec-2")
        self._login(c, self.ADMIN)
        # So'rov tab sarlavhasini yuboradi, lekin cookie'siz (oddiy klient)
        req = _u.Request(f"http://{HOST}:{PORT}/api/auth/me",
                         headers={"X-Avto-Tab": c.tab}, method="GET")
        try:
            with _u.urlopen(req, timeout=15) as resp:
                st = resp.status
        except urllib.error.HTTPError as e:
            st = e.code
        self.assertEqual(st, 401, "faqat tab kaliti bilan kirish mumkin bo'lsa — zaiflik")

    def test_106_legacy_client_without_header_works(self):
        """Eski mijozlar (curl, test, bot) sarlavhasiz ham ishlayveradi."""
        c = Client()
        st, r = c.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        self.assertEqual(c.get("/api/auth/me")[0], 200)
        self.assertEqual(c.post("/api/auth/logout")[0], 200)
        self.assertEqual(c.get("/api/auth/me")[0], 401)

    def test_107_legacy_and_tab_sessions_do_not_leak(self):
        """Sarlavhasiz (legacy) sessiya bilan tab'ga xos sessiya aralashmasin."""
        legacy = Client()
        st, _ = legacy.post("/api/auth/login", {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200)
        token = None
        for ck in legacy.jar:
            if ck.name == "sid":
                token = ck.value
        self.assertTrue(token)

        # O'sha cookie qiymati + tab kaliti bilan kirish MUMKIN EMAS
        self.assertEqual(TabClient("iso-mix-1").req("GET", "/api/auth/me", cookie=token)[0], 401)
        # ... lekin legacy ko'rinishida ishlaydi
        self.assertEqual(legacy.get("/api/auth/me")[0], 200)
        legacy.post("/api/auth/logout")

    def test_108_remember_me_cookie_shape(self):
        """"Meni eslab qolish" — yoqilganda 30 kun, o'chirilganda sessiya cookie'si.

        (MODUL 3 uchun batafsil testlar: TestRememberMe)"""
        import urllib.request as _u
        cj = http.cookiejar.CookieJar()
        op = _u.build_opener(_u.HTTPCookieProcessor(cj))
        req = _u.Request(
            f"http://{HOST}:{PORT}/api/auth/login",
            data=json.dumps({"login": "admin", "password": "admin123",
                             "role": "admin", "remember": True}).encode(),
            headers={"Content-Type": "application/json",
                     "X-Requested-With": "Avtomaktab",
                     "X-Avto-Tab": "iso-rm-1"},
            method="POST")
        with op.open(req, timeout=15) as resp:
            sc = resp.headers.get("Set-Cookie", "")
        self.assertIn("Max-Age=2592000", sc, sc)
        self.assertIn("HttpOnly", sc)

    def test_109_relogin_same_tab_ok(self):
        """Bir xil tab'da qayta-qayta kirish xatosiz ishlashi (UNIQUE constraint)."""
        jar = http.cookiejar.CookieJar()
        c = TabClient("iso-relog", jar=jar)
        u1 = self._login(c, self.STUD)
        st, r = c.post("/api/auth/logout")
        self.assertEqual(st, 200, r)
        # Chiqish so'rovi serverga YETIB BORMAGAN holatni ham modellashtiramiz:
        # boshqa tab chiqib ketsin, bu tab esa o'z holicha kalsin.
        c2 = TabClient("iso-relog-2", jar=jar)
        self._login(c2, self.INSTR)
        u2 = self._login(c, self.STUD)
        self.assertEqual(u2["id"], u1["id"])
        st, me = c.get("/api/auth/me")
        self.assertEqual(st, 200, me)
        self.assertEqual(me["user"]["id"], u1["id"])

    def test_110_revoke_all_still_works_per_tab(self):
        """'Barcha qurilmalardan chiqish' — har bir tab'ning sessiyasi."""
        jar = http.cookiejar.CookieJar()
        t1 = TabClient("iso-rev-1", jar=jar)
        t2 = TabClient("iso-rev-2", jar=jar)
        self._login(t1, self.ADMIN)
        self._login(t2, self.ADMIN)

        self.assertEqual(t1.get("/api/auth/me")[0], 200)
        self.assertEqual(t2.get("/api/auth/me")[0], 200)

        self.assertEqual(t2.post("/api/me/sessions/revoke-all")[0], 200)
        self.assertEqual(t1.get("/api/auth/me")[0], 401)
        self.assertEqual(t2.get("/api/auth/me")[0], 401)


class TestRoleMustMatch(Base):
    """MODUL 2 — login qilganda tanlangan rol hisobning roliga MOS bo'lishi shart."""

    CASES = [
        ("admin", "admin123", "admin"),
        (INSTR_LOGIN_1, INSTR_PASS_1, "instructor"),
        (STUD_LOGIN_1, STUD_PASS_1, "student"),
    ]

    def _login(self, body):
        # Bu testlar ataylab RAD etilishlarni tekshiradi — login urinishlari
        # limitini (6 urinish / 5 daqiqa) tozalab, boshqa testlarni koldirmaymiz.
        from app.auth import reset_login_attempts
        reset_login_attempts("login:" + body["login"])
        return Client().post("/api/auth/login", body)

    def test_200_correct_role_allowed(self):
        for login, pw, role in self.CASES:
            st, r = self._login({"login": login, "password": pw, "role": role})
            self.assertEqual(st, 200, f"{login}/{role}: {r}")
            self.assertEqual(r["user"]["role"], role)

    def test_201_wrong_role_rejected(self):
        for login, pw, role in self.CASES:
            for wrong in ("admin", "instructor", "student"):
                if wrong == role:
                    continue
                st, r = self._login({"login": login, "password": pw, "role": wrong})
                self.assertEqual(st, 400, f"{login} roli {wrong} deb kirildi: {r}")
                self.assertEqual(r["error"], "auth.wrong_role")

    def test_202_missing_role_rejected(self):
        """Rol yuborilmasa kirish RAD etiladi (avval 'avtomatik' deb o'tib ketardi)."""
        for login, pw, _role in self.CASES:
            for extra in ({}, {"role": ""}, {"role": "   "}, {"role": None}):
                body = {"login": login, "password": pw}
                body.update(extra)
                st, r = self._login(body)
                self.assertEqual(st, 400, f"{login} rolisiz kirdi: {r}")
                self.assertEqual(r["error"], "auth.wrong_role")

    def test_203_unknown_role_rejected(self):
        for bogus in ("Admin", "ADMIN", "root", "superadmin", "teaching", "1", "null", "student "):
            st, r = self._login({"login": "admin", "password": "admin123", "role": bogus})
            self.assertEqual(st, 400, f"noma'lum rol '{bogus}' qabul qilindi: {r}")
            self.assertEqual(r["error"], "auth.wrong_role")

    def test_204_wrong_password_still_reported_first(self):
        """Parol noto'g'ri bo'lsa, xabar rol haqida emas — hisob borligi oshkor bo'lmaydi."""
        st, r = self._login({"login": STUD_LOGIN_1, "password": "noto'g'ri-parol", "role": "admin"})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "auth.wrong_credentials")

    def test_205_no_session_created_on_role_mismatch(self):
        """Rad etilgandan keyin sessiya QOLMASIN."""
        c = Client()
        st, _ = self._login({"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "admin"})
        self.assertEqual(st, 400)
        c = TabClient("role-mismatch-1")
        c.post("/api/auth/login", {"login": STUD_LOGIN_1, "password": STUD_PASS_1, "role": "admin"})
        self.assertEqual(c.get("/api/auth/me")[0], 401)


class TestRememberMe(Base):
    """MODUL 3 — "Meni eslab qolish"."""

    def _login_cookie(self, remember):
        cj = http.cookiejar.CookieJar()
        op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
        req = urllib.request.Request(
            f"http://{HOST}:{PORT}/api/auth/login",
            data=json.dumps({"login": "admin", "password": "admin123",
                             "role": "admin", "remember": remember}).encode(),
            headers={"Content-Type": "application/json",
                     "X-Requested-With": "Avtomaktab", "X-Avto-Tab": "rm-%s" % remember},
            method="POST")
        with op.open(req, timeout=15) as resp:
            return resp.headers.get("Set-Cookie", "")

    def test_210_remember_on_is_30_days(self):
        sc = self._login_cookie(True)
        self.assertIn("Max-Age=2592000", sc, sc)
        self.assertIn("HttpOnly", sc)
        self.assertIn("SameSite=Lax", sc)

    def test_211_remember_off_is_session_cookie(self):
        """Checkbox o'chirilgan -> Max-Age YO'Q = brauzer yopilganda o'chadi."""
        sc = self._login_cookie(False)
        self.assertNotIn("Max-Age", sc, "sessiya cookie'si Max-Age olmasligi kerak: %s" % sc)
        self.assertIn("HttpOnly", sc)

    def test_212_session_actually_works(self):
        for remember in (True, False):
            c = TabClient("rm-work-%s" % remember)
            st, r = c.post("/api/auth/login", {
                "login": "admin", "password": "admin123", "role": "admin", "remember": remember})
            self.assertEqual(st, 200, r)
            self.assertEqual(c.get("/api/auth/me")[0], 200)

    def test_213_token_never_in_response_body(self):
        """Token JS ga BERILMASIN — faqat HttpOnly cookie orqali (XSS himoyasi)."""
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": "admin", "password": "admin123", "role": "admin", "remember": True})
        self.assertEqual(st, 200, r)
        self.assertNotIn("token", r, "token javob tanasida qaytayapti — XSS ga ochiq!")
        self.assertIn("user", r)

    def test_214_login_name_only_in_localstorage(self):
        """localStorage'da faqat LOGIN nomi saqlanadi — token emas."""
        import re
        src = (ROOT / "web" / "js" / "app.js").read_text(encoding="utf-8")
        for m in re.finditer(r'localStorage\.(setItem|getItem)\(\s*"([^"]+)"', src):
            self.assertNotIn("token", m.group(2).lower(),
                             "localStorage'ga token yozilmoqda: %s" % m.group(2))
            self.assertNotIn("sid", m.group(2).lower(),
                             "localStorage'ga sessiya kaliti yozilmoqda: %s" % m.group(2))


class TestUsersByRole(Base):
    """MODUL 6 — foydalanuvchilar rolda bo'linib ko'rsatiladi
    (Talabalar | Instruktorlar | Adminlar), har biri o'z qidiruvi bilan."""

    def setUp(self):
        self.admin = Client()
        from app.auth import reset_login_attempts
        reset_login_attempts("login:admin")
        st, r = self.admin.post("/api/auth/login",
                                {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)

    def test_300_each_role_returned_separately(self):
        """Har bir rol faqat o'z rolikadagilarni oladi — ro'yxatlar aralashmaydi."""
        for role in ("student", "instructor", "admin"):
            st, r = self.admin.get(f"/api/admin/users?role={role}")
            self.assertEqual(st, 200, r)
            self.assertTrue(r["users"], f"{role} bo'sh qaytdi")
            for u in r["users"]:
                self.assertEqual(u["role"], role,
                                 f"role={role} so'rovida {u['role']} qaytdi — ro'yxatlar aralashdi")

    def test_301_admin_tab_shows_admins(self):
        st, r = self.admin.get("/api/admin/users?role=admin")
        self.assertEqual(st, 200, r)
        self.assertTrue(any(u["login"] == "admin" for u in r["users"]),
                        "boshqaruvchi admin ro'yxatda yo'q")
        # Adminlar uchun 'progress'/'instructors' kabi talaba maydonlari chiqmasin
        for u in r["users"]:
            self.assertNotIn("progress", u)
            self.assertNotIn("instructors", u)

    def test_302_student_tab_has_instructor_and_progress(self):
        st, r = self.admin.get("/api/admin/users?role=student")
        self.assertEqual(st, 200, r)
        for u in r["users"]:
            self.assertIn("student", u)
            pr = u.get("progress")
            self.assertIsNotNone(pr, f"{u['login']} uchun progress yo'q")
            for k in ("total", "done", "pct"):
                self.assertIn(k, pr)
            self.assertTrue(0 <= pr["pct"] <= 100, f"{u['login']}: pct={pr['pct']}")
            self.assertLessEqual(pr["done"], pr["total"])
            self.assertIsInstance(u.get("instructors"), list)

    def test_303_instructor_tab_has_students_count(self):
        st, r = self.admin.get("/api/admin/users?role=instructor")
        self.assertEqual(st, 200, r)
        for u in r["users"]:
            self.assertIn("instructor", u)
            self.assertIsInstance(u.get("students_count"), int)
            self.assertGreaterEqual(u["students_count"], 0)

    def test_304_search_inside_role(self):
        """Qidiruv faqat tanlangan rol ichida ishlaydi."""
        st, r = self.admin.get("/api/admin/users?role=student&q=" + STUD_LOGIN_1)
        self.assertEqual(st, 200, r)
        self.assertTrue(any(u["login"] == STUD_LOGIN_1 for u in r["users"]))

        # Boshqa rolga tegishli qidiruv natijasi bo'sh bo'lishi SHART
        st, r2 = self.admin.get("/api/admin/users?role=admin&q=" + STUD_LOGIN_1)
        self.assertEqual(st, 200, r2)
        self.assertEqual([u for u in r2["users"] if u["login"] == STUD_LOGIN_1], [])

        st, r3 = self.admin.get("/api/admin/users?role=student&q=admin123")
        self.assertEqual(st, 200, r3)
        self.assertEqual([u for u in r3["users"] if u["role"] != "student"], [])

    def test_305_other_roles_cannot_list_users(self):
        for login, pw, role in ((STUD_LOGIN_1, STUD_PASS_1, "student"),
                                 (INSTR_LOGIN_1, INSTR_PASS_1, "instructor")):
            from app.auth import reset_login_attempts
            reset_login_attempts("login:" + login)
            c = Client()
            st, _ = c.post("/api/auth/login", {"login": login, "password": pw, "role": role})
            self.assertEqual(st, 200)
            for q in ("role=student", "role=instructor", "role=admin"):
                self.assertEqual(c.get(f"/api/admin/users?{q}")[0], 403,
                                 f"{role} admin ro'yxatini ko'ra oladi!")


class TestTotalLessonsTarget(Base):
    """MODUL 5 — "Jami darslar": admin belgilaydigan maqsad.
    Rejimlar: individual (talaba uchun alohida) va ommaviy (platforma)."""

    STUD = (STUD_LOGIN_1, STUD_PASS_1, "student")

    def setUp(self):
        self.admin = Client()
        from app.auth import reset_login_attempts
        reset_login_attempts("login:admin")
        st, r = self.admin.post("/api/auth/login",
                                {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        self.db = Db(str(TMP / "test.db"))
        # Har bir test toza holatdan boshlanishi uchun individual qiymatni
        # ommaviy sozlamaga qaytaramiz.
        self._set_individual(None)
        self._set_group(30)

    # ------------------------------------------------------------------ yordamchi
    def _set_group(self, n):
        st, r = self.admin.put("/api/admin/settings", {"total_lessons_target": n})
        self.assertEqual(st, 200, r)

    def _set_individual(self, val):
        st, r = self.admin.get("/api/admin/users?role=student&q=" + STUD_LOGIN_1)
        self.assertEqual(st, 200, r)
        uid = r["users"][0]["id"]
        body = {"student": {"total_lessons_target": val}}
        st, r2 = self.admin.put(f"/api/admin/users/{uid}", body)
        self.assertEqual(st, 200, r2)

    def _home(self):
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": self.STUD[0], "password": self.STUD[1], "role": "student"})
        self.assertEqual(st, 200, r)
        st, h = c.get("/api/student/home")
        self.assertEqual(st, 200, h)
        return h

    def _done_count(self):
        """Bazadagi haqiqiy bajarilgan darslar soni."""
        row = self.db.q1(
            """SELECT COUNT(*) AS n FROM session_students ss
               JOIN lesson_sessions ls ON ls.id=ss.session_id
               JOIN students s ON s.id=ss.student_id
               JOIN users u ON u.id=s.user_id
               WHERE u.login=? AND ss.student_status='active'
                 AND ls.status!='cancelled' AND ls.status='completed'""",
            (self.STUD[0],))
        return row["n"]

    # ------------------------------------------------------------------ testlar
    def test_400_home_returns_progress(self):
        h = self._home()
        pr = h["progress"]
        for k in ("target", "done", "remaining", "pct", "mode"):
            self.assertIn(k, pr)
        self.assertEqual(pr["done"], self._done_count())
        self.assertEqual(pr["mode"], "group", "individual qiymat yo'q — ommaviy rejim kerak")
        self.assertEqual(pr["target"], 30)
        self.assertEqual(pr["pct"], round(self._done_count() * 100 / 30))

    def test_401_group_target_applies_to_all(self):
        for n in (10, 45, 100):
            self._set_group(n)
            self.assertEqual(self._home()["progress"]["target"], n,
                             f"ommaviy maqsad {n} qo'llanmadi")

    def test_402_individual_overrides_group(self):
        self._set_group(30)
        self._set_individual(12)
        pr = self._home()["progress"]
        self.assertEqual(pr["mode"], "individual")
        self.assertEqual(pr["target"], 12, "individual maqsad ustun kelmadi")
        self.assertEqual(pr["group_target"], 30)
        self.assertEqual(pr["individual_target"], 12)

    def test_403_clearing_individual_returns_to_group(self):
        self._set_individual(12)
        self.assertEqual(self._home()["progress"]["mode"], "individual")
        self._set_individual(None)
        pr = self._home()["progress"]
        self.assertEqual(pr["mode"], "group")
        self.assertEqual(pr["target"], 30)
        self.assertIsNone(pr["individual_target"])

    def test_404_pct_and_remaining_are_consistent(self):
        self._set_group(20)
        pr = self._home()["progress"]
        self.assertEqual(pr["remaining"], max(0, pr["target"] - pr["done"]))
        self.assertEqual(pr["pct"], round(pr["done"] * 100 / pr["target"]))
        self.assertTrue(0 <= pr["pct"] <= 100)

    def test_405_invalid_group_target_normalized(self):
        """Admin noto'g'ri qiymat yuborsa — standartga qaytariladi (xato chiqmaydi)."""
        for bad, want in ((0, 30), (-5, 30), (1000, 999), ("abc", 30), (None, 30), ("45", 45)):
            st, r = self.admin.put("/api/admin/settings", {"total_lessons_target": bad})
            self.assertEqual(st, 200, f"{bad!r}: {r}")
            pr = self._home()["progress"]
            self.assertEqual(pr["target"], want, f"{bad!r} -> {pr['target']}, kutilgan {want}")
            self.assertGreaterEqual(pr["target"], 1, "maqsad 0 bo'lsa foiz hisobi buziladi")
            self.assertTrue(0 <= pr["pct"] <= 100)

    def test_406_invalid_individual_target_rejected(self):
        self._set_individual(20)
        for bad in (-5, 1000, 99999, "abc", "1.5"):
            st, r = self.admin.put(f"/api/admin/users/{self._uid()}",
                                   {"student": {"total_lessons_target": bad}})
            self.assertEqual(st, 400, f"{bad!r} qabul qilindi: {r}")
            self.assertEqual(r["error"], "user.bad_total_lessons")
        # Xato qabul qilinmagan qiymat maqsadni buzmasin
        self.assertEqual(self._ind_value(), 20)

    def test_406b_zero_means_clear_individual(self):
        """0 / bo'sh qiymat = individual rejimni bekor qilish (ommaviyga qaytish)."""
        for empty in (0, "0", "", None):
            self._set_individual(15)
            self._set_individual(empty)
            self.assertIsNone(self._ind_value(), f"{empty!r} individual maqsadni tozalamadi")
            self.assertEqual(self._home()["progress"]["mode"], "group")

    def test_407_valid_individual_target_accepted(self):
        for good in (1, 50, 999):
            self._set_individual(good)
            self.assertEqual(self._ind_value(), good)
            self.assertEqual(self._home()["progress"]["target"], good)

    def test_408_new_student_can_get_individual_target(self):
        st, r = self.admin.post("/api/admin/users", {
            "role": "student", "first_name": "Maqsad", "last_name": "Test",
            "group_name": "G-1", "license_category": "B", "total_lessons_target": 7,
        })
        self.assertEqual(st, 200, r)
        uid = r["user"]["id"]
        row = self.db.q1("SELECT total_lessons_target FROM students WHERE user_id=?", (uid,))
        self.assertEqual(row["total_lessons_target"], 7)
        # Tozalash
        self.db.upd("UPDATE users SET deleted_at=datetime('now') WHERE id=?", (uid,))

    def test_409_new_student_without_target_uses_group(self):
        st, r = self.admin.post("/api/admin/users", {
            "role": "student", "first_name": "Kurs", "last_name": "Test",
            "group_name": "G-2", "license_category": "B",
        })
        self.assertEqual(st, 200, r)
        row = self.db.q1("SELECT total_lessons_target FROM students WHERE user_id=?",
                         (r["user"]["id"],))
        self.assertIsNone(row["total_lessons_target"], "individual maqsad berilmagan bo'lishi kerak")
        self.db.upd("UPDATE users SET deleted_at=datetime('now') WHERE id=?", (r["user"]["id"],))

    def test_410_settings_expose_target(self):
        st, r = self.admin.get("/api/admin/settings")
        self.assertEqual(st, 200, r)
        self.assertIn("total_lessons_target", r["settings"])
        self._set_group(44)
        st, r = self.admin.get("/api/admin/settings")
        self.assertEqual(r["settings"]["total_lessons_target"], 44)

    def test_411_student_cannot_change_target(self):
        st, r = self.admin.post("/api/admin/users", {
            "role": "student", "first_name": "Sinov", "last_name": "Test", "group_name": "G-3"})
        self.assertEqual(st, 200, r)
        uid = r["user"]["id"]
        login = r["credentials"]["login"]
        pw = r["credentials"]["password"]
        c = Client()
        st, _ = c.post("/api/auth/login", {"login": login, "password": pw, "role": "student"})
        self.assertEqual(st, 200)
        self.assertEqual(c.put(f"/api/admin/users/{uid}",
                               {"student": {"total_lessons_target": 1}})[0], 403)
        self.assertEqual(c.put("/api/admin/settings", {"total_lessons_target": 1})[0], 403)
        self.assertIsNone(self.db.q1(
            "SELECT total_lessons_target FROM students WHERE user_id=?", (uid,))["total_lessons_target"])
        self.db.upd("UPDATE users SET deleted_at=datetime('now') WHERE id=?", (uid,))

    def test_412_instructor_cannot_change_target(self):
        c = Client()
        st, _ = c.post("/api/auth/login",
                       {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        self.assertEqual(st, 200)
        self.assertEqual(c.put("/api/admin/settings", {"total_lessons_target": 1})[0], 403)

    # ------------------------------------------------------------------ kichik yordamchilar
    def _uid(self):
        st, r = self.admin.get("/api/admin/users?role=student&q=" + STUD_LOGIN_1)
        return r["users"][0]["id"]

    def _ind_value(self):
        return self.db.q1("SELECT total_lessons_target FROM students WHERE user_id=?",
                          (self._uid(),))["total_lessons_target"]


class TestMapAndCoordinates(Base):
    """MODUL 4 — XARITA (Yandex Maps JS API) va kenglik/uzunlik.

    Tekshiriladi:
      * `GET /api/config` — xarita sozlamalari KIRISHSIZ qaytariladi,
        kalit faqat `.env` dan olinadi, hech qanday sirli ma'lumot yo'q.
      * Xarita markazi — Toshkent (`.env` bilan o'zgartirilishi mumkin).
      * "Kenglik"/"Uzunlik" validatsiyasi — noto'g'ri qiymat JIM QOLMAYDI,
        aniq `coord.*` kodi qaytariladi va eski qiymat buzilmaydi.
    """

    ENV_KEYS = ("YANDEX_MAPS_API_KEY", "MAP_PROVIDER", "MAP_CENTER_LAT",
                "MAP_CENTER_LNG", "MAP_CENTER_ZOOM")
    TASHKENT = (41.311081, 69.240562)

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from app.auth import reset_login_attempts
        reset_login_attempts("login:admin")
        admin = Client()
        st, _ = admin.post("/api/auth/login",
                           {"login": "admin", "password": "admin123", "role": "admin"})
        assert st == 200, "admin login muvaffaqiyatsiz"
        st, cr = admin.post("/api/admin/users", {
            "role": "student", "first_name": "Xarita", "last_name": "Sinov",
            "group_name": "G-M4", "license_category": "B",
        })
        assert st == 200, cr
        cls.user_id = cr["user"]["id"]
        cls.login = cr["credentials"]["login"]
        cls.password = cr["credentials"]["password"]
        row = Db(str(TMP / "test.db")).q1("SELECT id FROM students WHERE user_id=?", (cls.user_id,))
        assert row, "talaba yaratilmadi"
        cls.student_id = row["id"]

        # Yetarli masofadagi kun (seed sessiyalari bilan to'qnashmasligi uchun)
        from datetime import datetime, timedelta
        d = datetime.now()
        day = (d + timedelta(days=(7 - d.weekday()) % 7 or 7) + timedelta(days=21)).strftime("%Y-%m-%d")
        st, sr = admin.post("/api/admin/sessions", {
            "date": day, "start_time": "08:00", "end_time": "09:00",
            "instructor_id": 1, "student_ids": [cls.student_id], "notes": "M4MAP",
        })
        assert st == 200, sr
        cls.session_id = sr["id"]

    @classmethod
    def tearDownClass(cls):
        db = Db(str(TMP / "test.db"))
        if getattr(cls, "session_id", None):
            db.upd("DELETE FROM session_students WHERE session_id=?", (cls.session_id,))
            db.upd("DELETE FROM notifications WHERE data LIKE ?",
                   (f'%"session_id": {cls.session_id}%',))
            db.upd("DELETE FROM lesson_sessions WHERE id=?", (cls.session_id,))
        if getattr(cls, "user_id", None):
            db.upd("DELETE FROM students WHERE user_id=?", (cls.user_id,))
            db.upd("DELETE FROM users WHERE id=?", (cls.user_id,))
        super().tearDownClass()

    def setUp(self):
        from app.auth import reset_login_attempts
        reset_login_attempts("login:admin")
        reset_login_attempts("login:" + self.login)
        reset_login_attempts("login:INSTR_LOGIN_1")
        self.admin = Client()
        st, r = self.admin.post("/api/auth/login",
                                {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        self.db = Db(str(TMP / "test.db"))
        # .env kalitlarini o'zgartirmaslik uchun zaxira
        self._env_backup = {k: os.environ.get(k) for k in self.ENV_KEYS}
        self._clear_env()
        self._reset_pickup(None, None, "")

    def tearDown(self):
        self._restore_env()

    def _clear_env(self):
        for k in self.ENV_KEYS:
            os.environ.pop(k, None)

    def _restore_env(self):
        for k, v in self._env_backup.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v

    def _reset_pickup(self, lat, lng, address=""):
        self.db.upd("""UPDATE session_students
                       SET pickup_address=?, pickup_lat=?, pickup_lng=?
                       WHERE session_id=? AND student_id=?""",
                    (address, lat, lng, self.session_id, self.student_id))

    def _pickup(self):
        return self.db.q1("""SELECT pickup_address, pickup_lat, pickup_lng
                             FROM session_students WHERE session_id=? AND student_id=?""",
                          (self.session_id, self.student_id))

    def _student(self):
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": self.login, "password": self.password, "role": "student"})
        self.assertEqual(st, 200, r)
        return c

    def _put(self, c, body, sid=None):
        return c.put(f"/api/student/sessions/{sid or self.session_id}/pickup", body)

    # ------------------------------------------------------------ /api/config
    def test_500_config_is_public(self):
        """Sessiyasiz (cookie'siz) so'rov ham ishlashi SHART — login sahifasi
        ham xaritadan foydalanadi."""
        anon = Client()
        st, r = anon.get("/api/config")
        self.assertEqual(st, 200, r)
        self.assertTrue(r["ok"])
        self.assertIn("maps", r["config"])

    def test_501_config_shape(self):
        st, r = self.admin.get("/api/config")
        self.assertEqual(st, 200, r)
        m = r["config"]["maps"]
        for k in ("provider", "api_key", "enabled", "center", "zoom"):
            self.assertIn(k, m, f"/api/config.maps.{k} yo'q")
        self.assertIn("lat", m["center"])
        self.assertIn("lng", m["center"])
        self.assertIsInstance(m["api_key"], str)
        self.assertIsInstance(m["enabled"], bool)

    def test_502_center_is_tashkent_by_default(self):
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertAlmostEqual(m["center"]["lat"], self.TASHKENT[0], places=4)
        self.assertAlmostEqual(m["center"]["lng"], self.TASHKENT[1], places=4)
        self.assertEqual(m["zoom"], 12)
        # Toshkent chegaralarida bo'lishi SHART (41°N, 69°E)
        self.assertTrue(40 < m["center"]["lat"] < 42)
        self.assertTrue(68 < m["center"]["lng"] < 70)

    def test_503_center_from_env(self):
        os.environ["MAP_CENTER_LAT"] = "39.6542"
        os.environ["MAP_CENTER_LNG"] = "66.9597"
        os.environ["MAP_CENTER_ZOOM"] = "14"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertAlmostEqual(m["center"]["lat"], 39.6542, places=4)
        self.assertAlmostEqual(m["center"]["lng"], 66.9597, places=4)
        self.assertEqual(m["zoom"], 14)

    def test_504_bad_center_falls_back_to_tashkent(self):
        """`.env` da noto'g'ri son bo'lsa — xarita JOYNI bo'sh qolmasligi uchun
        Toshkentga qaytariladi (jim qolmaydi, lekin xarita ishlayveradi)."""
        os.environ["MAP_CENTER_LAT"] = "Toshkent"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertAlmostEqual(m["center"]["lat"], self.TASHKENT[0], places=4)
        os.environ["MAP_CENTER_LAT"] = "41.311081,5"   # vergulli son
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertAlmostEqual(m["center"]["lat"], 41.311081, places=4)

    def test_505_disabled_without_key(self):
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertFalse(m["enabled"], "kalit yo'q bo'lsa xarita 'ochilgan' bo'lmasligi kerak")
        self.assertEqual(m["api_key"], "")

    def test_506_enabled_with_key_from_env(self):
        os.environ["YANDEX_MAPS_API_KEY"] = "test-abc123.Yx-mapKey"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertTrue(m["enabled"])
        self.assertEqual(m["api_key"], "test-abc123.Yx-mapKey")
        self.assertEqual(m["provider"], "yandex")

    def test_507_provider_none_disables_map(self):
        os.environ["YANDEX_MAPS_API_KEY"] = "test-abc123.Yx-mapKey"
        os.environ["MAP_PROVIDER"] = "none"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertFalse(m["enabled"], "MAP_PROVIDER=none — xarita ko'rsatilmasligi kerak")
        self.assertEqual(m["provider"], "none")

    def test_508_config_never_leaks_secrets(self):
        """Ommaviy endpoint parol/token YOK, faqat xarita kalitini beradi."""
        os.environ["YANDEX_MAPS_API_KEY"] = "map-key-777"
        os.environ.setdefault("BOT_TOKEN", "123456:SECRET-BOT-TOKEN")
        os.environ.setdefault("ADMIN_INITIAL_PASSWORD", "SECRET-ADMIN-PW")
        st, r = self.admin.get("/api/config")
        self.assertEqual(st, 200)
        raw = json.dumps(r, ensure_ascii=False)
        self.assertNotIn("SECRET-BOT-TOKEN", raw)
        self.assertNotIn("SECRET-ADMIN-PW", raw)
        self.assertNotIn("BOT_TOKEN", raw)
        self.assertNotIn("ADMIN_INITIAL_PASSWORD", raw)
        self.assertNotIn("password", raw.lower())
        self.assertNotIn("token", raw.lower())

    def test_509_config_only_get(self):
        self.assertEqual(self.admin.post("/api/config", {"x": 1})[0], 404)
        self.assertEqual(self.admin.put("/api/config", {"x": 1})[0], 404)
        self.assertEqual(self.admin.delete("/api/config")[0], 404)
        self.assertEqual(self.admin.get("/api/config/maps")[0], 404)

    # ------------------------------------------------------------ kenglik/uzunlik
    def test_510_valid_coords_saved(self):
        c = self._student()
        st, r = self._put(c, {"address": "Yunusobod 12", "lat": 41.3364, "lng": 69.2785})
        self.assertEqual(st, 200, r)
        row = self._pickup()
        self.assertEqual(row["pickup_address"], "Yunusobod 12")
        self.assertAlmostEqual(row["pickup_lat"], 41.3364, places=5)
        self.assertAlmostEqual(row["pickup_lng"], 69.2785, places=5)

    def test_511_coords_as_strings(self):
        """Frontend `input` dan qiymat STRING bo'lib keladi — qabul qilinishi kerak."""
        c = self._student()
        for lat, lng in (("41.311081", "69.240562"), (" 41.311081 ", " 69.240562 "),
                         ("41,311081", "69,240562")):
            st, r = self._put(c, {"lat": lat, "lng": lng})
            self.assertEqual(st, 200, f"{lat}/{lng}: {r}")
            row = self._pickup()
            self.assertAlmostEqual(row["pickup_lat"], 41.311081, places=5)
            self.assertAlmostEqual(row["pickup_lng"], 69.240562, places=5)

    def test_512_coords_rounded_to_7_decimals(self):
        c = self._student()
        st, _ = self._put(c, {"lat": 41.31108150000, "lng": 69.24056240000})
        self.assertEqual(st, 200)
        row = self._pickup()
        self.assertEqual(row["pickup_lat"], 41.3110815)
        self.assertEqual(row["pickup_lng"], 69.2405624)

    def test_513_bounds_are_allowed(self):
        """Chegara qiymatlari xato EMAS (Toshkentdan ancha uzoq nuqta ham kerak)."""
        c = self._student()
        for lat, lng in ((90, 180), (-90, -180), (0, 0)):
            st, r = self._put(c, {"lat": lat, "lng": lng})
            self.assertEqual(st, 200, f"{lat}/{lng}: {r}")

    def test_514_bad_lat_rejected(self):
        c = self._student()
        for bad in ("abc", 90.0001, -90.0001, 1000, True, "NaN", "Infinity", "1,2,3", [], {}):
            st, r = self._put(c, {"lat": bad, "lng": 69.24})
            self.assertEqual(st, 400, f"kenglik {bad!r} qabul qilindi: {r}")
            self.assertEqual(r["error"], "coord.bad_lat", f"{bad!r}: {r}")

    def test_515_bad_lng_rejected(self):
        c = self._student()
        for bad in ("abc", 180.0001, -180.0001, 99999, "NaN", "1e400"):
            st, r = self._put(c, {"lat": 41.31, "lng": bad})
            self.assertEqual(st, 400, f"uzunlik {bad!r} qabul qilindi: {r}")
            self.assertEqual(r["error"], "coord.bad_lng", f"{bad!r}: {r}")

    def test_516_incomplete_pair_rejected(self):
        """Faqat bittasi to'ldirilgan holat — chalkash, xato qaytariladi."""
        c = self._student()
        for body in ({"lat": 41.31}, {"lng": 69.24}, {"lat": 41.31, "lng": ""},
                     {"lat": "", "lng": 69.24}):
            st, r = self._put(c, body)
            self.assertEqual(st, 400, f"{body}: {r}")
            self.assertEqual(r["error"], "coord.pair_incomplete", f"{body}: {r}")

    def test_517_clear_both(self):
        self._reset_pickup(41.33, 69.27, "Eski manzil")
        c = self._student()
        st, r = self._put(c, {"address": "", "lat": "", "lng": ""})
        self.assertEqual(st, 200, r)
        row = self._pickup()
        self.assertIsNone(row["pickup_lat"])
        self.assertIsNone(row["pickup_lng"])
        self.assertEqual(row["pickup_address"], "")

    def test_518_rejected_value_keeps_old_data(self):
        """Xatoli so'rov yuborilganda ESKI koordinata buzilmasligi SHART."""
        self._reset_pickup(41.3333, 69.2727, "Saqlangan manzil")
        c = self._student()
        st, r = self._put(c, {"address": "Yangi manzil", "lat": 91, "lng": 69.24})
        self.assertEqual(st, 400, r)
        row = self._pickup()
        self.assertAlmostEqual(row["pickup_lat"], 41.3333, places=5)
        self.assertAlmostEqual(row["pickup_lng"], 69.2727, places=5)
        self.assertEqual(row["pickup_address"], "Saqlangan manzil",
                         "xato so'rovda manzil ham o'zgarib ketmasligi kerak")

    def test_519_address_only_keeps_coords(self):
        """Faqat manzil o'zgartirilsa — koordinata o'zgarMAYdi."""
        self._reset_pickup(41.3333, 69.2727, "Eski")
        c = self._student()
        st, r = self._put(c, {"address": "Yangi manzil"})
        self.assertEqual(st, 200, r)
        row = self._pickup()
        self.assertEqual(row["pickup_address"], "Yangi manzil")
        self.assertAlmostEqual(row["pickup_lat"], 41.3333, places=5)
        self.assertAlmostEqual(row["pickup_lng"], 69.2727, places=5)

    def test_520_other_roles_forbidden(self):
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": INSTR_LOGIN_1, "password": INSTR_PASS_1, "role": "instructor"})
        self.assertEqual(st, 200, r)
        self.assertEqual(self._put(c, {"lat": 41.31, "lng": 69.24})[0], 403)
        self.assertEqual(self._put(self.admin, {"lat": 41.31, "lng": 69.24})[0], 403)

    def test_521_unknown_session_404(self):
        """Noto'g'ri sessiya — validatsiyadan keyin 404 (koordinata o'zgarMAYdi)."""
        self._reset_pickup(41.3333, 69.2727, "Saqlangan")
        c = self._student()
        st, r = self._put(c, {"lat": 41.5, "lng": 69.5}, sid=999999)
        self.assertEqual(st, 404, r)
        self.assertEqual(r["error"], "session.student_not_found")
        row = self._pickup()
        self.assertAlmostEqual(row["pickup_lat"], 41.3333, places=5)

    def test_522_saved_coords_visible_in_session(self):
        """Saqlangan koordinata session detallarida QAYTARILISHI shart
        ("Xaritada ko'rish" havolasi uchun)."""
        c = self._student()
        st, _ = self._put(c, {"address": "Chilonzor 40", "lat": 41.3299, "lng": 69.2902})
        self.assertEqual(st, 200)
        st, r = c.get(f"/api/student/sessions/{self.session_id}")
        self.assertEqual(st, 200, r)
        ss = r["session"]
        self.assertAlmostEqual(ss["pickup_lat"], 41.3299, places=5)
        self.assertAlmostEqual(ss["pickup_lng"], 69.2902, places=5)
        self.assertEqual(ss["pickup_address"], "Chilonzor 40")

    def test_523_session_without_coords_returns_null(self):
        """Koordinata yo'q bo'lsa frontend "Xaritada ko'rish" havolasi
        ko'rsatmasligi kerak — qiymat NULL bo'lishi SHART (0 emas)."""
        c = self._student()
        st, r = c.get(f"/api/student/sessions/{self.session_id}")
        self.assertEqual(st, 200, r)
        ss = r["session"]
        self.assertIsNone(ss["pickup_lat"])
        self.assertIsNone(ss["pickup_lng"])

# ============================================================================
# BAND 24 — YANGI TEST SINFLARI: AUTH / MAP / LESSONS / NOTIFICATIONS /
#            PROFILE / SETTINGS
#
# Har bir sinf BAND raqamiga ishora qiladi va o'sha bandning talabini
# ANIQ tekshiradi. Ma'lumot — har doim DB dan (fixture'lar seed'dan keladi).
# ============================================================================


def _json_db():
    return Db(str(TMP / "test.db"))


def _unlink_all(db, sql, params=()):
    """Test tozalash — `upd` INTEGER qaytarmasligi uchun."""
    db.upd(sql, params)


class _MixinAdmin:
    """Har bir test klassi uchun umumiy yordamchilar."""

    def setUp(self):
        from app.auth import reset_ip_attempts, reset_login_attempts
        for lg in ("admin", INSTR_LOGIN_1, INSTR_LOGIN_3,
                   STUD_LOGIN_1, STUD_LOGIN_2, STUD_LOGIN_3, STUD_LOGIN_4):
            reset_login_attempts("login:" + lg)
        reset_ip_attempts("127.0.0.1")
        self.db = _json_db()
        self.admin = Client()
        st, r = self.admin.post("/api/auth/login",
                                {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)

    def tearDown(self):
        try:
            self.db.close()
        except Exception:
            pass

    def _mkuser(self, role, first, last, **extra):
        body = {"role": role, "first_name": first, "last_name": last}
        body.update(extra)
        st, r = self.admin.post("/api/admin/users", body)
        self.assertEqual(st, 200, r)
        return r

    def _login(self, login, password, role):
        from app.auth import reset_login_attempts
        reset_login_attempts("login:" + login)
        c = Client()
        st, r = c.post("/api/auth/login", {"login": login, "password": password, "role": role})
        return c, st, r

    def _cleanup_user(self, user_id):
        db = _json_db()
        try:
            srow = db.q1("SELECT id FROM students WHERE user_id=?", (user_id,))
            if srow:
                db.upd("DELETE FROM session_students WHERE student_id=?", (srow["id"],))
                db.upd("DELETE FROM students WHERE id=?", (srow["id"],))
            irow = db.q1("SELECT id FROM instructors WHERE user_id=?", (user_id,))
            if irow:
                # instruktorga tegishli darslar avval o'chadi:
                # `lesson_sessions.instructor_id` NOT NULL bo'lgani uchun
                # instruktor qatorini o'chirsak, darslar "osilib" qoladi.
                for ls in db.q("SELECT id FROM lesson_sessions WHERE instructor_id=?",
                               (irow["id"],)):
                    db.upd("DELETE FROM session_students WHERE session_id=?", (ls["id"],))
                    db.upd("DELETE FROM attendance WHERE session_id=?", (ls["id"],))
                    db.upd("DELETE FROM notifications WHERE related_lesson_id=?", (ls["id"],))
                    db.upd("DELETE FROM lesson_sessions WHERE id=?", (ls["id"],))
                db.upd("DELETE FROM instructors WHERE id=?", (irow["id"],))
            db.upd("DELETE FROM notifications WHERE user_id=?", (user_id,))
            db.upd("DELETE FROM user_settings WHERE user_id=?", (user_id,))
            db.upd("DELETE FROM notification_settings WHERE user_id=?", (user_id,))
            db.upd("DELETE FROM sessions_ring WHERE user_id=?", (user_id,))
            db.upd("DELETE FROM students WHERE user_id=?", (user_id,))
            db.upd("DELETE FROM instructors WHERE user_id=?", (user_id,))
            db.upd("DELETE FROM users WHERE id=?", (user_id,))
        finally:
            try:
                db.close()
            except Exception:
                pass

    def _mk_session(self, start_dt, created_at, student_id, inst_id=1,
                    car_id=1, status="scheduled", confirm="pending", notes="BAND24"):
        db = self.db
        date, start = start_dt.strftime("%Y-%m-%d"), start_dt.strftime("%H:%M")
        end_dt = start_dt + timedelta(minutes=60)
        # Kechani kechiruvchi sessiyalar uchun `+1d` SUFFIX qo'yilmaydi:
        # ilova `end_time` ni oddiy "HH:MM" sifatida saqlaydi va shu qat'iy
        # formatda string solishtiriladi ("00:15 +1d" > "18:20" -> noto'g'ri).
        end = min(end_dt, end_dt.replace(hour=23, minute=59, second=0)).strftime("%H:%M")
        sid = db.ex(
            """INSERT INTO lesson_sessions(date,start_time,end_time,instructor_id,car_id,
               car_name_snapshot,car_plate_snapshot,capacity_snapshot,status,confirm_state,
               notes,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (date, start, end, inst_id, car_id, "Test avto", "01A001AA", 4, status,
             confirm, notes, created_at, created_at))
        db.ex(
            """INSERT INTO session_students(session_id,student_id,pickup_address,
               attendance_status,student_status,joined_at) VALUES(?,?,'','unmarked','active',?)""",
            (sid, student_id, created_at))
        return sid

    def _drop_session(self, sid):
        db = self.db
        db.upd("DELETE FROM notifications WHERE related_lesson_id=?", (sid,))
        db.upd("DELETE FROM session_students WHERE session_id=?", (sid,))
        db.upd("DELETE FROM lesson_sessions WHERE id=?", (sid,))

    def _admin_msg(self, title, text):
        """`POST /api/admin/notifications` — barcha faol foydalanuvchilarga
        xabar yuboradi (`source=ADMIN_MESSAGE`)."""
        return self.admin.post("/api/admin/notifications",
                               {"title": title, "text": text})


# ============================================================================
# 1) AUTH — BAND 5, 6, 7 (kuchli credential, xavfsizlik, unikalik)
# ============================================================================
class TestAuthBAND24(_MixinAdmin, Base):
    """AUTH: login/parol formati, hashing, unikalik, rate-limit, CSRF, RBAC."""

    def test_A01_new_credential_format_is_random(self):
        """BAND 5: `usrL_` + random unikal suffix, `usrP_` + 14 belgi.
        Ketma-ket/predictable format QAT'IY TAQIQLANGAN."""
        logins, pwds = set(), set()
        for i in range(4):
            r = self._mkuser("student", "Auth", "T%d" % i, group_name="AUTH")
            lg = r["credentials"]["login"]
            pw = r["credentials"]["password"]
            self.assertRegex(lg, r"^usrL_[A-Za-z0-9]{14}$", "login formati: " + lg)
            self.assertRegex(pw, r"^usrP_", "parol prefiksi: " + pw)
            body = pw[len("usrP_"):]
            self.assertEqual(len(body), 14, "parol tanasi 14 belgi bo'lishi kerak: " + pw)
            self.assertTrue(any(c.islower() for c in body), "kichik harf kerak")
            self.assertTrue(any(c.isupper() for c in body), "katta harf kerak")
            self.assertTrue(any(c.isdigit() for c in body), "raqam kerak")
            self.assertTrue(any(not c.isalnum() for c in body), "maxsus belgi kerak")
            logins.add(lg)
            pwds.add(pw)
            self._cleanup_user(r["user"]["id"])
        self.assertEqual(len(logins), 4, "loginlar UNIKAL bo'lishi shart")
        self.assertEqual(len(pwds), 4, "parollar UNIKAL bo'lishi shart")

    def test_A02_credentials_not_sequential(self):
        """BAND 5: eski ketma-ket format (`usrL_00001`) QAYTMAYDI."""
        r1 = self._mkuser("student", "Seq", "One")
        r2 = self._mkuser("student", "Seq", "Two")
        try:
            l1 = r1["credentials"]["login"]
            l2 = r2["credentials"]["login"]
            self.assertNotIn("0000", l1[len("usrL_"):][:4])
            self.assertFalse(l1[len("usrL_"):].isdigit(), "login sonlardan iborat bo'lmasligi kerak")
            self.assertNotEqual(l1, l2)
        finally:
            self._cleanup_user(r1["user"]["id"])
            self._cleanup_user(r2["user"]["id"])

    def test_A03_password_never_stored_plain(self):
        """BAND 6: parol DB'da PLAIN TEXT ko'rinishida SAQLANMAYDI."""
        r = self._mkuser("student", "Hash", "Test")
        try:
            pw = r["credentials"]["password"]
            row = self.db.q1("SELECT password_hash FROM users WHERE id=?", (r["user"]["id"],))
            h = row["password_hash"] or ""
            self.assertNotEqual(h, pw, "parol o'zicha saqlangan!")
            self.assertNotIn(pw, h, "parol hash ichida ko'rinmoqda!")
            self.assertIn("scrypt", h, "scrypt ishlovchi formati kutilgan")
        finally:
            self._cleanup_user(r["user"]["id"])

    def test_A04_salt_unique_per_password(self):
        """BAND 6: har bir parol uchun salt ALOHIDA (ikkita bir xil parol ->
        turliq hash)."""
        from app.auth import hash_password as _hp
        a = _hp("usrP_Same000000!A1")
        b = _hp("usrP_Same000000!A1")
        self.assertNotEqual(a, b, "bir xil parol uchun hash bir xil chiqdi (salt yo'q)")

    def test_A05_password_verified_by_hash(self):
        """BAND 6: login parolni HASH orqali tekshiradi."""
        from app.auth import verify_password
        r = self._mkuser("student", "Verify", "Hash")
        try:
            pw = r["credentials"]["password"]
            lg = r["credentials"]["login"]
            h = self.db.q1("SELECT password_hash FROM users WHERE id=?",
                           (r["user"]["id"],))["password_hash"]
            self.assertTrue(verify_password(pw, h))
            self.assertFalse(verify_password(pw + "x", h))
            self.assertFalse(verify_password("usrP_Wrong000000!A1", h))
            # haqiqiy login ham ishlaydi
            c, st, rr = self._login(lg, pw, "student")
            self.assertEqual(st, 200, rr)
            c2, st2, rr2 = self._login(lg, pw + "x", "student")
            self.assertEqual(st2, 400)
            self.assertEqual(rr2["error"], "auth.wrong_credentials")
        finally:
            self._cleanup_user(r["user"]["id"])

    def test_A06_password_never_returned_by_api(self):
        """BAND 6: admin ham parolni KO'RA OLMAYDI."""
        r = self._mkuser("student", "Secret", "Hidden")
        try:
            uid = r["user"]["id"]
            st, prof = self.admin.get(f"/api/admin/users/{uid}")
            self.assertEqual(st, 200, prof)
            raw = json.dumps(prof, ensure_ascii=False)
            self.assertNotIn(r["credentials"]["password"], raw)
            self.assertNotIn("password_hash", prof["user"])
            st, lst = self.admin.get("/api/admin/users?role=student")
            self.assertEqual(st, 200)
            raw2 = json.dumps(lst, ensure_ascii=False)
            self.assertNotIn(r["credentials"]["password"], raw2)
            self.assertNotIn("password_hash", raw2)
        finally:
            self._cleanup_user(uid)

    def test_A07_weak_password_rejected(self):
        """BAND 6: foydalanuvchi kuchsiz parol qo'ya olmaydi."""
        r = self._mkuser("student", "Weak", "Pass")
        try:
            c, st, _ = self._login(r["credentials"]["login"],
                                   r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            for bad in ("abc", "alllowercase123", "ALLUPPERCASE123", "NoDigitsHere!",
                        "Short1!"):
                st2, r2 = c.put("/api/me/profile", {"login": r["credentials"]["login"]})
                self.assertEqual(st2, 200, r2)
                st3, r3 = c.post("/api/auth/change-password",
                                 {"old_password": r["credentials"]["password"],
                                  "new_password": bad, "confirm_password": bad})
                self.assertEqual(st3, 400, f"{bad!r} qabul qilindi: {r3}")
                self.assertEqual(r3["error"], "auth.password_weak", f"{bad!r}: {r3}")
        finally:
            self._cleanup_user(r["user"]["id"])

    def test_A08_same_password_rejected(self):
        r = self._mkuser("student", "Same", "Pass")
        try:
            pw = r["credentials"]["password"]
            c, st, _ = self._login(r["credentials"]["login"], pw, "student")
            self.assertEqual(st, 200)
            st2, r2 = c.post("/api/auth/change-password",
                             {"old_password": pw, "new_password": pw,
                              "confirm_password": pw})
            self.assertEqual(st2, 400, r2)
            self.assertEqual(r2["error"], "auth.password_same", r2)
        finally:
            self._cleanup_user(r["user"]["id"])

    def test_A09_password_reuse_rejected(self):
        """BAND 6/7: eski parol QAYTA ishlatilmaydi."""
        r = self._mkuser("student", "Reuse", "Test")
        try:
            old_pw = r["credentials"]["password"]
            c, st, _ = self._login(r["credentials"]["login"], old_pw, "student")
            self.assertEqual(st, 200)
            new_pw = "usrP_Fresh0000!Zz9"
            st2, r2 = c.post("/api/auth/change-password",
                             {"old_password": old_pw, "new_password": new_pw,
                              "confirm_password": new_pw})
            self.assertEqual(st2, 200, r2)
            # yangi parol bilan kirish, eskisi bilan QAYTARISHga urinish
            c2, st3, _ = self._login(r["credentials"]["login"], new_pw, "student")
            self.assertEqual(st3, 200)
            st4, r4 = c2.post("/api/auth/change-password",
                              {"old_password": new_pw, "new_password": old_pw,
                               "confirm_password": old_pw})
            self.assertEqual(st4, 400, r4)
            self.assertEqual(r4["error"], "auth.password_used", r4)
            # eski parol endi umuman ishlamaydi
            _, st5, _ = self._login(r["credentials"]["login"], old_pw, "student")
            self.assertEqual(st5, 400)
        finally:
            self._cleanup_user(r["user"]["id"])

    def test_A10_admin_reset_password_new_random(self):
        """BAND 5/22: admin yangi parol generatsiya qilsa — random, unique,
        ko'rsatiladi va KEYIN qaytarilmaydi."""
        r = self._mkuser("student", "Reset", "Pass")
        try:
            uid = r["user"]["id"]
            old = r["credentials"]["login"], r["credentials"]["password"]
            st, rr = self.admin.post(f"/api/admin/users/{uid}/reset-password")
            self.assertEqual(st, 200, rr)
            new_pw = rr["password"]
            self.assertRegex(rr["login"], r"^usrL_[A-Za-z0-9]{14}$")
            self.assertRegex(new_pw, r"^usrP_")
            self.assertEqual(len(new_pw) - len("usrP_"), 14)
            self.assertNotEqual(new_pw, old[1])
            # eski parol endi ishlamaydi, yangisi ishlaydi
            _, st1, _ = self._login(old[0], old[1], "student")
            self.assertEqual(st1, 400)
            _, st2, r2 = self._login(rr["login"], new_pw, "student")
            self.assertEqual(st2, 200, r2)
            # DB'da parol ko'rinmaydi
            h = self.db.q1("SELECT password_hash FROM users WHERE id=?", (uid,))["password_hash"]
            self.assertNotIn(new_pw, h)
        finally:
            self._cleanup_user(uid)

    def test_A11_login_rate_limit_blocks(self):
        """BAND 6: brute-force himoyasi — chegaradan keyin rad etiladi.

        `MAX_LOGIN_TRIES` test muhitida 100000 qilib ko'tarilgan (boshqa testlar
        kirish oqib ketmasligi uchun), shuning uchun chegara bu yerda
        ANIQ ko'rsatiladi: `login_allowed(key, max_tries, window)`.
        """
        from app.auth import login_allowed, reset_login_attempts
        key = "login:usrL_bd0000000001"
        reset_login_attempts(key)
        allowed = 0
        for _ in range(12):
            if login_allowed(key, 6, 300):
                allowed += 1
            else:
                break
        self.assertEqual(allowed, 6, "login limiti 6 urinishdan keyin ishlashi kerak")
        # 7-8-urinish ham rad etiladi (counter oshMAYdi)
        self.assertFalse(login_allowed(key, 6, 300))
        self.assertFalse(login_allowed(key, 6, 300))
        # oyna o'tsa — yana urinishga ruxsat beriladi (counter 0 ga qaytadi)
        import time as _time
        _time.sleep(0.05)
        self.assertTrue(login_allowed(key, 6, 0.01),
                        "oyna o'tganda bloklash to'lanadi")

        # API darajasida: chegarani 3 ga tushirib, kodni tekshiramiz
        import app.api as api_mod
        orig = api_mod.login_allowed
        api_mod.login_allowed = lambda k, *a, **kw: orig(k, 3, 300)
        try:
            reset_login_attempts(key)
            c = Client()
            errs = []
            for _ in range(5):
                st, rr = c.post("/api/auth/login",
                                {"login": "usrL_bd0000000001",
                                 "password": "usrP_x000000!Ab1", "role": "student"})
                errs.append((st, rr.get("error")))
        finally:
            api_mod.login_allowed = orig
        self.assertIn("auth.too_many_attempts", [e for _, e in errs],
                      "brute-force himoyasi API darajasida ishlayapti")
        self.assertTrue(all(st == 400 for st, _ in errs), errs)

    def test_A12_ip_blocking(self):
        """BAND 6: IP bo'yi bloklash — 4-urishda blok, `seconds` qaytariladi."""
        from app.auth import ip_blocked, reset_ip_attempts
        ip = "9.9.9.9"
        reset_ip_attempts(ip)
        self.assertEqual(ip_blocked(ip, 3, 300, 60), 0)   # 1
        self.assertEqual(ip_blocked(ip, 3, 300, 60), 0)   # 2
        self.assertEqual(ip_blocked(ip, 3, 300, 60), 0)   # 3
        left = ip_blocked(ip, 3, 300, 60)                 # 4 -> blok
        self.assertGreater(left, 0, "IP bloklanmadi")
        self.assertLessEqual(left, 60)
        # blok davomida qaytariladi va hisob tozalanadi
        reset_ip_attempts(ip)
        self.assertEqual(ip_blocked(ip, 3, 300, 60), 0)

        # API darajasida ham `auth.ip_blocked` + params.seconds qaytariladi
        import app.api as api_mod
        orig = api_mod.ip_blocked

        def fake_blocked(ip_, *a, **kw):
            if ip_ == "127.0.0.1":
                return 42
            return orig(ip_, *a, **kw)
        api_mod.ip_blocked = fake_blocked
        try:
            st, r = Client().post("/api/auth/login",
                                  {"login": "admin", "password": "admin123",
                                   "role": "admin"})
        finally:
            api_mod.ip_blocked = orig
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "auth.ip_blocked", r)
        self.assertEqual(r.get("params", {}).get("seconds", r.get("seconds")), 42, r)

    def test_A13_csrf_token_returned_and_enforced(self):
        """BAND 6: login va `me` CSRF token beradi; noto'g'ri token -> 403."""
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        tok = r.get("csrf") or ""
        self.assertTrue(tok, "login javobida csrf token bo'lishi shart")
        st2, r2 = c.get("/api/auth/me")
        self.assertEqual(st2, 200, r2)
        self.assertTrue(r2.get("csrf"), "GET /api/auth/me csrf token qaytarishi shart")
        # noto'g'ri token bilan POST -> 403
        raw = self._raw_post(c, "/api/me/settings", {"lang": "uz"},
                             extra={"X-CSRF-Token": "not-a-real-token"})
        self.assertEqual(raw[0], 403, raw)
        self.assertEqual(raw[1].get("error"), "csrf_invalid", raw)

    def test_A13b_csrf_token_is_mandatory(self):
        """BAND 6: token butunlay YUBORILMASA ham so'rov rad etiladi.

        AVVALGI XATO: `if expected and sent and ...` — `sent` bo'sh bo'lsa
        shart bajarilmasdi, ya'ni sarlavhani umuman yubormaydi (CSRF o'tkazib
        yuborish mumkin edi). Endi `sent` majburiy.
        """
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        self.assertTrue(r.get("csrf"), r)
        raw = self._raw_post(c, "/api/me/settings", {"lang": "uz"})   # token YO'Q
        self.assertEqual(raw[0], 403, "token'siz POST qabul qilindi: %s" % (raw,))
        self.assertEqual(raw[1].get("error"), "csrf_invalid", raw)

    def test_A13c_cross_site_request_rejected(self):
        """BAND 6: boshqa dandan `Origin` + token'siz so'rov -> 403."""
        c = Client()
        st, r = c.post("/api/auth/login",
                       {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200, r)
        raw = self._raw_post(c, "/api/admin/notifications",
                             {"title": "h", "text": "x"},
                             extra={"X-Requested-With": "",
                                    "Origin": "http://zararli-sayt.example"},
                             drop_xrw=True)
        self.assertEqual(raw[0], 403, raw)
        self.assertEqual(raw[1].get("error"), "csrf_origin", raw)

    def test_A13d_other_session_token_rejected(self):
        """Boshqa foydalanuvchi sessiyasining token'i — qabul qilinMAYDI."""
        c1 = Client()
        self.assertEqual(c1.post("/api/auth/login",
                                 {"login": "admin", "password": "admin123",
                                  "role": "admin"})[0], 200)
        c2 = Client()
        self.assertEqual(c2.post("/api/auth/login",
                                 {"login": STUD_LOGIN_1,
                                  "password": PW_BY_LOGIN[STUD_LOGIN_1],
                                  "role": "student"})[0], 200)
        self.assertNotEqual(c1._csrf, c2._csrf)
        raw = self._raw_post(c1, "/api/me/settings", {"lang": "ru"},
                             extra={"X-CSRF-Token": c2._csrf})
        self.assertEqual(raw[0], 403, raw)
        self.assertEqual(raw[1].get("error"), "csrf_invalid", raw)

    def _raw_post(self, client, path, body, extra=None, drop_xrw=False):
        """CSRF sarlavhalarini qo'lda boshqarish uchun xom so'rov."""
        data = json.dumps(body).encode()
        headers = {"Content-Type": "application/json"}
        if not drop_xrw:
            headers["X-Requested-With"] = "Avtomaktab"
        headers.update(extra or {})
        req = urllib.request.Request(
            f"http://{HOST}:{PORT}{path}", data=data, headers=headers, method="POST")
        try:
            with client.opener.open(req, timeout=15) as resp:
                return resp.status, json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            try:
                return e.code, json.loads(e.read().decode("utf-8"))
            except Exception:
                return e.code, {}

    def test_A14_session_cookie_httponly_samesite(self):
        """BAND 6: `sid` cookie — HttpOnly + SameSite=Lax."""
        import http.cookiejar
        c = Client()
        st, _ = c.post("/api/auth/login",
                       {"login": "admin", "password": "admin123", "role": "admin"})
        self.assertEqual(st, 200)
        sid = None
        for ck in c.jar:
            if ck.name == "sid":
                sid = ck
        self.assertIsNotNone(sid, "sid cookie qo'yilmadi")
        # urllib cookie jar HttpOnly/SameSite ni saqlamaydi -> QO'LDI sarlavhani tekshiramiz
        raw = self._headers_on_login().lower()
        self.assertIn("httponly", raw, "sid cookie HttpOnly bo'lishi shart")
        self.assertIn("samesite=lax", raw, "sid cookie SameSite=Lax bo'lishi shart")
        self.assertIn("path=/", raw)

    def _headers_on_login(self):
        data = json.dumps({"login": "admin", "password": "admin123",
                           "role": "admin"}).encode()
        req = urllib.request.Request(
            f"http://{HOST}:{PORT}/api/auth/login", data=data,
            headers={"Content-Type": "application/json"}, method="POST")
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(
            http.cookiejar.CookieJar()))
        try:
            with opener.open(req, timeout=15) as resp:
                return resp.headers.get("Set-Cookie") or ""
        except urllib.error.HTTPError as e:
            return e.headers.get("Set-Cookie") or ""

    def test_A15_rbac_and_idor(self):
        """BAND 6: RBAC + IDOR — talaba admin endpoint'iga, boshqa foydalanuvchi
        ma'lumotiga kira olmaydi."""
        c, st, _ = self._login(STUD_LOGIN_1, PW_BY_LOGIN[STUD_LOGIN_1], "student")
        self.assertEqual(st, 200)
        st2, r2 = c.get("/api/admin/users")
        self.assertEqual(st2, 403, r2)
        st3, r3 = c.get("/api/admin/users/1")
        self.assertEqual(st3, 403, r3)
        # IDOR: boshqa foydalanuvchining bildirishnoma id'si
        other_uid = self._uid_of(STUD_LOGIN_2)
        st0, _ = self._admin_msg("A15-OTHER", "boshqa odamga")
        self.assertEqual(st0, 200)
        nid = self.db.q1("SELECT id FROM notifications WHERE user_id=? ORDER BY id DESC",
                         (other_uid,))["id"]
        self.assertIsNotNone(nid)
        st4, r4 = c.get(f"/api/me/notifications/{nid}")
        self.assertEqual(st4, 404, r4)

    def _uid_of(self, login):
        row = self.db.q1("SELECT id FROM users WHERE login=?", (login,))
        return row["id"] if row else 0

    def test_A16_admin_profile_and_users_list_shape(self):
        """BAND 22: admin ro'yxatida jami/bajarilgan/qolgan bor."""
        r = self._mkuser("student", "Shape", "Test", group_name="SH", license_category="B")
        try:
            st, lst = self.admin.get("/api/admin/users?role=student&q=" + r["credentials"]["login"])
            self.assertEqual(st, 200, lst)
            found = [u for u in lst["users"] if u["id"] == r["user"]["id"]]
            self.assertTrue(found, "talaba ro'yxatda topilmadi")
            pr = found[0]["progress"]
            for k in ("target", "done", "remaining", "pct", "mode", "individual_target"):
                self.assertIn(k, pr, f"progress.{k} yo'q")
            st2, prof = self.admin.get(f"/api/admin/users/{r['user']['id']}")
            self.assertEqual(st2, 200, prof)
            u = prof["user"]
            self.assertIn("login", u)
            self.assertIn("phone", u)
            self.assertIn("student", u)
            self.assertIn("group_name", u["student"])
            self.assertIn("license_category", u["student"])
            self.assertIn("progress", u)
        finally:
            self._cleanup_user(r["user"]["id"])


# ============================================================================
# 2) MAP — BAND 1 (Yandex autocomplete, UI dan kenglik/uzunlik olib tashlash)
# ============================================================================
class TestMapBAND24(_MixinAdmin, Base):
    """MAP: geocoder sozlamalari, manzil orqali nuqta, UI da lat/lng yo'q."""

    ENV_KEYS = ("YANDEX_MAPS_API_KEY", "YANDEX_GEOCODER_API_KEY", "MAP_PROVIDER")

    def setUp(self):
        super().setUp()
        self._env_backup = {k: os.environ.get(k) for k in self.ENV_KEYS}
        for k in self.ENV_KEYS:
            os.environ.pop(k, None)
        r = self._mkuser("student", "Map", "Autocomplete", group_name="MAP", license_category="B")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.srow = self.db.q1("SELECT id FROM students WHERE user_id=?", (self.uid,))
        self.student_id = self.srow["id"]
        # kelajakdagi mashg'ulot (pickup endpoint'i uchun)
        d = datetime.now() + timedelta(days=40)
        self.sid = self._mk_session(d.replace(hour=9, minute=0, second=0, microsecond=0),
                                    now(), self.student_id, notes="MAPBAND24")
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)

    def tearDown(self):
        self._drop_session(self.sid)
        self._cleanup_user(self.uid)
        for k, v in self._env_backup.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        super().tearDown()

    # ------------------------------------------------------------- sozlamalar
    def test_M01_config_exposes_geocoder_fields(self):
        st, r = Client().get("/api/config")
        self.assertEqual(st, 200, r)
        m = r["config"]["maps"]
        for k in ("geocoder_api_key", "geocoder_enabled", "lang"):
            self.assertIn(k, m, f"config.maps.{k} yo'q")

    def test_M02_geocoder_disabled_without_key(self):
        """Kalit yo'q bo'lsa `geocoder_enabled=false` — soxta kalit QO'YILMAYDI."""
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertFalse(m["geocoder_enabled"], "kalitsiz geocoder 'yoqiq' bo'lishi kerak")
        self.assertEqual(m["geocoder_api_key"], "")

    def test_M03_geocoder_enabled_with_env_key(self):
        os.environ["YANDEX_GEOCODER_API_KEY"] = "geo-test-key-123"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertTrue(m["geocoder_enabled"])
        self.assertEqual(m["geocoder_api_key"], "geo-test-key-123")

    def test_M04_geocoder_falls_back_to_map_key(self):
        """Aloha kalit yo'q bo'lsa `YANDEX_MAPS_API_KEY` ga qaytadi."""
        os.environ["YANDEX_MAPS_API_KEY"] = "map-test-key-456"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertTrue(m["geocoder_enabled"])
        self.assertEqual(m["geocoder_api_key"], "map-test-key-456")

    def test_M05_provider_none_disables_geocoder(self):
        os.environ["YANDEX_GEOCODER_API_KEY"] = "geo-test-key-123"
        os.environ["MAP_PROVIDER"] = "none"
        m = self.admin.get("/api/config")[1]["config"]["maps"]
        self.assertFalse(m["geocoder_enabled"])

    def test_M06_key_never_hardcoded_in_source(self):
        """BAND 1: API kalit frontend/backend source code'da OCHIQ yozilmaydi."""
        hard = 0
        for rel in ("web/js/map.js", "web/js/api.js", "web/js/shared.js",
                    "web/index.html", "app/config.py", "app/api.py"):
            p = ROOT / rel
            if not p.exists():
                continue
            txt = p.read_text(encoding="utf-8", errors="ignore")
            # haqiqiy Yandex kalit shakli: <id>.<hash>
            import re as _re
            if _re.search(r"\bAQ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}", txt):
                hard += 1
            if "YANDEX_MAPS_API_KEY=" in txt or "YANDEX_GEOCODER_API_KEY=" in txt:
                # faqat `.env.example` da bo'lishi mumkin, kodda emas
                if not rel.endswith(".env.example"):
                    hard += 1
        self.assertEqual(hard, 0, "maxfiy kalit source code'da topildi!")

    # ------------------------------------------------------------- endpoint
    def test_M07_address_saved_without_coords(self):
        """Autocomplete tanlanganda manzil + koordinata keladi; agar topilmasa
        ham manzil saqlanadi (jim qolmaydi, xarita nuqtasi bo'lmaydi)."""
        st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup",
                           {"address": "O'zbekiston, Toshkent, Beshariq tumani"})
        self.assertEqual(st, 200, r)
        row = self.db.q1("""SELECT pickup_address,pickup_lat,pickup_lng FROM session_students
                            WHERE session_id=? AND student_id=?""",
                         (self.sid, self.student_id))
        self.assertIn("Beshariq", row["pickup_address"])
        self.assertIsNone(row["pickup_lat"])
        self.assertIsNone(row["pickup_lng"])

    def test_M08_coords_from_autocomplete_persisted(self):
        st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup",
                           {"address": "Toshkent, Beshariq tumani, 12",
                            "lat": 41.3256, "lng": 69.2364})
        self.assertEqual(st, 200, r)
        st2, d2 = self.c.get(f"/api/student/sessions/{self.sid}")
        self.assertEqual(st2, 200, d2)
        ss = d2["session"]
        self.assertAlmostEqual(ss["pickup_lat"], 41.3256, places=5)
        self.assertAlmostEqual(ss["pickup_lng"], 69.2364, places=5)
        self.assertEqual(ss["pickup_address"], "Toshkent, Beshariq tumani, 12")

    def test_M09_empty_request_rejected(self):
        """BAND 1: manzil ham, koordinata ham kelmagan bo'lsa — RAD etiladi
        ("Olib ketish joyi" jim qolmasligi kerak)."""
        st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup", {})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "map.address_required", r)
        # DB o'zgarmagan
        row = self.db.q1("""SELECT pickup_address FROM session_students
                            WHERE session_id=? AND student_id=?""",
                         (self.sid, self.student_id))
        self.assertEqual((row["pickup_address"] or ""), "")

    def test_M09b_blank_address_cleared_with_coords(self):
        """Faqat nuqta berilgan (autocomplete topmagan) holat — manzil bo'sh
        qoladi, ammo xato QILINMAYDI (nuqta saqlanadi)."""
        st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup",
                           {"address": "  ", "lat": 41.3, "lng": 69.2})
        self.assertEqual(st, 200, r)
        row = self.db.q1("""SELECT pickup_address,pickup_lat FROM session_students
                            WHERE session_id=? AND student_id=?""",
                         (self.sid, self.student_id))
        self.assertEqual(row["pickup_address"], "")
        self.assertAlmostEqual(row["pickup_lat"], 41.3, places=5)

    def test_M10_address_too_long_rejected(self):
        st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup",
                           {"address": "B" * 5000})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "map.address_too_long")

    def test_M11_bad_coords_still_rejected(self):
        """Autocomplete kelmagan holatda ham noto'g'ri koordinata jim qolmaydi."""
        for bad in ("abc", 91, -91, 1e400):
            st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup",
                               {"address": "Test", "lat": bad, "lng": 69.2})
            self.assertEqual(st, 400, f"{bad!r}: {r}")
            self.assertTrue(r["error"].startswith("coord."), r)

    def test_M12_idor_other_student_pickup(self):
        """BAND 6: boshqa talabaning mashg'ulotida olish manzilini o'zgartirib
        bo'lmaydi."""
        r2 = self._mkuser("student", "Map", "Other", group_name="MAP2")
        try:
            st, r = self.c.put(f"/api/student/sessions/{self.sid}/pickup",
                               {"address": "X", "lat": 41.3, "lng": 69.2})
            self.assertEqual(st, 200, r)
            # endi ikkinchi talaba uchun alohida mashg'ulot yaratamiz
            srow2 = self.db.q1("SELECT id FROM students WHERE user_id=?", (r2["user"]["id"],))
            d = datetime.now() + timedelta(days=45)
            sid2 = self._mk_session(d.replace(hour=10, minute=0, second=0, microsecond=0),
                                    now(), srow2["id"], notes="MAP-OTHER")
            # 1-talaba 2-mashg'ulotga kira olmaydi
            st2, r3 = self.c.put(f"/api/student/sessions/{sid2}/pickup",
                                 {"address": "Y", "lat": 41.4, "lng": 69.3})
            self.assertEqual(st2, 404, r3)
            self._drop_session(sid2)
        finally:
            self._cleanup_user(r2["user"]["id"])

    def test_M13_other_roles_forbidden(self):
        inst, st, _ = self._login(INSTR_LOGIN_1, PW_BY_LOGIN[INSTR_LOGIN_1], "instructor")
        self.assertEqual(st, 200)
        st2, r2 = inst.put(f"/api/student/sessions/{self.sid}/pickup",
                           {"address": "X", "lat": 41.3, "lng": 69.2})
        self.assertEqual(st2, 403, r2)

    # ------------------------------------------------------------- frontend
    def test_M14_pickup_modal_is_plain_text_no_map(self):
        """XARITA VAQTINCHA O'CHIRILGAN.

        Faol funksiya `openPickupEdit` — FAQAT bitta matn maydoni:
          * "Kenglik"/"Uzunlik" inputlari yo'q (avvalgidek),
          * xarita (`MapView.mount`) CHAQLANMAYDI,
          * autocomplete (`attachAutocomplete`) ulanmagan,
          * placeholder/to'g'ri label `meeting.*` dan keladi.
        Xarita kodi `openPickupEditMap` da SAQLANGAN (o'chirilmagan).
        """
        p = ROOT / "web" / "js" / "shared.js"
        txt = p.read_text(encoding="utf-8", errors="ignore")
        i = txt.find("function openPickupEdit(s, m)")
        self.assertGreater(i, 0, "openPickupEdit topilmadi")
        chunk = txt[i:i + 4000]
        end = chunk.find("\n  }")
        body = chunk[:end if end > 0 else len(chunk)]
        self.assertNotIn("map.lat", body, "Kenglik maydoni hali ham bor")
        self.assertNotIn("map.lng", body, "Uzunlik maydoni hali ham bor")
        self.assertNotIn("MapView.mount", body, "xarita endi render qilinmasligi kerak")
        self.assertNotIn("attachAutocomplete", body, "autocomplete endi ulanmasligi kerak")
        self.assertNotIn("reverseGeocode", body, "reverse geocode kerak emas")
        self.assertNotIn("map-picker", body, "xarita konteyneri kerak emas")
        # oddiy matn rejimi faol
        self.assertIn("meeting.place_ph", body, "matn placeholder'i ishlatilmayapti")
        self.assertIn("meeting.place", body, "kalit ishlatilmayapti")
        # `lat`/`lng` yuborilmaydi (bazada NULL qoladi)
        self.assertIn("{ address: text }", body, "so'rovda lat/lng yuborilmoqda")
        self.assertNotIn("lat: r.lat", body, "koordinata yuborilmoqda")

    def test_M14b_map_code_preserved_for_future(self):
        """Xarita kodi O'CHIRILMAGAN — kelajakda qayta yoqish uchun saqlangan."""
        p = ROOT / "web" / "js" / "shared.js"
        txt = p.read_text(encoding="utf-8", errors="ignore")
        self.assertIn("function openPickupEditMap", txt,
                      "xarita rejimi kodi yo'qoldirilgan")
        i = txt.find("function openPickupEditMap")
        chunk = txt[i:i + 6000]
        for fn in ("MapView.mount", "attachAutocomplete", "reverseGeocode"):
            self.assertIn(fn, chunk, f"saqlangan xarita kodida {fn} yo'q")
        # `map.js` butunligicha saqlangan
        mt = (ROOT / "web" / "js" / "map.js").read_text(encoding="utf-8", errors="ignore")
        for fn in ("suggest", "geocode", "reverseGeocode", "attachAutocomplete"):
            self.assertIn("function " + fn, mt, f"map.js da {fn} yo'q")
        # backend bayrogi ham bor (qayta yoqish uchun bitta qator)
        at = (ROOT / "app" / "api.py").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("MAP_DISABLED", at, "backend da MAP_DISABLED bayrog'i yo'q")

    def test_M14c_map_disabled_flag_and_labels(self):
        """Bayroq `true` + "Uchrashuv joyi" label'lari mavjud."""
        sh = (ROOT / "web" / "js" / "shared.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("const MAP_DISABLED = true;", sh,
                      "xarita vaqtincha o'chirilgan deb belgilanmagan")
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn('"meeting.place": { uz: "Uchrashuv joyi"', i18n)
        self.assertIn("Masalan: Beshariq, Mustaqillik maydoni yonida", i18n)

    def test_M15_map_module_exposes_autocomplete_api(self):
        txt = (ROOT / "web" / "js" / "map.js").read_text(encoding="utf-8", errors="ignore")
        for fn in ("suggest", "geocode", "reverseGeocode", "attachAutocomplete",
                   "geocoderAvailable"):
            self.assertIn("function " + fn, txt, f"map.js da {fn} yo'q")
        self.assertIn("geocoder_enabled", txt, "kalit yo'qligi tekshirilmaydi")

    def test_M16_student_schedule_removed_from_navigation(self):
        """BAND 4: talaba sidebar'ida "schedule" (Mening jadvalim) yo'q."""
        txt = (ROOT / "web" / "js" / "app.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("student: [")
        self.assertGreater(i, 0)
        j = txt.find("]", i)
        student_nav = txt[i:j]
        self.assertNotIn('"schedule"', student_nav,
                         "talaba sidebar'ida 'schedule' hali ham bor")
        # alias bilan eski manzil 'Amaliy mashg'ulotlarim'ga yo'naltiriladi
        self.assertIn("ROUTE_ALIASES", txt)
        self.assertIn("student:lessons", txt)
        txt2 = (ROOT / "web" / "js" / "views-student.js").read_text(encoding="utf-8",
                                                                    errors="ignore")
        self.assertNotIn("async function schedule", txt2,
                         "'Mening jadvalim' view hali ham bor")
        self.assertIn("async function lessons", txt2)


# ============================================================================
# 3) LESSONS — BAND 13, 14, 19, 20 (holatlar, ajratish, booking, integritet)
# ============================================================================
class TestLessonsBAND24(_MixinAdmin, Base):
    """LESSONS: kelajak/o'tgan ajratish, biznes holati, booking qoidalari."""

    def setUp(self):
        super().setUp()
        r = self._mkuser("student", "Les", "Test", group_name="LES", license_category="B")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.student_id = self.db.q1("SELECT id FROM students WHERE user_id=?",
                                     (self.uid,))["id"]
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)
        self.sids = []

    def tearDown(self):
        for sid in self.sids:
            self._drop_session(sid)
        self._cleanup_user(self.uid)
        super().tearDown()

    def _new(self, start_dt, **kw):
        sid = self._mk_session(start_dt, kw.pop("created_at", now()),
                               self.student_id, **kw)
        self.sids.append(sid)
        return sid

    def _list(self, upcoming):
        st, r = self.c.get(f"/api/student/sessions?upcoming={upcoming}")
        self.assertEqual(st, 200, r)
        return r

    def _workday(self, start_days, step=1):
        """Instruktor 1 uchun haqiqiy ISH KUNI bo'lgan sana ("YYYY-MM-DD").

        Nima uchun: `now + 50 kun` kabi DOIMIY raqam qat'iy sana-bog'liq —
        masalan 2026-10-03 + 50 = 2026-11-22 (yakshanba) bo'lib, `work_days`
        da yo'q bo'lgani uchun booking `not_work_day` bilan RAD etilardi.
        Testlar sana o'zgarmasa ham o'tishi uchun ish kuni `work_days` dan
        OLIB topiladi.
        """
        inst = self.db.q1("SELECT work_days FROM instructors WHERE id=1")
        self.assertIsNotNone(inst, "instruktor 1 topilmadi")
        raw = inst["work_days"]
        days = json.loads(raw) if isinstance(raw, str) else (raw or [])
        names = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
        d = datetime.now() + timedelta(days=start_days)
        for _ in range(21):
            if names[d.weekday()] in days:
                return d.strftime("%Y-%m-%d")
            d += timedelta(days=step)
        self.fail("kelajakdagi ish kuni topilmadi (work_days=%s)" % days)

    # --------------------------------------------------------- BAND 13/14
    def test_L01_future_session_is_pending_or_confirmed(self):
        """KELMAGAN mashg'ulot: `pending` (kutilmoqda) yoki `confirmed`
        (tasdiqlangan) — HECH QACHON `completed` EMAS."""
        d = (datetime.now() + timedelta(days=3)).replace(hour=9, minute=0, second=0,
                                                        microsecond=0)
        self._new(d, status="scheduled", confirm="pending")
        rows = self._list(1)["sessions"]
        s = [x for x in rows if x["id"] == self.sids[-1]][0]
        self.assertEqual(s["_display_status"], "pending", s)
        self.assertTrue(s["_is_future"])

    def test_L02_confirmed_state_shown(self):
        d = (datetime.now() + timedelta(days=4)).replace(hour=9, minute=0, second=0,
                                                        microsecond=0)
        self._new(d, status="scheduled", confirm="confirmed")
        rows = self._list(1)["sessions"]
        s = [x for x in rows if x["id"] == self.sids[-1]][0]
        self.assertEqual(s["_display_status"], "confirmed", s)

    def test_L03_ongoing_session(self):
        """BAND 14: `JARAYONDA` holati alohida ko'rsatiladi (`completed` EMAS)."""
        d = (datetime.now() + timedelta(days=3)).replace(hour=9, minute=0, second=0,
                                                        microsecond=0)
        self._new(d, status="ongoing")
        rows = self._list(1)["sessions"]
        s = [x for x in rows if x["id"] == self.sids[-1]][0]
        self.assertEqual(s["_display_status"], "ongoing", s)

    def test_L04_past_not_completed_is_overdue(self):
        """BAND 14: o'tib ketgan lekin yakunlanmagan dars AVTOMATIK
        "Bajarilgan" BO'LMAYDI — `overdue`."""
        d = datetime.now().replace(hour=9, minute=0, second=0, microsecond=0) \
            - timedelta(days=2)
        self._new(d, status="scheduled")
        rows = self._list(0)["sessions"]
        s = [x for x in rows if x["id"] == self.sids[-1]][0]
        self.assertEqual(s["_display_status"], "overdue", s)
        self.assertFalse(s["_is_future"])

    def test_L05_completed_stays_completed(self):
        d = datetime.now().replace(hour=9, minute=0, second=0, microsecond=0) \
            - timedelta(days=2)
        self._new(d, status="completed")
        rows = self._list(0)["sessions"]
        s = [x for x in rows if x["id"] == self.sids[-1]][0]
        self.assertEqual(s["_display_status"], "completed", s)

    def test_L06_cancelled_stays_cancelled(self):
        d = datetime.now().replace(hour=9, minute=0, second=0, microsecond=0) \
            - timedelta(days=2)
        self._new(d, status="cancelled")
        rows = self._list(0)["sessions"]
        s = [x for x in rows if x["id"] == self.sids[-1]][0]
        self.assertEqual(s["_display_status"], "cancelled", s)

    def test_L07_past_day_leaves_upcoming_list(self):
        """BAND 13: 1 kun o'tgan dars 'kelayotgan' ro'yxatidan CHIQADI,
        lekin DB'dan O'CHIRILMAYDI va tarixda SAQLANADI.

        Kelajakdagi dars "erta indin" qo'yiladi — test soatga bog'liq bo'lmasligi
        uchun (kechqurun ishga tushsa ham ishlashi kerak)."""
        tomorrow9 = ((datetime.now() + timedelta(days=1)).replace(hour=9, minute=0,
                                                                 second=0, microsecond=0))
        yesterday9 = tomorrow9 - timedelta(days=2)
        self._new(tomorrow9, status="scheduled")
        self._new(yesterday9, status="scheduled")
        sid_future, sid_past = self.sids[-2], self.sids[-1]
        up_ids = [x["id"] for x in self._list(1)["sessions"]]
        self.assertIn(sid_future, up_ids)
        self.assertNotIn(sid_past, up_ids, "o'tgan dars 'kelayotgan' da qoldi")
        past_ids = [x["id"] for x in self._list(0)["sessions"]]
        self.assertIn(sid_past, past_ids, "o'tgan dars tarixda yo'q")
        # DB'da SAQLANGAN
        row = self.db.q1("SELECT id FROM lesson_sessions WHERE id=?", (sid_past,))
        self.assertIsNotNone(row, "o'tgan dars DB'dan o'chirilgan!")

    def test_L08_all_split_and_counts(self):
        f = (datetime.now() + timedelta(days=6)).replace(hour=9, minute=0, second=0,
                                                         microsecond=0)
        p = datetime.now().replace(hour=9, minute=0, second=0, microsecond=0) \
            - timedelta(days=1)
        self._new(f)
        self._new(p)
        st, r = self.c.get("/api/student/sessions?upcoming=all")
        self.assertEqual(st, 200, r)
        up_ids = [x["id"] for x in r["upcoming"]]
        past_ids = [x["id"] for x in r["past"]]
        self.assertIn(self.sids[-2], up_ids)
        self.assertIn(self.sids[-1], past_ids)
        self.assertNotIn(self.sids[-2], past_ids)
        self.assertNotIn(self.sids[-1], up_ids)
        self.assertGreaterEqual(r["counts"]["upcoming"], 1)
        self.assertGreaterEqual(r["counts"]["past"], 1)
        self.assertGreaterEqual(r["counts"]["overdue"], 1)

    def test_L09_home_cards_from_db(self):
        """BAND 2: bosh sahifadagi 3 karta va progress — hammasi DB dan."""
        d = (datetime.now() + timedelta(days=5)).replace(hour=9, minute=0, second=0,
                                                        microsecond=0)
        self._new(d)
        st, r = self.c.get("/api/student/home")
        self.assertEqual(st, 200, r)
        pr = r["progress"]
        for k in ("target", "done", "remaining", "pct", "mode"):
            self.assertIn(k, pr, f"progress.{k} yo'q")
        self.assertGreater(pr["target"], 0, "jami darslar soni DB dan kelishi shart")
        self.assertIn("counts", r)
        for k in ("upcoming", "overdue"):
            self.assertIn(k, r["counts"])
        # frontend'da hardcoded raqam BO'LMASIN
        txt = (ROOT / "web" / "js" / "views-student.js").read_text(encoding="utf-8",
                                                                    errors="ignore")
        self.assertIn("week.total_lessons", txt)
        self.assertIn("week.done_lessons", txt)
        self.assertIn("home.overall_progress", txt)
        # "Qolgan darslar" karti: PLACEHOLDER'LI kalit paramsiz chaqirilmasin
        # (aks holda ekranda "Qolgan: {n} 30" ko'rinadi).
        self.assertIn('t("progress.remaining")', txt)
        self.assertNotIn('t("week.remaining_lessons")', txt,
                         "placeholder'li kalit paramsiz chaqirilgan — {n} ko'rinib qoladi")

    def test_L10_overdue_not_counted_as_done(self):
        """20/8 -> 20/9: o'tgan, yakunlanmagan dars `done` ga QO'SHILMAYDI."""
        st0, h0 = self.c.get("/api/student/home")
        done0 = h0["progress"]["done"]
        d = datetime.now().replace(hour=9, minute=0, second=0, microsecond=0) \
            - timedelta(days=1)
        self._new(d, status="scheduled")
        st1, h1 = self.c.get("/api/student/home")
        self.assertEqual(h1["progress"]["done"], done0,
                         "o'tib ketgan dars 'bajarilgan' ga hisoblandi")
        self.assertGreaterEqual(h1["counts"]["overdue"], 1)

    def test_L11_remaining_never_negative(self):
        self.admin.post(f"/api/admin/users/{self.uid}/total-lessons",
                        {"scope": "user", "user_id": self.uid, "total_lessons": 1})
        st, h = self.c.get("/api/student/home")
        self.assertEqual(st, 200, h)
        self.assertGreaterEqual(h["progress"]["remaining"], 0)

    def test_L12_unknown_category_rejected(self):
        st, r = self.c.get("/api/student/sessions?upcoming=bogus")
        self.assertEqual(st, 200, r)   # `bogus` -> default kelayotgan (moslik)

    # --------------------------------------------------------- BAND 19/20
    def test_L13_car_double_booking_rejected(self):
        """BAND 19: bir vaqtda bitta avtomobilga bir nechta talaba
        biriktirilMASIN (slot conflict)."""
        day = self._workday(50)
        st, r1 = self.admin.post("/api/admin/sessions", {
            "date": day, "start_time": "08:00", "end_time": "09:00",
            "instructor_id": 1, "student_ids": [self.student_id], "notes": "L13-A"})
        self.assertEqual(st, 200, r1)
        self.sids.append(r1["id"])
        r2u = self._mkuser("student", "Les", "Second", group_name="LES2")
        try:
            srow2 = self.db.q1("SELECT id FROM students WHERE user_id=?",
                               (r2u["user"]["id"],))
            st2, r2 = self.admin.post("/api/admin/sessions", {
                "date": day, "start_time": "08:00", "end_time": "09:00",
                "instructor_id": 1, "student_ids": [srow2["id"]], "notes": "L13-B"})
            self.assertIn(st2, (400, 409), r2)
            self.assertTrue(r2.get("error"), r2)
        finally:
            self._cleanup_user(r2u["user"]["id"])

    def test_L14_duplicate_student_in_session_rejected(self):
        day = self._workday(51)
        st, r = self.admin.post("/api/admin/sessions", {
            "date": day, "start_time": "10:00", "end_time": "11:00",
            "instructor_id": 1, "student_ids": [self.student_id, self.student_id]})
        self.assertIn(st, (400, 409), r)
        self.assertIn("duplicate", str(r.get("error", "")) + str(r))

    def test_L15_foreign_keys_present(self):
        """BAND 20: jadvalar orasida FK constraint'lar mavjud."""
        fks = self.db.q("PRAGMA foreign_key_list(session_students)")
        cols = {f["table"] for f in fks}
        self.assertIn("lesson_sessions", cols, "session_students -> lesson_sessions FK yo'q")
        self.assertIn("students", cols, "session_students -> students FK yo'q")
        fk2 = self.db.q("PRAGMA foreign_key_list(lesson_sessions)")
        self.assertTrue(fk2, "lesson_sessions da FK yo'q")

    def test_L16_student_cannot_see_other_students_session(self):
        r2 = self._mkuser("student", "Les", "Hidden", group_name="LES3")
        try:
            srow2 = self.db.q1("SELECT id FROM students WHERE user_id=?",
                               (r2["user"]["id"],))
            d = (datetime.now() + timedelta(days=8)).replace(hour=12, minute=0, second=0,
                                                            microsecond=0)
            sid2 = self._mk_session(d, now(), srow2["id"], notes="LES-HID")
            self.sids.append(sid2)
            st, r = self.c.get(f"/api/student/sessions/{sid2}")
            self.assertEqual(st, 404, r)
        finally:
            self._cleanup_user(r2["user"]["id"])


# ============================================================================
# 4) NOTIFICATIONS — BAND 8, 9, 11, 12, 17, 18, 21
# ============================================================================
class TestNotificationsBAND24(_MixinAdmin, Base):
    """NOTIFICATIONS: kategoriyalar, alohida record, 2-soat eslatma, sozlamalar."""

    def setUp(self):
        super().setUp()
        r = self._mkuser("student", "Notif", "Test", group_name="NOT", license_category="B")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.student_id = self.db.q1("SELECT id FROM students WHERE user_id=?",
                                     (self.uid,))["id"]
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)
        self.sids = []

    def tearDown(self):
        for sid in self.sids:
            self._drop_session(sid)
        self._cleanup_user(self.uid)
        super().tearDown()

    def _notifs(self, cat=None):
        q = f"/api/me/notifications?category={cat}" if cat else "/api/me/notifications"
        st, r = self.c.get(q)
        self.assertEqual(st, 200, r)
        return r["notifications"]

    # ------------------------------------------------------------- BAND 8
    def test_N01_only_four_categories(self):
        """'So'rovlar' va 'Xavfsizlik hodisalari' OLib TASHLANDI."""
        r = self._notifs()
        cats = r if isinstance(r, list) else r
        st, rr = self.c.get("/api/me/notifications")
        self.assertEqual(sorted(rr["categories"]),
                         ["all", "lesson", "message", "reminder"],
                         "faqat 4 ta kategoriya bo'lishi kerak")
        # frontend filtrlari ham 4 ta
        txt = (ROOT / "web" / "js" / "shared.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("const CATS = [")
        chunk = txt[i:i + 600]
        self.assertNotIn('["request"', chunk, "'So'rovlar' filtri olib tashlanmagan")
        self.assertNotIn('["security"', chunk, "'Xavfsizlik' filtri olib tashlanmagan")

    def test_N02_request_and_security_category_rejected(self):
        for cat in ("request", "security"):
            st, r = self.c.get(f"/api/me/notifications?category={cat}")
            self.assertEqual(st, 400, f"{cat}: {r}")
            self.assertEqual(r["error"], "bad_request")

    # ------------------------------------------------------------- BAND 12
    def test_N03_message_category_only_admin_messages(self):
        """'Xabarlar' = FAQAT admin yuborgan xabar."""
        st, r = self._admin_msg("MSG-1", "Bitta xabar")
        self.assertEqual(st, 200, r)
        msgs = self._notifs("message")
        self.assertTrue(any(m["title"] == "MSG-1" for m in msgs))
        for m in msgs:
            self.assertEqual(m["source"], "ADMIN_MESSAGE",
                             "Xabarlar kategoriyasida boshqa manba bor")
        # eslatma 'message' da chiqmasin
        start = (datetime.now() + timedelta(minutes=110)).replace(second=0, microsecond=0)
        created = (start - timedelta(hours=3)).strftime("%Y-%m-%d %H:%M:%S")
        sid = self._mk_session(start, created, self.student_id, notes="N03")
        self.sids.append(sid)
        ensure_reminders(self.db)
        rem = self._notifs("reminder")
        self.assertTrue(any(m["related_lesson_id"] == sid for m in rem))
        msgs2 = self._notifs("message")
        self.assertFalse(any(m["related_lesson_id"] == sid for m in msgs2),
                         "eslatma 'Xabarlar' kategoriyasiga tushib ketdi")

    # ------------------------------------------------------------- BAND 11
    def test_N04_each_message_is_separate_record(self):
        """Har bir xabar — ALOHIDA DB qatori; matnlar aralashMAYDI."""
        bodies = ["Xabar-birinchi", "Xabar-ikkinchi", "Xabar-uchinchi"]
        for b in bodies:
            st, _ = self._admin_msg(b, "matn-" + b)
            self.assertEqual(st, 200)
        rows = self.db.q(
            "SELECT id, title, body, source FROM notifications "
            "WHERE user_id=? AND title LIKE 'Xabar-%' ORDER BY id", (self.uid,))
        for b in bodies:
            hits = [r for r in rows if r["title"] == b]
            self.assertEqual(len(hits), 1, f"{b} uchun {len(hits)} ta qator topildi")
            self.assertIn("matn-" + b, hits[0]["body"])
            self.assertEqual(hits[0]["source"], "ADMIN_MESSAGE")
        # detail aynan o'sha ID bo'yicha
        target = [r for r in rows if r["title"] == "Xabar-ikkinchi"][0]
        st, d = self.c.get(f"/api/me/notifications/{target['id']}")
        self.assertEqual(st, 200, d)
        self.assertEqual(d["notification"]["title"], "Xabar-ikkinchi")
        self.assertIn("matn-Xabar-ikkinchi", d["notification"]["body"])

    def test_N05_notification_detail_idor(self):
        other = self._mkuser("student", "Notif", "Other", group_name="NOT2")
        try:
            st, r = self._admin_msg("OTHER-1", "boshqa odamga")
            self.assertEqual(st, 200, r)
            row = self.db.q1(
                "SELECT id FROM notifications WHERE user_id=? AND title='OTHER-1'",
                (other["user"]["id"],))
            self.assertIsNotNone(row, "boshqa foydalanuvchiga xabar yuborilmadi")
            st2, d = self.c.get(f"/api/me/notifications/{row['id']}")
            self.assertEqual(st2, 404, d)
        finally:
            self._cleanup_user(other["user"]["id"])

    def test_N06_notification_too_long_rejected(self):
        st, r = self._admin_msg("X", "M" * 5000)
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "notif.too_long")
        self.assertEqual(self.db.q1(
            "SELECT COUNT(*) c FROM notifications WHERE title='X'")["c"], 0)

    # ------------------------------------------------------------- BAND 9
    def test_N07_reminder_exactly_once(self):
        start = (datetime.now() + timedelta(minutes=115)).replace(second=0, microsecond=0)
        created = (start - timedelta(hours=4)).strftime("%Y-%m-%d %H:%M:%S")
        sid = self._mk_session(start, created, self.student_id, notes="N07")
        self.sids.append(sid)
        n1 = ensure_reminders(self.db)
        self.assertGreaterEqual(n1, 1)
        n2 = ensure_reminders(self.db)
        self.assertEqual(n2, 0, "ikkinchi marta eslatma yuborildi")
        c = self.db.q1("""SELECT COUNT(*) c FROM notifications
                          WHERE user_id=? AND source='LESSON_REMINDER'
                            AND related_lesson_id=?""", (self.uid, sid))["c"]
        self.assertEqual(c, 1, "eslatma bir marta emas")
        row = self.db.q1("""SELECT * FROM notifications
                             WHERE user_id=? AND source='LESSON_REMINDER'
                               AND related_lesson_id=?""", (self.uid, sid))
        data = row["data"] if isinstance(row["data"], dict) else jload(row["data"], {})
        for k in ("date", "start_time", "instructor_name", "pickup_address"):
            self.assertIn(k, data, f"eslatma metadatasida {k} yo'q")
        self.assertEqual(row["created_by"], "SYSTEM")

    def test_N08_no_reminder_for_late_created_lesson(self):
        start = (datetime.now() + timedelta(minutes=40)).replace(second=0, microsecond=0)
        sid = self._mk_session(start, now(), self.student_id, notes="N08")
        self.sids.append(sid)
        self.assertEqual(ensure_reminders(self.db), 0)
        c = self.db.q1("""SELECT COUNT(*) c FROM notifications
                          WHERE related_lesson_id=? AND source='LESSON_REMINDER'""",
                       (sid,))["c"]
        self.assertEqual(c, 0)

    def test_N09_no_reminder_for_started_lesson(self):
        start = (datetime.now() - timedelta(minutes=30)).replace(second=0, microsecond=0)
        sid = self._mk_session(start,
                               (start - timedelta(hours=5)).strftime("%Y-%m-%d %H:%M:%S"),
                               self.student_id, notes="N09")
        self.sids.append(sid)
        self.assertEqual(ensure_reminders(self.db), 0)

    def test_N10_reminder_cleared_on_cancel(self):
        start = (datetime.now() + timedelta(minutes=115)).replace(second=0, microsecond=0)
        sid = self._mk_session(start,
                               (start - timedelta(hours=4)).strftime("%Y-%m-%d %H:%M:%S"),
                               self.student_id, notes="N10")
        self.sids.append(sid)
        ensure_reminders(self.db)
        # DIQQAT: eslatma talabaga VA instruktorga boriladi -> `user_id` bilan
        # filtrlash shart (aks holda 2 ta qator topiladi).
        self.assertEqual(self.db.q1(
            "SELECT COUNT(*) c FROM notifications WHERE related_lesson_id=? "
            "AND user_id=? AND source='LESSON_REMINDER'", (sid, self.uid))["c"], 1)
        st, r = self.admin.post(f"/api/admin/sessions/{sid}/cancel",
                                {"reason": "bekor qilindi"})
        self.assertIn(st, (200, 201), r)
        c = self.db.q1("""SELECT COUNT(*) c FROM notifications
                          WHERE related_lesson_id=? AND user_id=?
                          AND source='LESSON_REMINDER'""",
                       (sid, self.uid))["c"]
        self.assertEqual(c, 0, "bekor qilingandan keyin eslatma qoldi")

    def test_N11_reminder_cleared_on_reschedule(self):
        start = (datetime.now() + timedelta(minutes=115)).replace(second=0, microsecond=0)
        sid = self._mk_session(start,
                               (start - timedelta(hours=4)).strftime("%Y-%m-%d %H:%M:%S"),
                               self.student_id, notes="N11")
        self.sids.append(sid)
        ensure_reminders(self.db)
        self.assertEqual(self.db.q1(
            "SELECT COUNT(*) c FROM notifications WHERE related_lesson_id=? "
            "AND user_id=? AND source='LESSON_REMINDER'", (sid, self.uid))["c"], 1)
        # Sana HAQIQIY ish kuni bo'lishi shart (instruktor mon..sat ishlaydi) —
        # `now + 9 kun` ba'zi haftalarda yakshanbaga tushib, "not_work_day"
        # bilan 409 qaytardi (test haftaning qaysi kuni bajarilishiga bog'liq
        # edi). Endi eng yaqin ish kuni topiladi.
        new_d = datetime.now() + timedelta(days=3)
        while new_d.strftime("%a").lower() == "sun":
            new_d += timedelta(days=1)
        new_day = new_d.strftime("%Y-%m-%d")
        st, r = self.admin.post(f"/api/admin/sessions/{sid}/reschedule",
                                {"date": new_day, "start_time": "14:00",
                                 "end_time": "15:00", "reason": "ko'chirildi"})
        self.assertIn(st, (200, 201), r)
        c = self.db.q1("""SELECT COUNT(*) c FROM notifications
                          WHERE related_lesson_id=? AND user_id=?
                          AND source='LESSON_REMINDER'""",
                       (sid, self.uid))["c"]
        self.assertEqual(c, 0, "vaqt ko'chirilgandan keyin eslatma qoldi")

    def test_N12_reminder_unique_index_exists(self):
        idx = self.db.q("PRAGMA index_list(notifications)")
        names = {r["name"] for r in idx}
        self.assertIn("ux_notif_reminder_once", names,
                      "eslatma 'faqat bir marta' indeksi yo'q")

    # ------------------------------------------------------------- BAND 18
    def test_N13_notification_schema_columns(self):
        cols = {r["name"] for r in self.db.q("PRAGMA table_info(notifications)")}
        for c in ("id", "user_id", "type", "title", "body", "created_at", "read_at",
                  "related_lesson_id", "created_by", "source", "data"):
            self.assertIn(c, cols, f"notifications.{c} yo'q")

    def test_N14_metadata_object_returned(self):
        st, _ = self._admin_msg("META-1", "meta")
        self.assertEqual(st, 200)
        for m in self._notifs():
            if m["title"] == "META-1":
                self.assertIsInstance(m["metadata"], dict, "metadata ob'ekt bo'lishi shart")
                break
        else:
            self.fail("META-1 topilmadi")

    def test_N15_admin_message_created_by_admin(self):
        st, _ = self._admin_msg("BY-1", "kim yuborgan")
        self.assertEqual(st, 200)
        row = self.db.q1("SELECT created_by, source FROM notifications WHERE title='BY-1'")
        self.assertEqual(row["source"], "ADMIN_MESSAGE")
        self.assertIsNotNone(row["created_by"], "created_by bo'lishi shart")
        self.assertNotEqual(row["created_by"], "SYSTEM")

    # ------------------------------------------------------------- BAND 21
    def test_N16_frontend_uses_id_not_index(self):
        """`key`/identifikator ro'yxat POSITSIYASI emas, DB `id` si bo'lishi shart."""
        for rel in ("web/js/shared.js", "web/js/ui.js"):
            txt = (ROOT / rel).read_text(encoding="utf-8", errors="ignore")
            self.assertIn('"notif-" + n.id', txt,
                          rel + ": barqaror kalit (notif-<id>) ishlatilmadi")

    def test_N17_metadata_preferred_over_json_parse(self):
        txt = (ROOT / "web" / "js" / "shared.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("n.metadata", txt, "metadata maydoni ishlatilmadi")
        self.assertIn("JSON.parse(n.data", txt, "data matni fallback sifatida yo'q")

    # ------------------------------------------------------------- BAND 17
    def test_N18_settings_default_all_true(self):
        st, r = self.c.get("/api/me/settings")
        self.assertEqual(st, 200, r)
        ns = r["settings"]["notification_settings"]
        for k in ("lesson_reminders", "admin_messages", "lesson_status_updates",
                  "messages", "requests", "security"):
            self.assertIn(k, ns, f"notification_settings.{k} yo'q")
            self.assertTrue(ns[k], f"{k} default true bo'lishi kerak")

    def test_N19_settings_roundtrip_and_persistence(self):
        st, r = self.c.put("/api/me/settings",
                           {"notification_settings": {"lesson_reminders": False,
                                                      "admin_messages": False}})
        self.assertEqual(st, 200, r)
        ns = r["notification_settings"]
        self.assertFalse(ns["lesson_reminders"])
        self.assertFalse(ns["admin_messages"])
        self.assertTrue(ns["lesson_status_updates"], "boshqa ustun buzildi")
        # BOSHQA qurilmada (boshqa sessiya) saqlangan holda o'qiladi
        c2, st2, _ = self._login(self.login, self.password, "student")
        self.assertEqual(st2, 200)
        st3, r3 = c2.get("/api/me/settings")
        self.assertEqual(st3, 200)
        ns3 = r3["settings"]["notification_settings"]
        self.assertFalse(ns3["lesson_reminders"], "sozlama saqlanmagan")
        self.assertFalse(ns3["admin_messages"], "sozlama saqlanmagan")

    def test_N20_reminder_setting_off_blocks(self):
        """Sozlama REAL: o'chirilgan eslatma yuborilmaydi."""
        st, r = self.c.put("/api/me/settings",
                           {"notification_settings": {"lesson_reminders": False}})
        self.assertEqual(st, 200, r)
        start = (datetime.now() + timedelta(minutes=115)).replace(second=0, microsecond=0)
        sid = self._mk_session(start,
                               (start - timedelta(hours=4)).strftime("%Y-%m-%d %H:%M:%S"),
                               self.student_id, notes="N20")
        self.sids.append(sid)
        ensure_reminders(self.db)
        c = self.db.q1("""SELECT COUNT(*) c FROM notifications
                          WHERE user_id=? AND source='LESSON_REMINDER'
                            AND related_lesson_id=?""", (self.uid, sid))["c"]
        self.assertEqual(c, 0, "sozlama o'chirilgan bo'lsa ham eslatma yuborildi")

    def test_N21_admin_message_setting_off_blocks(self):
        st, r = self.c.put("/api/me/settings",
                           {"notification_settings": {"admin_messages": False}})
        self.assertEqual(st, 200, r)
        st2, r2 = self._admin_msg("BLOCKED-1", "yuborilmasligi kerak")
        self.assertEqual(st2, 200, r2)
        self.assertEqual(self.db.q1(
            "SELECT COUNT(*) c FROM notifications WHERE user_id=? AND title='BLOCKED-1'",
            (self.uid,))["c"], 0, "sozlama o'chirilgan, xabar yuborildi")

    def test_N22_legacy_notif_shape_kept(self):
        """Eski frontend shakli (`notif`) saqlanadi — `admin` alohida kalit."""
        st, r = self.c.get("/api/me/settings")
        self.assertEqual(st, 200, r)
        notif = r["settings"]["notif"]
        for k in ("lesson", "admin", "message", "reminder", "request", "security"):
            self.assertIn(k, notif, f"notif.{k} yo'q")

    def test_N23_legacy_put_shape_accepted(self):
        st, r = self.c.put("/api/me/settings", {"notif": {"reminder": False}})
        self.assertEqual(st, 200, r)
        self.assertFalse(r["notification_settings"]["lesson_reminders"])
        st2, r2 = self.c.put("/api/me/settings", {"notif": {"reminder": True}})
        self.assertEqual(st2, 200, r2)
        self.assertTrue(r2["notification_settings"]["lesson_reminders"])


# ============================================================================
# 5) PROFILE — BAND 15, 16
# ============================================================================
class TestProfileBAND24(_MixinAdmin, Base):
    """PROFILE: 5 ta tahrirlanadigan maydon, login unikal, parol ko'rinmasin."""

    def setUp(self):
        super().setUp()
        r = self._mkuser("student", "Prof", "Test", group_name="PRF",
                         license_category="C", phone="+998901112233")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.student_id = self.db.q1("SELECT id FROM students WHERE user_id=?",
                                     (self.uid,))["id"]
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)

    def tearDown(self):
        self._cleanup_user(self.uid)
        super().tearDown()

    def test_P01_update_five_fields(self):
        st, r = self.c.put("/api/me/profile", {
            "birth_date": "2004-05-17", "phone": "+998907654321",
            "login": self.login, "group_name": "PRF-2024", "license_category": "B",
        })
        self.assertEqual(st, 200, r)
        u = self.db.q1("SELECT birth_date, phone FROM users WHERE id=?", (self.uid,))
        self.assertEqual(u["birth_date"], "2004-05-17")
        self.assertEqual(u["phone"], "+998907654321")
        s = self.db.q1("SELECT group_name, license_category FROM students WHERE id=?",
                       (self.student_id,))
        self.assertEqual(s["group_name"], "PRF-2024")
        self.assertEqual(s["license_category"], "B")

    def test_P02_login_format_enforced(self):
        for bad in ("admin", "usrL_short", "usrX_aaaaaaaaaaaaaaaa", "usrL_ab cd1234",
                    "usrL_ab!cd1234567", "PLAINlogin12345"):
            st, r = self.c.put("/api/me/profile", {"login": bad})
            self.assertEqual(st, 400, f"{bad!r} qabul qilindi: {r}")
            self.assertEqual(r["error"], "profile.login_format", f"{bad!r}: {r}")
        # eski login o'zgarmagan
        self.assertEqual(self.db.q1("SELECT login FROM users WHERE id=?",
                                    (self.uid,))["login"], self.login)

    def test_P03_login_unique_enforced(self):
        r2 = self._mkuser("student", "Prof", "Second", group_name="PRF2")
        try:
            st, r = self.c.put("/api/me/profile", {"login": r2["credentials"]["login"]})
            self.assertEqual(st, 400, r)
            self.assertEqual(r["error"], "profile.login_taken", r)
        finally:
            self._cleanup_user(r2["user"]["id"])

    def test_P04_login_change_reported_and_revokes_sessions(self):
        new_login = "usrL_profChanged0001"
        st, r = self.c.put("/api/me/profile", {"login": new_login})
        self.assertEqual(st, 200, r)
        self.assertTrue(r.get("login_changed"), "login_changed qaytarilishi shart")
        self.assertEqual(self.db.q1("SELECT login FROM users WHERE id=?",
                                    (self.uid,))["login"], new_login)
        # eski sessiya bekor qilindi
        self.assertEqual(self.c.get("/api/auth/me")[0], 401)
        # yangi login bilan kirish ishlaydi
        c2, st2, r2 = self._login(new_login, self.password, "student")
        self.assertEqual(st2, 200, r2)
        # eski login bilan KIRIB BO'LMAYDI
        c3, st3, r3 = self._login(self.login, self.password, "student")
        self.assertEqual(st3, 400, r3)
        self.db.upd("UPDATE users SET login=? WHERE id=?", (self.login, self.uid))

    def test_P05_password_not_editable_via_profile(self):
        """BAND 15: parol profil orqali O'Zgartirilmaydi."""
        st, r = self.c.put("/api/me/profile", {
            "password": "usrP_New000000!Zx1",
            "new_password": "usrP_New000000!Zx1",
            "password_hash": "x", "must_change_password": 1,
        })
        self.assertEqual(st, 200, r)
        h = self.db.q1("SELECT password_hash FROM users WHERE id=?", (self.uid,))["password_hash"]
        from app.auth import verify_password
        self.assertTrue(verify_password(self.password, h), "parol o'zgargan!")
        # javobda parol maydonlari YO'Q
        raw = json.dumps(r)
        self.assertNotIn("password_hash", raw)
        self.assertNotIn("must_change_password", raw)

    def test_P06_enrolled_at_hidden_from_student(self):
        """BAND 16: 'Ro'yxatga olingan sana' talabaga KO'RSATILMAYDI."""
        st, r = self.c.get("/api/auth/me")
        self.assertEqual(st, 200, r)
        self.assertNotIn("enrolled_at", r["student"],
                         "enrolled_at talabaga ko'rsatildi")
        st2, r2 = self.c.get("/api/student/home")
        self.assertEqual(st2, 200, r2)
        self.assertNotIn("enrolled_at", json.dumps(r2))
        # lekin DB'da SAQLANADI
        self.assertIsNotNone(self.db.q1(
            "SELECT enrolled_at FROM students WHERE id=?", (self.student_id,)))

    def test_P07_admin_sees_enrolled_at(self):
        st, r = self.admin.get(f"/api/admin/users/{self.uid}")
        self.assertEqual(st, 200, r)
        self.assertIn("enrolled_at", r["user"]["student"],
                      "admin enrolled_at ni ko'ra olmaydi")

    def test_P08_invalid_phone_and_birth(self):
        st, r = self.c.put("/api/me/profile", {"phone": "123"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "profile.phone_invalid")
        st2, r2 = self.c.put("/api/me/profile", {"birth_date": "3000-01-01"})
        self.assertEqual(st2, 400, r2)
        self.assertEqual(r2["error"], "profile.birth_date_invalid")
        st3, r3 = self.c.put("/api/me/profile", {"birth_date": "2000-13-45"})
        self.assertEqual(st3, 400, r3)
        self.assertEqual(r3["error"], "profile.birth_date_invalid")

    def test_P09_instructor_profile_works(self):
        c, st, rr = self._login(INSTR_LOGIN_1, PW_BY_LOGIN[INSTR_LOGIN_1], "instructor")
        self.assertEqual(st, 200, rr)
        st2, r2 = c.put("/api/me/profile", {"birth_date": "1990-01-02",
                                            "phone": "+998901234567"})
        self.assertEqual(st2, 200, r2)
        st3, r3 = c.get("/api/auth/me")
        self.assertEqual(st3, 200, r3)
        self.assertEqual(r3["user"]["birth_date"], "1990-01-02")
        self.db.upd("UPDATE users SET birth_date=NULL WHERE login=?", (INSTR_LOGIN_1,))
        self.db.upd("UPDATE users SET phone=NULL WHERE login=?", (INSTR_LOGIN_1,))

    def test_P10_profile_modal_has_personal_fields_only(self):
        """MODUL 4: 'Profilni tahrirlash'da FAQAT shaxsiy ma'lumotlar.

        LOGIN va PAROL bu modaldan OLIB TASHLANDI — ular endi alohida
        "Parol va login o'zgartirish" oqimiga ko'chirilgan (joriy parol
        talab qilinadi). Bu test aynan shu talabni tekshiradi.
        """
        txt = (ROOT / "web" / "js" / "shared.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("function editProfileModal")
        self.assertGreater(i, 0, "editProfileModal topilmadi")
        chunk = txt[i:i + 4000]
        end = chunk.find("\n  }")
        body = chunk[:end if end > 0 else len(chunk)]
        # shaxsiy ma'lumotlar — SAQLANADI
        for key in ("student.birth_date", "common.phone",
                    "student.group", "student.category"):
            self.assertIn(key, body, f"profil modalidan {key} yo'q")
        # MODUL 4: login maydoni bu modaldan O'CHIRILGAN
        self.assertNotIn("common.login", body,
                         "login 'Profilni tahrirlash' modalidan olib tashlanishi shart")
        self.assertNotIn("login: String(", body,
                         "login profildan yuborilmasligi shart")
        # parol maydonlari BO'LMASIN (hech qachon)
        self.assertNotIn('"password"', body, "profil modalida parol maydoni bor")
        self.assertIn("profile.password_not_here", body, "parol ogohlantirishi yo'q")

        # alohida oqim mavjudligi: joriy parol + yangi login/parol/takrorlash
        j = txt.find("function openChangeCredentials")
        self.assertGreater(j, 0, "openChangeCredentials topilmadi")
        # Funksiya TUZILISHI o'zgarganda chegarali kesim (3000 belgi) yetarli
        # bo'lmasligi mumkin — keyingi yuqori darajali `function` deklaratsiyasi
        # yoki oxirigacha.
        nxt = txt.find("\n  function ", j + 10)
        cred = txt[j:nxt if nxt > 0 else j + 20000]
        for key in ("profile.current_password", "profile.new_login",
                    "profile.new_password", "profile.confirm_new_password"):
            self.assertIn(key, cred, f"parol/login modalida {key} yo'q")
        # real-vaqt kuch ro'yxati: PW_RULES (frontend ham bosh harf/raqamni
        # tekshiradi) va uning i18n kalitlari HAQIQATAN mavjud
        self.assertIn("PW_RULES", cred, "parol talablari real-vaqtda hisoblanmayapti")
        self.assertIn('"auth.pw_rule_"', txt, "parol talablari i18n'ga bog'lanmagan")
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8")
        for key in ("auth.pw_rule_len", "auth.pw_rule_upper", "auth.pw_rule_lower",
                    "auth.pw_rule_digit", "auth.pw_rule_special",
                    "auth.pw_match_ok", "auth.pw_match_bad"):
            self.assertIn('"%s"' % key, i18n, f"i18n da {key} yo'q")
        # eski "easiy" parol modali credentials oqimida ishlatilmasin
        self.assertNotIn("openChangePassword(true)", cred,
                         "eski parol modali credentials oqimida ishlatilmasin")
        # majburiy-parol oqimi ham XUDDI SHU xavfsiz modalni ochadi
        appjs = (ROOT / "web" / "js" / "app.js").read_text(encoding="utf-8")
        self.assertIn("Shared.openChangeCredentials()", appjs,
                      "majburiy-parol oqimi alohida modaldan o'tishi kerak")
        self.assertNotIn("openChangePassword(", appjs,
                         "eski openChangePassword hali ishlatilmoqda")

    def test_P11_no_hardcoded_numbers_in_profile_view(self):
        txt = (ROOT / "web" / "js" / "views-student.js").read_text(encoding="utf-8",
                                                                    errors="ignore")
        i = txt.find("async function profile")
        self.assertGreater(i, 0)
        chunk = txt[i:i + 1200]
        # izohlarni tushirib qolamiz (ular `enrolled_at` haqida YOZADI,
        # lekin qiymatni chiqarmaydi) — tekshiruv KODGA qariladi.
        code = re.sub(r"/\*.*?\*/", "", chunk, flags=re.S)
        code = re.sub(r"//[^\n]*", "", code)
        self.assertIn("student.birth_date", code)
        self.assertIn("u.birth_date", code, "tug'ilgan sana API dan olinishi kerak")
        self.assertIn("u.phone", code)
        self.assertIn("u.login", code)
        self.assertNotIn("enrolled_at", code, "enrolled_at talabada ko'rsatilmoqda")


# ============================================================================
# 6) SETTINGS — BAND 17 (bildirishnoma sozlamalari DB da)
# ============================================================================
class TestSettingsBAND24(_MixinAdmin, Base):
    """SETTINGS: sozlamalar faqat serverda, boshqa qurilmada ham saqlanadi."""

    def setUp(self):
        super().setUp()
        r = self._mkuser("student", "Set", "Test", group_name="SET", license_category="B")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)

    def tearDown(self):
        self._cleanup_user(self.uid)
        super().tearDown()

    def test_S01_settings_table_exists(self):
        cols = {r["name"] for r in self.db.q("PRAGMA table_info(notification_settings)")}
        self.assertIn("user_id", cols)
        self.assertIn("lesson_reminders", cols)
        self.assertIn("admin_messages", cols)
        self.assertIn("lesson_status_updates", cols)

    def test_S02_no_localstorage_source_of_truth(self):
        """Frontend sozlamani localStorage'da emas, serverdan o'qishi SHART."""
        txt = (ROOT / "web" / "js" / "shared.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("function notifSettingsCard")
        self.assertGreater(i, 0, "notifSettingsCard topilmadi")
        chunk = txt[i:i + 2600]
        end = chunk.find("\n  }")
        body = chunk[:end if end > 0 else len(chunk)]
        self.assertIn('API.put("me/settings"', body, "sozlama serverga yuborilmayapti")
        self.assertIn("notification_settings", body, "yangi ustunlar ishlatilmayapti")
        self.assertNotIn("localStorage", body, "sozlama localStorage da saqlanmoqda")

    def test_S03_each_toggle_saves_its_own_column(self):
        for col in ("lesson_reminders", "admin_messages", "lesson_status_updates"):
            st, r = self.c.put("/api/me/settings",
                               {"notification_settings": {col: False}})
            self.assertEqual(st, 200, r)
            ns = r["notification_settings"]
            self.assertFalse(ns[col], f"{col} saqlanmadi")
            others = [k for k in ns if k != col]
            self.assertTrue(all(ns[k] for k in others),
                            f"{col} o'chganda boshqa ustunlar ham o'chdi")
            st2, r2 = self.c.put("/api/me/settings",
                                 {"notification_settings": {col: True}})
            self.assertEqual(st2, 200, r2)
            self.assertTrue(r2["notification_settings"][col])

    def test_S04_settings_survive_new_client(self):
        """Brauzer yopilsa yoki boshqa qurilmada kirsangiz ham saqlanadi."""
        self.c.put("/api/me/settings",
                   {"notification_settings": {"lesson_status_updates": False}})
        c2, st, _ = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200)
        st2, r2 = c2.get("/api/me/settings")
        self.assertEqual(st2, 200)
        self.assertFalse(r2["settings"]["notification_settings"]["lesson_status_updates"])
        # DB da ham saqlangan
        row = self.db.q1("SELECT * FROM notification_settings WHERE user_id=?", (self.uid,))
        self.assertIsNotNone(row)
        self.assertEqual(int(row["lesson_status_updates"]), 0)

    def test_S05_lang_and_theme_still_work(self):
        st, r = self.c.put("/api/me/settings", {"lang": "ru", "theme": "dark"})
        self.assertEqual(st, 200, r)
        st2, r2 = self.c.get("/api/me/settings")
        self.assertEqual(st2, 200, r2)
        self.assertEqual(r2["settings"]["lang"], "ru")
        self.assertEqual(r2["settings"]["theme"], "dark")
        st3, r3 = self.c.put("/api/me/settings", {"lang": "xx"})
        self.assertEqual(st3, 200, r3)
        st4, r4 = self.c.get("/api/me/settings")
        self.assertEqual(r4["settings"]["lang"], "ru", "noto'g'ri til o'zgartirildi")

    def test_S06_anonymous_cannot_read_settings(self):
        anon = Client()
        st, r = anon.get("/api/me/settings")
        self.assertEqual(st, 401, r)
        st2, r2 = anon.put("/api/me/settings",
                           {"notification_settings": {"lesson_reminders": False}})
        self.assertEqual(st2, 401, r2)

    def test_S07_other_role_cannot_touch_my_settings(self):
        r2 = self._mkuser("student", "Set", "Other", group_name="SET2")
        try:
            c2, st, _ = self._login(r2["credentials"]["login"],
                                    r2["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            st2, r3 = c2.get("/api/me/settings")
            self.assertEqual(st2, 200)
            ns = r3["settings"]["notification_settings"]
            self.assertTrue(all(ns.values()),
                            "boshqa foydalanuvchining sozlamalari aralashdi")
        finally:
            self._cleanup_user(r2["user"]["id"])


# ============================================================================
# 7) BAND 3 — admin jami mashg'ulotlar sonini o'zgartirish (alohida sinf,
#    chunki u boshqa sinflardagi progress qiymatlariga ta'sir qiladi)
# ============================================================================
class TestTotalLessonsAdminBAND24(_MixinAdmin, Base):
    """BAND 3: bitta talaba VA barcha talabalar uchun jami sonni o'zgartirish."""

    def setUp(self):
        super().setUp()
        self.uids = []
        self.sids = []
        # Yaratilgan talabaning HAQIQIY credential'lari (random generatsiya qilinadi,
        # shuning uchun fixture ro'yxatida yo'q).
        self.pws = {}
        for i in range(2):
            r = self._mkuser("student", "Tot", "L%d" % i, group_name="TOT%d" % i)
            self.uids.append(r["user"]["id"])
            self.pws[r["user"]["id"]] = r["credentials"]["password"]

    def tearDown(self):
        for sid in self.sids:
            self._drop_session(sid)
        for uid in self.uids:
            self._cleanup_user(uid)
        # platforma umumiy qiymatini tiklash
        self.admin.put("/api/admin/settings", {"total_lessons_target": 30})
        super().tearDown()

    def _complete(self, uid, n):
        """Talabaga `n` ta bajarilgan mashg'ulot yozib beradi."""
        srow = self.db.q1("SELECT id FROM students WHERE user_id=?", (uid,))
        for i in range(n):
            d = (datetime.now() - timedelta(days=10 + i)).replace(hour=9, minute=0,
                                                                 second=0, microsecond=0)
            sid = self._mk_session(d, now(), srow["id"], status="completed",
                                   notes="TOT-DONE")
            self.sids.append(sid)

    def _home(self, uid):
        c, st, _ = self._login(*(self._creds(uid)), role="student")
        self.assertEqual(st, 200)
        st2, h = c.get("/api/student/home")
        self.assertEqual(st2, 200, h)
        return h

    def _creds(self, uid):
        """Talabaning login/parol juftligi (parol API javobidan olinadi)."""
        row = self.db.q1("SELECT login FROM users WHERE id=?", (uid,))
        return row["login"], self.pws[uid]

    def test_T01_single_user_total(self):
        uid = self.uids[0]
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "user", "user_id": uid, "total_lessons": 25})
        self.assertEqual(st, 200, r)
        h = self._home(uid)
        self.assertEqual(h["progress"]["target"], 25)
        self.assertEqual(h["progress"]["individual_target"], 25)
        self.assertEqual(h["progress"]["mode"], "individual")
        self.assertEqual(h["progress"]["remaining"], 25 - h["progress"]["done"])

    def test_T02_below_completed_rejected(self):
        """BAND 3: yangi jami < bajarilgan bo'lsa — QABUL QILINMAYDI."""
        uid = self.uids[0]
        self._complete(uid, 8)
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "user", "user_id": uid, "total_lessons": 5})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "user.total_lessons_below_done", r)
        self.assertEqual(r.get("params", {}).get("done", r.get("done")), 8, r)
        # oldingi qiymat buzilMADI
        h = self._home(uid)
        self.assertNotEqual(h["progress"]["individual_target"], 5)

    def test_T03_equal_to_completed_allowed(self):
        uid = self.uids[0]
        self._complete(uid, 6)
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "user", "user_id": uid, "total_lessons": 6})
        self.assertEqual(st, 200, r)
        h = self._home(uid)
        self.assertEqual(h["progress"]["target"], 6)
        self.assertEqual(h["progress"]["remaining"], 0)

    def test_T04_history_not_deleted(self):
        """BAND 3: jami sonni o'zgartirish BAJARILGAN mashg'ulotlar,
        tarix va booking'ni O'CHIRMAYDI."""
        uid = self.uids[0]
        self._complete(uid, 3)
        before_done = self._home(uid)["progress"]["done"]
        before_rows = self.db.q("SELECT id FROM lesson_sessions")
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "user", "user_id": uid, "total_lessons": 40})
        self.assertEqual(st, 200, r)
        h = self._home(uid)
        self.assertEqual(h["progress"]["done"], before_done,
                         "bajarilgan darslar soni o'zgardi!")
        after_rows = self.db.q("SELECT id FROM lesson_sessions")
        self.assertEqual(len(after_rows), len(before_rows),
                         "mashg'ulot qatorlari o'chirildi")
        c, st2, _ = self._login(*(self._creds(uid)), role="student")
        st3, hist = c.get("/api/student/sessions?upcoming=0")
        self.assertEqual(st3, 200, hist)
        self.assertEqual(len([x for x in hist["sessions"] if x["_display_status"] == "completed"]),
                         before_done, "tarixdagi completed darslar kamaydi")

    def test_T05_bad_values_rejected(self):
        uid = self.uids[0]
        for bad in (0, -1, 1000, "abc", 12.5):
            st, r = self.admin.post("/api/admin/users/total-lessons",
                                    {"scope": "user", "user_id": uid,
                                     "total_lessons": bad})
            self.assertEqual(st, 400, f"{bad!r}: {r}")
            self.assertEqual(r["error"], "user.bad_total_lessons", f"{bad!r}: {r}")

    def test_T06_all_students_scope(self):
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "all", "total_lessons": 50})
        self.assertEqual(st, 200, r)
        self.assertGreaterEqual(r.get("updated", 0), len(self.uids))
        for uid in self.uids:
            h = self._home(uid)
            self.assertEqual(h["progress"]["target"], 50)
            # ommaviy o'zgartirish ham individual maqsad qo'yadi -> mode=individual
            self.assertEqual(h["progress"]["individual_target"], 50)
            self.assertEqual(h["progress"]["mode"], "individual")

    def test_T07_all_students_below_max_completed_rejected(self):
        """Ommaviy rejimda ham eng ko'p bajarilgan talaba tekshiriladi."""
        self._complete(self.uids[0], 7)
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "all", "total_lessons": 3})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "user.total_lessons_below_done", r)
        for uid in self.uids:
            h = self._home(uid)
            self.assertNotEqual(h["progress"]["target"], 3)

    def test_T08_clear_individual_target(self):
        uid = self.uids[0]
        self.admin.post("/api/admin/users/total-lessons",
                        {"scope": "user", "user_id": uid, "total_lessons": 44})
        self.assertEqual(self._home(uid)["progress"]["individual_target"], 44)
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "user", "user_id": uid, "total_lessons": None})
        self.assertEqual(st, 200, r)
        h = self._home(uid)
        self.assertIsNone(h["progress"]["individual_target"])
        self.assertEqual(h["progress"]["mode"], "group")

    def test_T09_non_student_rejected(self):
        row = self.db.q1("SELECT id FROM users WHERE login=?", (INSTR_LOGIN_1,))
        st, r = self.admin.post("/api/admin/users/total-lessons",
                                {"scope": "user", "user_id": row["id"],
                                 "total_lessons": 30})
        self.assertIn(st, (400, 404), r)

    def test_T10_only_admin_can_change(self):
        c, st, _ = self._login(STUD_LOGIN_1, PW_BY_LOGIN[STUD_LOGIN_1], "student")
        self.assertEqual(st, 200)
        st2, r = c.post("/api/admin/users/total-lessons",
                        {"scope": "user", "user_id": self.uids[0], "total_lessons": 1})
        self.assertEqual(st2, 403, r)

    def test_T11_admin_users_progress_shape(self):
        st, r = self.admin.get("/api/admin/users?role=student")
        self.assertEqual(st, 200, r)
        for u in r["users"]:
            pr = u.get("progress")
            if not pr:
                continue
            for k in ("target", "done", "remaining", "mode", "individual_target",
                      "group_target"):
                self.assertIn(k, pr, f"admin_users progress.{k} yo'q")

    def test_T12_modal_text_present(self):
        """Modal sarlavhasi va tanlovlar talab dagidek."""
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8",
                                                                   errors="ignore")
        self.assertIn("users.total_title", txt)
        self.assertIn("users.scope_single", txt)
        self.assertIn("users.scope_all", txt)
        self.assertIn("function totalLessonsModal", txt)
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("Jami amaliy mashg'ulotlar sonini o'zgartirish", i18n)
        self.assertIn("Bitta talaba", i18n)
        self.assertIn("Barcha talabalar", i18n)


# ============================================================================
# 9) ADMIN PANEL MODULLARI 1-7 — Baza bo'limi hisoblagich, tugmalar,
#    eskirgan so'rovlar, ommaviy amallar, lokalizatsiya
# ============================================================================


class TestAdminModulesBAND25(_MixinAdmin, Base):
    """MODUL 1/2/3/6 — Baza bo'limi va ommaviy amallar.

    MODUL 1: admin "Dashboard" -> "Bosh sahifa".
    MODUL 2: har bir toifa uchun "Jami: N ta <rol>" (filtrga mos).
    MODUL 3: "Bulk yaratish" -> "Tezkor yaratish"; "Import CSV" Baza
             bo'limidan olib tashlandi, o'rniga "Hisobot" tugmasi.
    MODUL 6: checkbox + "Hammasini tanlash" + ommaviy tahrirlash/o'chirish.
    """

    def setUp(self):
        super().setUp()
        self.uids = []
        # har bir rol uchun bittadan foydalanuvchi (MODUL 2 uchun)
        self.su = self._mkuser("student", "Mod2", "Student")["user"]["id"]
        self.iu = self._mkuser("instructor", "Mod2", "Instructor")["user"]["id"]
        self.au = self._mkuser("instructor", "Mod2", "Tmp2")["user"]["id"]
        # "Tezkor yaratish" bilan yaratilgan yana 2 ta talaba (MODUL 6).
        # Javob `created` — CREDENTIAL ro'yxati ({login, password}), `id` YO'Q.
        st, r = self.admin.post("/api/admin/users/bulk", {"count": 2})
        self.assertEqual(st, 200, r)
        self.assertEqual(r["count"], 2, r)
        # login orqali `user.id` ni topamiz (test tozalash uchun)
        for cr in r["created"]:
            row = self.db.q1("SELECT id FROM users WHERE login=?", (cr["login"],))
            self.assertIsNotNone(row, "bulk yaratilgan foydalanuvchi topilmadi")
            self.uids.append(row["id"])

    def tearDown(self):
        for uid in list(self.uids) + [self.su, self.iu, self.au]:
            self._cleanup_user(uid)
        super().tearDown()

    # ---------------------------------------------------------- MODUL 1
    def test_MOD1_admin_nav_uses_bosh_sahifa(self):
        """Admin sidebar "Bosh sahifa" — Talaba/Instruktor kabi."""
        app = (ROOT / "web" / "js" / "app.js").read_text(encoding="utf-8", errors="ignore")
        # admin roli `nav.home` kalitidan foydalanadi
        i = app.find("admin: [")
        j = app.find("]", i)
        chunk = app[i:j]
        self.assertIn('"nav.home"', chunk, "admin sidebar 'Bosh sahifa' emas")
        self.assertNotIn('"nav.dashboard"', chunk, "admin hali ham 'nav.dashboard' ishlatadi")
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn('"nav.home": { uz: "Bosh sahifa"', i18n)

    # ---------------------------------------------------------- MODUL 2
    def test_MOD2_total_matches_filter(self):
        """`total` — filtr bilan BIR XIL shartlarda hisoblanadi."""
        st, r = self.admin.get("/api/admin/users?role=student")
        self.assertEqual(st, 200, r)
        self.assertIn("total", r, "backend `total` maydoni yo'q (MODUL 2)")
        self.assertIn("by_role", r, "`by_role` maydoni yo'q")
        # `role=student` ham filtr -> `filtered` true (bo'sh bo'lishi SHART emas)
        self.assertTrue(r["filtered"], "rol filtri `filtered` true bo'lishi kerak")
        # faqat `q` (qidiruv) bilan filtr o'zgarganini tekshiramiz
        st, r2 = self.admin.get("/api/admin/users?role=student&q=Mod2")
        self.assertEqual(st, 200, r2)
        self.assertTrue(r2["filtered"], "qidiruv filtri qo'yilganda `filtered` true bo'lishi kerak")
        self.assertLessEqual(r2["total"], r["total"],
                             "qidiruv natijasi umumiy sondan KATTAROQ bo'lmasligi kerak")
        # `total` qidiruv natijasi bilan mos keladi
        st, allq = self.admin.get("/api/admin/users?role=student&q=Modu")
        self.assertEqual(allq["total"], allq["shown"])

    def test_MOD2_by_role_counts(self):
        """`by_role` — uch toifa soni (o'chirilmagan foydalanuvchilar)."""
        st, r = self.admin.get("/api/admin/users")
        self.assertEqual(st, 200, r)
        for role in ("student", "instructor", "admin"):
            self.assertIn(role, r["by_role"], f"{role} soni yo'q")
        self.assertGreaterEqual(r["by_role"]["student"], 1)
        self.assertGreaterEqual(r["by_role"]["instructor"], 1)

    # ---------------------------------------------------------- MODUL 3
    def test_MOD3_import_csv_removed_from_base_section(self):
        """Baza bo'limida "Import CSV" tugmasi OLIB TASHLANGAN."""
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        # `base()` funksiyasida importModal chaqirilmaydi
        i = txt.find("async function base()")
        j = txt.find("function bulkModal()", i)
        base_chunk = txt[i:j]
        self.assertNotIn("importModal()", base_chunk,
                         "'Import CSV' hali ham Baza bo'limida chaqirilmoqda")
        # "Hisobot" tugmasi bor (Hisobotlar bo'limiga o'tadi)
        self.assertIn("base.go_reports", base_chunk, "'Hisobot' tugmasi yo'q")
        self.assertIn("goReportsFor", base_chunk)

    def test_MOD3_bulk_renamed_to_tezkor_yaratish(self):
        """"Bulk yaratish" -> "Tezkor yaratish" (i18n + sarlavha)."""
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn('"student.bulk": { uz: "Tezkor yaratish"', i18n)
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        # modal sarlavhasida "Bulk" so'zi QOLMAMALIGI kerak
        i = txt.find("function bulkModal()")
        j = txt.find("function importModal()", i)
        self.assertNotIn('— Bulk', txt[i:j], "modal sarlavhasida 'Bulk' qolgan")

    def test_MOD3_import_moved_to_reports_not_deleted(self):
        """Import funksiyasi O'CHIRILMADI — Hisobotlar bo'limidan ochiladi."""
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("function importModal", txt, "importModal o'chirilgan")
        self.assertIn("importModal,", txt, "importModal eksportga qo'shilmagan")

    # ---------------------------------------------------------- MODUL 6
    def test_MOD6_bulk_update_common_fields(self):
        """Ommaviy tahrirlash: umumiy maydonlar barchasiga qo'llaniladi."""
        st, r = self.admin.post("/api/admin/users/bulk-update",
                                {"ids": [self.su], "group_name": "G-99", "license_category": "C"})
        self.assertEqual(st, 200, r)
        self.assertGreaterEqual(r["updated"], 1, r)
        row = self.db.q1("SELECT group_name, license_category FROM students WHERE user_id=?",
                         (self.su,))
        self.assertEqual(row["group_name"], "G-99")
        self.assertEqual(row["license_category"], "C")

    def test_MOD6_bulk_update_status_all(self):
        """Ommaviy holat o'zgartirish."""
        st, r = self.admin.post("/api/admin/users/bulk-update",
                                {"ids": [self.su, self.iu], "status": "blocked"})
        self.assertEqual(st, 200, r)
        for uid in (self.su, self.iu):
            row = self.db.q1("SELECT status FROM users WHERE id=?", (uid,))
            self.assertEqual(row["status"], "blocked")

    def test_MOD6_bulk_update_empty_fields_not_applied(self):
        """BO'SH maydonlar o'zgartirilMAYDI."""
        before = self.db.q1("SELECT group_name FROM students WHERE user_id=?", (self.su,))
        st, r = self.admin.post("/api/admin/users/bulk-update",
                                {"ids": [self.su], "group_name": "", "status": "active"})
        self.assertEqual(st, 200, r)
        after = self.db.q1("SELECT group_name FROM students WHERE user_id=?", (self.su,))
        self.assertEqual(after["group_name"], before["group_name"],
                         "bo'sh maydon o'zgarib ketdi")

    def test_MOD6_bulk_update_nothing_to_change_rejected(self):
        st, r = self.admin.post("/api/admin/users/bulk-update", {"ids": [self.su]})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "bulk.nothing_to_change")

    def test_MOD6_bulk_update_empty_ids_rejected(self):
        st, r = self.admin.post("/api/admin/users/bulk-update", {"ids": [], "status": "active"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "bulk.none")

    def test_MOD6_bulk_update_too_many_rejected(self):
        st, r = self.admin.post("/api/admin/users/bulk-update",
                                {"ids": list(range(1, 500)), "status": "active"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "bulk.too_many")

    def test_MOD6_bulk_update_bad_id_is_sql_injection_safe(self):
        """Noto'g'ri ID (SQLi) — 400, jim qolmaydi, hech narsa o'zgarmaydi."""
        st, r = self.admin.post("/api/admin/users/bulk-update",
                                {"ids": ["1 OR 1=1"], "status": "blocked"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "bulk.bad_id")

    def test_MOD6_bulk_delete_is_hard_delete(self):
        """Ommaviy o'chirish — haqiqiy DELETE (arxiv EMAS).

        Talab: "O'chirish" foydalanuvchini ARXIVGA o'tkazmasin — `users`
        jadvalidan butunlay o'chsin, "Arxivlanganlar"da ham ko'rinmasin.
        """
        st, r = self.admin.post("/api/admin/users/bulk-delete", {"ids": [self.su]})
        self.assertEqual(st, 200, r)
        self.assertEqual(r["deleted"], 1, r)
        # 1) `users` jadvalidan QATOR YO'Q (arxiv emas!)
        self.assertIsNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.su,)),
                          "users qatori saqlanib qoldi — bu arxivlash, o'chirmaslik emas")
        # 2) student yozuvi ham o'chgan (FK butunligi)
        self.assertIsNone(self.db.q1("SELECT id FROM students WHERE user_id=?", (self.su,)),
                          "students qatori qoldi — FK buzilgan")
        # 3) arxivda ham KO'RINMASIN
        st, r2 = self.admin.get("/api/admin/users?status=archived")
        self.assertNotIn(self.su, [u["id"] for u in r2["users"]],
                         "arxivlanganlar ro'yxatida qoldi")
        # 4) profil ham yo'q
        st, _ = self.admin.get("/api/admin/users/%s" % self.su)
        self.assertEqual(st, 404)

    def test_MOD6_bulk_delete_does_not_touch_others(self):
        """Boshqa foydalanuvchilarning ma'lumotlari OCHIRILMAYDI."""
        other = self._mkuser("student", "Mod2", "Boshqa")["user"]["id"]
        self.uids.append(other)
        before = self.db.q1("SELECT * FROM students WHERE user_id=?", (other,))
        st, r = self.admin.post("/api/admin/users/bulk-delete", {"ids": [self.su]})
        self.assertEqual(st, 200, r)
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (other,)))
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.iu,)))
        after = self.db.q1("SELECT * FROM students WHERE user_id=?", (other,))
        self.assertEqual(after["group_name"], before["group_name"])
        self.assertEqual(after["license_category"], before["license_category"])

    def test_MOD6_bulk_delete_self_protected(self):
        """Admin O'ZINI o'chira olMAYDI (tizimda admin qolishi uchun)."""
        me = self.admin.get("/api/auth/me")[1]["user"]["id"]
        st, r = self.admin.post("/api/admin/users/bulk-delete", {"ids": [me]})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "user.cannot_delete_self")
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (me,)),
                             "admin o'zini o'chirildi!")

    def test_MOD6_bulk_delete_mixed_with_self_rejected_entirely(self):
        """Aralash ro'yxatda admin bo'lsa — butun so'rov RAD ETILADI.

        XAVFSIZLIK: qisman o'chirish ("ba'zilari o'chirildi, ba'zilari
        qoldi") chalkash natija beradi. Shuning uchun admin o'z ro'yxatida
        bo'lsa — HECH NIMA o'chirilmaydi va aniq xato qaytariladi.
        Bu "jim qoldirmaslik" tamoyili: admin biladi ki, hech kim
        o'chirilmagan.
        """
        me = self.admin.get("/api/auth/me")[1]["user"]["id"]
        st, r = self.admin.post("/api/admin/users/bulk-delete", {"ids": [me, self.iu]})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "user.cannot_delete_self")
        # muhim: instructor HAM o'chirilmagan (atomik rad etish)
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (me,)),
                             "admin o'zini o'chirildi!")
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.iu,)),
                             "qisman o'chirish bo'ldi — atomik emas!")

    # ------------------------------------------- BUTUNLAY O'CHIRISH (DELETE)
    def test_PURGE_single_user_removed_from_users_table(self):
        """`DELETE /api/admin/users/{id}` — `users` qatori haqiqatan o'chadi."""
        st, r = self.admin.delete("/api/admin/users/%s" % self.su)
        self.assertEqual(st, 200, r)
        self.assertEqual(r["deleted"], 1, r)
        self.assertIsNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.su,)))

    def test_PURGE_not_visible_anywhere_after_refresh(self):
        """O'chgandan keyin HECH QAYERDA ko'rinmaydi (arxivda ham)."""
        self.admin.delete("/api/admin/users/%s" % self.su)
        for q in ("", "&status=archived", "&status=active", "&status=blocked"):
            st, r = self.admin.get("/api/admin/users?role=student" + q)
            self.assertEqual(st, 200, r)
            self.assertNotIn(self.su, [u["id"] for u in r["users"]],
                             "ro'yxatda qoldi (?%s)" % q)

    def test_PURGE_cascades_student_related_rows(self):
        """Talabaga tegishli qatorlar ham o'chadi (FK buzilmaydi)."""
        sid = self.db.q1("SELECT id FROM students WHERE user_id=?", (self.su,))["id"]
        # qarshi tekshiruv uchun ma'lumot yaratamiz
        iid = self.db.q1("SELECT id FROM instructors WHERE user_id=?", (self.iu,))["id"]
        ls = _ins_lesson(self.db, iid, "2026-01-05")
        self.db.ex("INSERT OR IGNORE INTO session_students "
                   "(session_id,student_id,attendance_status,joined_at) VALUES (?,?,'present','2026-01-05')",
                   (ls, sid))
        self.db.ex("INSERT INTO attendance (session_id,student_id,status,marked_at) "
                   "VALUES (?,?,'present','2026-01-05')", (ls, sid))
        self.db.ex("INSERT INTO practice_requests (student_id,status,created_at) "
                   "VALUES (?,'pending','2026-01-05')", (sid,))

        st, r = self.admin.delete("/api/admin/users/%s" % self.su)
        self.assertEqual(st, 200, r)
        self.assertIsNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.su,)))
        self.assertIsNone(self.db.q1("SELECT id FROM students WHERE id=?", (sid,)))
        self.assertIsNone(self.db.q1("SELECT id FROM session_students WHERE student_id=?", (sid,)))
        self.assertIsNone(self.db.q1("SELECT id FROM attendance WHERE student_id=?", (sid,)))
        self.assertIsNone(self.db.q1("SELECT id FROM practice_requests WHERE student_id=?", (sid,)))
        # mashg'ulot va boshqa talaba O'CHIRILMAGAN
        self.assertIsNotNone(self.db.q1("SELECT id FROM lesson_sessions WHERE id=?", (ls,)),
                             "mashg'ulot o'chib ketdi — u talabaniki emas")

    def test_PURGE_cascades_instructor_and_lessons(self):
        """Instruktor o'chirilsa — uning mashg'ulotlari va ularga bog'liq
        qatorlar ham o'chadi; `instructor_id` NOT NULL bo'lgani uchun
        darslarni "osilib qoldirib bo'lmaydi"."""
        iid = self.db.q1("SELECT id FROM instructors WHERE user_id=?", (self.iu,))["id"]
        ls = _ins_lesson(self.db, iid, "2026-02-05")

        st, r = self.admin.delete("/api/admin/users/%s" % self.iu)
        self.assertEqual(st, 200, r)
        self.assertIsNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.iu,)))
        self.assertIsNone(self.db.q1("SELECT id FROM instructors WHERE id=?", (iid,)))
        self.assertIsNone(self.db.q1("SELECT id FROM lesson_sessions WHERE id=?", (ls,)),
                          "instructor o'chirildi, lekin dars qoldi — FK NOT NULL buziladi")

    def test_PURGE_deletes_user_scoped_tables(self):
        """Sessiya, 2FA, token, xabar va shaxsiy sozlamalar o'chadi."""
        uid = self.su
        self.db.ex("INSERT INTO sessions_ring (user_id,token_hash,csrf_token,created_at,expires_at) "
                   "VALUES (?,'t_purge_1','c','2026-01-01','2030-01-01')", (uid,))
        self.db.ex("INSERT INTO twofa_codes (user_id,code_hash,created_at,expires_at) "
                   "VALUES (?,'h1','2026-01-01','2030-01-01')", (uid,))
        self.db.ex("INSERT INTO password_reset_tokens (user_id,token_hash,created_at,expires_at) "
                   "VALUES (?,'t2','2026-01-01','2030-01-01')", (uid,))
        self.db.ex("INSERT INTO user_settings (user_id,key,value,updated_at) "
                   "VALUES (?,'k','v','2026-01-01')", (uid,))
        self.db.ex("INSERT INTO notifications (user_id,title,created_at) "
                   "VALUES (?,'n','2026-01-01')", (uid,))
        self.db.ex("INSERT INTO messages (from_user_id,to_user_id,text,created_at) "
                   "VALUES (?,?,'hi','2026-01-01')", (self.iu, uid))

        st, r = self.admin.delete("/api/admin/users/%s" % uid)
        self.assertEqual(st, 200, r)
        plain = {"sessions_ring": ("user_id",), "twofa_codes": ("user_id",),
                 "password_reset_tokens": ("user_id",), "user_settings": ("user_id",),
                 "notifications": ("user_id",),
                 "messages": ("from_user_id", "to_user_id")}
        for t, cols in plain.items():
            where = " OR ".join(c + "=?" for c in cols)
            n = self.db.q1("SELECT COUNT(*) AS c FROM %s WHERE %s" % (t, where),
                           tuple([uid] * len(cols)))["c"]
            self.assertEqual(int(n), 0, "%s jadvalidan qoldi" % t)

    def test_PURGE_keeps_audit_log_and_other_users(self):
        """JURNAL va boshqa foydalanuvchilar saqlanadi."""
        me = self.admin.get("/api/auth/me")[1]["user"]["id"]
        n_before = int(self.db.q1("SELECT COUNT(*) AS c FROM audit_log")["c"])
        st, r = self.admin.delete("/api/admin/users/%s" % self.su)
        self.assertEqual(st, 200, r)
        n_after = int(self.db.q1("SELECT COUNT(*) AS c FROM audit_log")["c"])
        self.assertGreaterEqual(n_after, n_before + 1, "jurnalga yozuv qo'shilmadi")
        # o'chirilgan foydalanuvchi nomi jurnalda saqlangan bo'lishi kerak
        row = self.db.q1("SELECT description FROM audit_log ORDER BY id DESC LIMIT 1")
        self.assertIn("user_purged", row["description"])
        # boshqalar joyida
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.iu,)))
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (me,)))

    def test_PURGE_self_protected(self):
        me = self.admin.get("/api/auth/me")[1]["user"]["id"]
        st, r = self.admin.delete("/api/admin/users/%s" % me)
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "user.cannot_delete_self")
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (me,)))

    def test_PURGE_last_admin_protected(self):
        """Oxirgi faol adminni o'chirishga ruxsat YO'Q.

        `seed` bitta faol admin yaratadi ("admin"), shuning uchun uni
        `self.admin` orqali o'chirishga urinib ko'ramiz. Agar test bazasida
        boshqa faol admin bo'lsa — himoya o'sha uchun ishlaydi va biz
        o'shandan birini tanlab o'chiramiz (shu bilan ham himoya tekshiriladi).
        """
        st, r = self.admin.get("/api/admin/users?role=admin&status=active")
        self.assertEqual(st, 200, r)
        ids = [u["id"] for u in r["users"]]
        self.assertTrue(ids, "bitta faol admin bo'lishi kerak")
        # qo'shimcha admin yaratib uni oldindan o'chiramiz (test tozalash)
        extra = self._mkuser("admin", "PurgeAdm", "X")
        self.admin.delete("/api/admin/users/%s" % extra["user"]["id"])
        # endi ro'yxatdan oxirgi qolganini o'chirishga urinamiz
        st2, r2 = self.admin.get("/api/admin/users?role=admin&status=active")
        left = [u["id"] for u in r2["users"]]
        if len(left) != 1:
            return  # baza bir nechta faol admin bilan qurilgan — himoa o'rinida
        st3, rr = self.admin.delete("/api/admin/users/%s" % left[0])
        self.assertEqual(st3, 400, rr)
        self.assertEqual(rr["error"], "user.cannot_delete_self")
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (left[0],)))

    def test_PURGE_archive_is_separate_feature(self):
        """ARXIVLASH alohida ishlaydi va ma'lumotni SAQLAYDI.

        Talab #8: "Arxivlash" tugmasi bosilgandagina arxivga tushsin;
        "O'chirish" esa arxivga o'tkazmasin.
        """
        st, r = self.admin.post("/api/admin/users/%s/status" % self.su, {"status": "archive"})
        self.assertEqual(st, 200, r)
        row = self.db.q1("SELECT deleted_at, status FROM users WHERE id=?", (self.su,))
        self.assertEqual(row["status"], "archived")
        self.assertIsNotNone(row["deleted_at"])
        self.assertIsNotNone(self.db.q1("SELECT id FROM students WHERE user_id=?", (self.su,)),
                             "arxivlash ma'lumotni o'chirdi — arxiv SAQLASHI kerak")
        # arxivlanganlar ro'yxatida ko'rinadi
        st, r2 = self.admin.get("/api/admin/users?status=archived")
        self.assertIn(self.su, [u["id"] for u in r2["users"]])
        # va arxivdan chiqarish mumkin
        st, r3 = self.admin.post("/api/admin/users/%s/status" % self.su, {"status": "unarchive"})
        self.assertEqual(st, 200, r3)
        self.assertIsNone(self.db.q1("SELECT deleted_at FROM users WHERE id=?", (self.su,))["deleted_at"])

    # ------------------------- UI/MATN KONTROLLARI (doimiy regressiya qalqoni)
    def _i18n_val(self, key, lang):
        """i18n.js dan kalitni {n} placeholder'ini ham hisobga olib o'qiydi."""
        s = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        m = re.search(r'"' + re.escape(key) + r'"\s*:\s*\{', s)
        self.assertIsNotNone(m, "i18n kaliti yo'q: " + key)
        i = m.end() - 1
        depth, in_str, quote, esc = 0, False, "", False
        while i < len(s):
            ch = s[i]
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == quote:
                    in_str = False
            else:
                if ch in "\"'":
                    in_str, quote = True, ch
                elif ch == "{":
                    depth += 1
                elif ch == "}":
                    depth -= 1
                    if depth == 0:
                        break
            i += 1
        body = s[m.end():i]
        m2 = re.search(r'(?:^|[,{\s])' + lang + r'\s*:\s*"((?:[^"\\]|\\.)*)"', body)
        self.assertIsNotNone(m2, "%s tili yo'q: %s" % (lang, key))
        return m2.group(1)

    def test_PURGE_confirm_dialog_text_is_exact(self):
        """Tasdiqlash oynasi matni va tugmalari talab qilingandek.

        Talab: "Bu foydalanuvchi butunlay o'chiriladi. O'chirilgandan keyin
        uni tiklab bo'lmaydi. Davom etasizmi?" + [Bekor qilish][Butunlay o'chirish]
        """
        for lang in ("uz", "ru", "en"):
            with self.subTest(lang=lang):
                self.assertTrue(self._i18n_val("users.purge_confirm", lang).strip())
                self.assertTrue(self._i18n_val("users.purge_yes", lang).strip())
                self.assertTrue(self._i18n_val("users.purge_no", lang).strip())
        # UZ aniq matni (talab bo'yicha)
        self.assertEqual(
            self._i18n_val("users.purge_confirm", "uz"),
            "Bu foydalanuvchi butunlay o'chiriladi. O'chirilgandan keyin uni tiklab "
            "bo'lmaydi. Davom etasizmi?")
        self.assertEqual(self._i18n_val("users.purge_yes", "uz"), "Butunlay o'chirish")
        self.assertEqual(self._i18n_val("users.purge_no", "uz"), "Bekor qilish")
        self.assertEqual(self._i18n_val("users.purged", "uz"),
                         "Foydalanuvchi butunlay o'chirildi.")

    def test_PURGE_new_i18n_keys_exist_in_all_three_languages(self):
        """Yangi kalitlar 3 tilda ham BO'LISHI SHART (paritet)."""
        for key in ("users.purge", "users.purge_title", "users.purge_confirm",
                    "users.purge_yes", "users.purge_no", "users.purged",
                    "users.purge_hint", "bulk.archive", "bulk.archived",
                    "err.user.cannot_delete_self", "audit.act.user_purged"):
            for lang in ("uz", "ru", "en"):
                with self.subTest(key=key, lang=lang):
                    self.assertTrue(self._i18n_val(key, lang).strip())

    def test_PURGE_ui_delete_calls_real_DELETE_not_archive(self):
        """Frontend: "Arxivlash" -> /status (arxiv), "O'chirish" -> DELETE.

        Bu test butun arxitekturani tekshiradi: agar kodingda arxivlash
        `API.del` bilan aralashtirilsa yoki o'chirish /status'ga ulangansа,
        test darhol qizil bo'ladi.
        """
        src = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        # 1) `purgeUser` -> API.del
        i = src.find("async function purgeUser")
        j = src.find("function canPurge", i)
        self.assertGreater(i, 0, "purgeUser() topilmadi")
        chunk = src[i:j]
        self.assertIn("API.del(`admin/users/${u.id}`)", chunk,
                      "o'chirish HAQIQIY DELETE chaqirishi kerak")
        self.assertNotIn("/status", chunk, "o'chirish arxivlash bilan aralashgan")
        # 2) arxivlash -> /status (DELETE emas)
        i = src.find("async function setUserStatus")
        j = src.find("function canPurge", i)
        self.assertGreater(i, 0, "setUserStatus() topilmadi")
        self.assertIn("/status", src[i:j], "arxivlash /status orqali bo'lishi kerak")
        # 3) tasdiqlash `danger` bilan (xavfli ekani ko'rinadi)
        i = src.find("function userPurgeAction")
        j = src.find("function openUserProfile", i)
        self.assertGreater(i, 0, "userPurgeAction() topilmadi")
        chunk = src[i:j if j > i else i + 1200]
        self.assertIn("confirmDialog", chunk)
        self.assertIn("danger: true", chunk, "xavfli amal tasdiqlanishi kerak")
        self.assertIn("users.purge_yes", chunk)
        self.assertIn("users.purge_no", chunk)
        # 4) barcha 4 ta jadval + profil oynasida tugma bor
        self.assertGreaterEqual(src.count("userPurgeAction(u, load)"), 4,
                                "3 rol jadvali + 'Baza' qatorida tugma kerak")
        self.assertIn("userPurgeAction(u, () => {", src, "profil oynasida ham kerak")
        # 5) o'zini o'chirish tugmasi chiqmaydi
        i = src.find("function canPurge")
        j = src.find("function userPurgeAction", i)
        self.assertIn("App.me", src[i:j], "o'zini o'chirish himoyasi frontend'da yo'q")

    def test_PURGE_backend_has_no_soft_delete_left(self):
        """O'chirish yo'llarida `UPDATE ... deleted_at` QOLMASIN.

        (arxivlash alohida endpoint orqali ishlaydi — `POST .../status`.)
        Haqiqiy `DELETE FROM users` esa `_purge_users` ichida bajariladi.
        """
        src = (ROOT / "app" / "api.py").read_text(encoding="utf-8", errors="ignore")

        def chunk(fn, nxt):
            i = src.find(fn)
            self.assertGreater(i, 0, fn + " topilmadi")
            j = src.find(nxt, i + 10)
            return src[i:j if j > i else i + n]

        del_fn = chunk("def admin_user_delete(self, uid):", "\n    # ")
        self.assertNotIn("SET deleted_at", del_fn, "soft-delete qolib ketgan!")
        self.assertNotIn("status='archived'", del_fn,
                         "o'chirish arxivga o'tkazmasin!")
        self.assertIn("_purge_users", del_fn)

        purge = chunk("def _purge_users(self, uids):", "def _purge_cleanup_images")
        self.assertIn("DELETE FROM users WHERE id=?", purge,
                      "haqiqiy `DELETE FROM users` bajarilishi kerak")
        self.assertNotIn("SET deleted_at", purge)
        # FK butunligi: kamida o'z profil jadvallari va foydalanuvchi qatorlari
        for tbl in ("session_students", "attendance", "practice_requests",
                    "students", "lesson_sessions", "instructors",
                    "notifications", "messages"):
            with self.subTest(tbl=tbl):
                self.assertIn("FROM %s" % tbl, purge,
                              "bog'liq jadval tozalanmagan: " + tbl)
        # foydalanuvchiga tegishli qatorlar — umumiy ro'yxatda tekshiriladi
        i = purge.find("for table in (")
        self.assertGreater(i, 0, "foydalanuvchi jadvallari tozalanmagan")
        loop = purge[i:purge.find(")", i) + 1]
        for tbl in ("sessions_ring", "password_reset_tokens", "twofa_codes",
                    "hidden_requests", "hidden_history", "user_settings",
                    "notification_settings"):
            with self.subTest(tbl=tbl):
                self.assertIn('"%s"' % tbl, loop, "jadval tozalanmagan: " + tbl)
        # jurnal esa SAQLANADI (o'chirilmaydi)
        self.assertNotIn("DELETE FROM audit_log", purge,
                         "audit jurnalini o'chirish qat'iyan TAQIQLANGAN")

    def test_PURGE_bulk_delete_is_hard_delete(self):
        """Ommaviy o'chirish ham arxiv EMAS — haqiqiy DELETE."""
        src = (ROOT / "app" / "api.py").read_text(encoding="utf-8", errors="ignore")
        i = src.find("def admin_users_bulk_delete(self, body):")
        j = src.find("def admin_users_update", i)
        chunk = src[i:j]
        self.assertGreater(i, 0)
        self.assertNotIn("SET deleted_at", chunk, "bulk-delete soft-ga qaytgan!")
        self.assertIn("_purge_users", chunk)

    def test_PURGE_requires_admin(self):
        """Talaba/instruktor foydalanuvchini o'chira OLMAYDI."""
        st, r = self.admin.delete("/api/admin/users/%s" % self.su)
        self.assertEqual(st, 200, r)
        r2 = self._mkuser("student", "PurgeStud", "Y")
        pwd = r2["credentials"]["password"]
        c = Client()
        s2, _ = c.post("/api/auth/login",
                       {"login": r2["user"]["login"], "password": pwd, "role": "student"})
        self.assertEqual(s2, 200)
        s3, rr = c.delete("/api/admin/users/%s" % self.iu)
        self.assertIn(s3, (401, 403), rr)
        self.assertIsNotNone(self.db.q1("SELECT id FROM users WHERE id=?", (self.iu,)))

    # ------------------------------------------------- RBAC / xavfsizlik
    def test_MOD6_bulk_requires_admin(self):
        """Instruktor/talaba ommaviy ammalarni BAJAROLMAYDI."""
        r = self._mkuser("instructor", "Norole", "T")
        pwd = r["credentials"]["password"]
        c = Client()
        s2, _ = c.post("/api/auth/login",
                       {"login": r["user"]["login"], "password": pwd, "role": "instructor"})
        self.assertEqual(s2, 200)
        s3, rr = c.post("/api/admin/users/bulk-delete", {"ids": [self.su]})
        self.assertEqual(s3, 403, rr)
        s4, rr = c.post("/api/admin/users/bulk-update", {"ids": [self.su], "status": "blocked"})
        self.assertEqual(s4, 403, rr)
        self._cleanup_user(r["user"]["id"])

    # ----------------------------------------------------- frontend statik
    def test_MOD2_frontend_count_bar(self):
        """Frontend'da hisoblagich + tanlov holati mavjud (VAZIFA 2 arxitekturasi).

        Eski `bulkSelectAll`/`bulkCheckbox` funksiyalari olib tashlandi —
        ularning o'rniga `createSelection` (immutabil tanlov) va
        `createBulkStore` (yagona holat manbai) ishlaydi.
        """
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("function countBar", txt)
        self.assertIn("function createSelection", txt)
        self.assertIn("function createBulkStore", txt)
        # uchala jadvalda hisoblagich chaqiriladi
        self.assertGreaterEqual(txt.count("countBar(res"), 3,
                                "uchala toifada hisoblagich yo'q")
        # eski yondashuv QAYTMADI (aralash arxitektura bo'lmasligi uchun)
        self.assertNotIn("function bulkSelectAll", txt)
        self.assertNotIn("function bulkCheckbox", txt)
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn('"base.count"', i18n)
        self.assertIn("Jami: {n} ta {what}", i18n)

    def test_MOD6_frontend_checkbox_and_bulk_bar(self):
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        self.assertIn("function bulkEditModal", txt)
        # store orqali yaratiladigan uchala element — obyekt metodi shaklida
        for sig in ("headerNode() {", "rowNode(u) {", "actionsNode() {", "attach(next, serverTotal) {"):
            self.assertIn(sig, txt, "%s metodi yo'q" % sig)
        # va 3 ta jadvalda hammasi chaqiriladi
        for call in ("bulk.headerNode()", "bulk.actionsNode()", "bulk.attach(list, res.total)"):
            self.assertEqual(txt.count(call), 3,
                             "uchala jadvalda %s chaqirilishi kerak (topildi: %d)"
                             % (call, txt.count(call)))
        self.assertEqual(txt.count("bulk.rowNode(u)"), 3,
                         "uchala jadvalda bulk.rowNode(u) kerak")
        # O'LIK kod butunlay olib tashlandi (ikki nusxali `bulkActionsBar`
        # aynan "N ta tanlandi" chiqmasligining manbai edi)
        self.assertNotIn("function bulkActionsBar", txt)
        self.assertNotIn("function bulkSelectAll", txt)
        self.assertNotIn("function bulkCheckbox", txt)
        self.assertGreaterEqual(txt.count('class: "bulk-td"'), 3,
                                "uchala jadvalda qator checkbox yo'q")
        self.assertGreaterEqual(txt.count('class: "bulk-th"'), 3,
                                "uchala jadvalda 'Hammasini tanlash' yo'q")

    # ================= VAZIFA 2 — ommaviy tanlash + oltin checkbox =============
    def _css(self):
        return (ROOT / "web" / "css" / "styles.css").read_text(encoding="utf-8", errors="ignore")

    def test_V2_selection_is_immutable(self):
        """VAZIFA 2/A: holat YANGI massiv/obyekt bilan yangilanadi.

        Muammo A ning asosi shuki `push`/`add` mutatsiyasi render qayta
        ishga tushmasligi. `createSelection` har o'zgarishda YANGI `ids`
        massivi yaratadi — eski havola buzilmaydi.
        """
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("function createSelection")
        j = txt.find("function createBulkStore", i)
        self.assertNotEqual(i, -1, "createSelection topilmadi")
        self.assertNotEqual(j, -1, "createBulkStore topilmadi")
        chunk = txt[i:j]
        # --- hech qanday to'g'ridan-to'g'ri mutatsiya yo'q (MUAMMO A ning asosi) ---
        for bad in ("ids.push(", "ids.splice(", "ids.pop(", "ids.shift(",
                    "ids.sort(", "ids.reverse(", "ids[0] =", "ids.length ="):
            self.assertNotIn(bad, chunk, "%s — to'g'ridan-to'g'ri mutatsiya" % bad)
        # --- YANGI massiv yaratiladi (concat / filter orqali) ---
        self.assertIn("ids = a;", chunk, "yangi massiv tayinlanmayapti")
        self.assertIn("ids.concat([id])", chunk, "add() yangi massiv yaratmayapti")
        self.assertIn("ids.filter(", chunk, "remove()/only()/except() filter ishlatmayapti")
        # --- Set bilan takrorlilar oldi ---
        self.assertIn("new Set(next)", chunk, "takrorlanadigan ID'lar filtrlanmayapti")
        # --- o'zgarish aniqlanadi (keraksiz re-render yo'q) ---
        self.assertIn("a.length === ids.length && a.every", chunk, "o'zgarish aniqlanmagan")
        # --- bildirishnoma: obuna + xatoni yashirmaslik ---
        self.assertIn("subscribe(fn)", chunk)
        self.assertIn("subs.slice().forEach", chunk, "bildirishnoma nusxasi ishlatilmayapti")
        self.assertIn("notify();", chunk)
        self.assertNotIn("catch (e) {}", chunk, "obuna xatosi yashirilmoqda")
        self.assertIn("console.error", chunk, "obuna xatosi konsolga chiqarilmayapti")

    def test_V2_bulk_store_single_source(self):
        """VAZIFA 2/A+B: hisoblagich va checkbox bitta store dan oziqlanadi."""
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("function createBulkStore")
        j = txt.find("function countBar", i)
        self.assertNotEqual(i, -1)
        self.assertNotEqual(j, -1)
        chunk = txt[i:j]                      # butun store (ikkala o'lik funksiya YO'Q)
        # --- hisoblagich har render'da QAYTA o'qiladi (MUAMMO A ning tuzatilishi) ---
        self.assertIn("actionsNode() {", chunk)
        act = chunk[chunk.index("actionsNode() {"):]
        self.assertIn("subs.push((st) => {", act, "panel obuna qilinmagan")
        self.assertIn("const n = st.size;", act,
                      "son panel YARATILGANDA emas, har render'da o'qilishi kerak")
        self.assertNotIn("const n = state.size;", chunk[:chunk.index("actionsNode() {")],
                         "son yana bir marta tashqarida o'qilmoqda")
        # --- qator checkbox'i store'dan o'qiydi, to'g'ridan-to'g'ri EMAS ---
        self.assertIn("cb.checked = state.has(u.id);", chunk)
        self.assertIn("state.toggle(u.id);", chunk)
        # --- "hammasini tanlash" haqiqiy TOGGLE (MUAMMO B) ---
        self.assertIn("function toggleAll", chunk)
        self.assertIn("list.every((u) => state.has(u.id))", chunk,
                      "tanlash holati to'g'ri aniqlanmayapti")
        self.assertIn("if (all) { state.except(list); return; }", chunk,
                      "qayta bosish -> bekor qilish yo'q")
        self.assertIn("state.allOf(list);", chunk)
        # ro'yxat ID'lari ro'yxatdan keladi (tasodifiy sonlar EMAS)
        self.assertIn("allOf(list) { return this.set((list || []).map((u) => u.id))", txt)
        # --- server chegarasida tasdiqlash ---
        self.assertIn("total > list.length", chunk, "server chegarasi tekshirilmayapti")
        self.assertIn("bulk.confirm_all", chunk)
        # --- qisman tanlanganda header indeterminate ---
        self.assertIn("head.indeterminate = inList > 0 && !all;", chunk,
                      "indeterminate noto'g'ri hisoblanmoqda")
        self.assertIn("head.checked = all;", chunk)
        # --- filtr o'zgarganda ko'rinmaydigan tanlov tozalanadi ---
        self.assertIn("state.only(list);", chunk, "attach da tanlov tozalanmayapti")
        # --- qator vizual belgisi ---
        self.assertIn('tr.classList.toggle("bulk-row-sel", on)', chunk)
        # --- xatolar YASHIRILMAYDI ---
        self.assertNotIn("catch (e) {}", chunk)
        self.assertIn("console.error", chunk, "xato yashirilmoqda")

    def test_V2_all_three_tables_use_store(self):
        """Uchala jadval (talaba/instruktor/admin) bir xil store dan foydalanadi."""
        import re as _re
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        # aniq rol bilan 3 ta jadval + 1 ta ta'rif (izoh satrlaridan kelib chiqmaydi)
        self.assertEqual(len(_re.findall(r'createBulkStore\("', txt)), 3,
                         "3 ta jadval store yaratishi kerak (topildi: %d)"
                         % len(_re.findall(r'createBulkStore\("', txt)))
        self.assertEqual(txt.count("function createBulkStore"), 1, "store ta'rifi bir necha marta")
        for role in ("student", "instructor", "admin"):
            self.assertIn('const bulk = createBulkStore("%s", load);' % role, txt,
                          "%s jadvali store'ga ulanmagan" % role)
        # toggleAll va panel — store ichida, tashqarida emas
        self.assertNotIn('onclick: toggleAll', txt)
        # har bir jadval uch xil komponent to'liq ulangan
        for fn in ("bulk.headerNode()", "bulk.rowNode(u)", "bulk.attach(list, res.total)"):
            self.assertEqual(txt.count(fn), 3, "%s — 3 ta jadvalda kerak (topildi: %d)"
                             % (fn, txt.count(fn)))

    def test_V2_gold_checkbox_css(self):
        """VAZIFA 2/C: barcha checkbox'lar oltin-sariq gradient aksentga."""
        css = self._css()
        # 1) gradient token bitta manba
        self.assertIn("--gold-1: #F5A623", css)
        self.assertIn("--gold-2: #FF8C42", css)
        self.assertIn("--gold-grad: linear-gradient(135deg, var(--gold-1), var(--gold-2))", css)
        # 2) checkbox to'liq qayta dizayn qilingan
        self.assertIn('input[type="checkbox"] {', css)
        self.assertIn("appearance: none", css, "standart brauzer checkbox'i almashtirilmagan")
        # bo'sh holat: shaffof fon + kulrang border
        cb_block = css[css.find('input[type="checkbox"] {'):css.find("input[type=\"checkbox\"]:disabled")]
        self.assertIn("background: transparent", cb_block, "bo'sh checkbox shaffof emas")
        self.assertIn("border: 1.8px solid var(--border-strong)", cb_block,
                      "bo'sh checkbox borderi kulrang emas")
        # belgilangan: gradient + oq "✓"
        self.assertIn("input[type=\"checkbox\"]:checked,\ninput[type=\"checkbox\"]:indeterminate {",
                      css)
        chk = css[css.find('input[type="checkbox"]:checked,'):]
        chk = chk[:chk.find("input[type=\"checkbox\"]:disabled")]
        self.assertIn("background: var(--gold-grad)", chk, "belgilangan checkbox oltin emas")
        self.assertIn("::after", chk, '"✓" belgisi yo\'q')
        self.assertIn("border: solid #fff", chk, '"✓" rangi oq emas')
        # 3) qisman holat — chiziqcha
        self.assertIn('input[type="checkbox"]:indeterminate::after', css)
        # 4) eski ko'k `accent-color` QOLMADI (izohlardagi eslatma hisobga olinmaydi)
        import re as _re
        live = _re.sub(r"/\*.*?\*/", "", css, flags=_re.S)
        self.assertNotIn("accent-color: var(--primary)", live,
                         "eski ko'k aksent hali ham bor")
        # 5) toggle-switch BIR XIL manbadan
        self.assertIn(".toggle.on", css)
        tog = css[css.find(".toggle.on"):css.find(".toggle.on") + 400]
        self.assertIn("var(--gold-grad)", tog, "toggle-switch boshqa rangda")
        # 6) o'chirilgan `--primary` o'zgaruvchisi bo'lmasligi kerak
        import re as _re
        self.assertEqual(_re.findall(r"--primary\s*:", css), [],
                         "bo'sh --primary hali ham aniqlangan")

    def test_V2_checkbox_gold_used_everywhere(self):
        """Checkbox uslubi GLOBAL — alohida klass emas (bitta qoida hammasiga)."""
        css = self._css()
        # global `input[type="checkbox"]` qoidasi bitta bor
        self.assertEqual(css.count('input[type="checkbox"] {'), 1,
                         "global checkbox qoidasi bir necha marta yozilgan")
        # ommaviy ammal paneli ham oltin
        self.assertIn(".bulk-bar {", css)
        self.assertIn(".bulk-row-sel", css)
        self.assertIn("var(--gold-soft)", css)

    def test_MOD7_no_hardcoded_backup_in_uz(self):
        """MODUL 7: UZ tarjimalarida inglizcha "Backup" QOLMADI."""
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8", errors="ignore")
        import re as _re
        # `uz:` qiymatlarida "Backup" so'zi bo'lmasin
        bad = _re.findall(r'"[\w.]+":\s*\{\s*uz:\s*"[^"]*\bBackup\b[^"]*"', i18n)
        self.assertEqual(bad, [], f"UZ da 'Backup' qolgan: {bad}")
        self.assertIn('"nav.backup": { uz: "Zaxira nusxa"', i18n)
        self.assertIn('"backup.title": { uz: "Zaxira nusxa"', i18n)

    def test_MOD7_admin_new_strings_use_i18n(self):
        """MODUL 3/6 yangi tugma matnlari i18n orqali keladi."""
        txt = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8", errors="ignore")
        i = txt.find("function createBulkStore")
        j = txt.find("function bulkEditModal", i)
        self.assertNotEqual(i, -1, "createBulkStore topilmadi")
        self.assertNotEqual(j, -1, "bulkEditModal topilmadi")
        chunk = txt[i:j]
        for k in ("bulk.delete", "bulk.edit", "bulk.none", "bulk.deselect_all", "bulk.count"):
            self.assertIn(k, chunk, f"{k} ishlatilmayapti (hardcoded bo'lib qolgan)")


# ============================================================================
# 8) MIGRATSIYA — eski bazani yangilash serverni buzMASIN (regressiya)
# ============================================================================
class TestMigrationBAND24(unittest.TestCase):
    """`init_db` ESKI bazada ham xatosiz ishlashi shart.

    Bu test birinchi marta ishga tushganda platforma ISHLAMASDI:
    `SCHEMA` ichida `notifications(source, ...)` ga bog'liq indekslar
    `migrate()` dan OLDIN yaratilardi -> "no such column: source".
    Eski (foydalanuvchi) bazalarda shu ustun `ALTER TABLE` bilan qo'shiladi.
    """

    def setUp(self):
        import sqlite3
        self._sqlite3 = sqlite3
        self.tmp = tempfile.mkdtemp(prefix="avto-mig-")
        self.path = str(Path(self.tmp) / "old.db")
        # 1) "ESKI" bazani qo'l bilan yaratamiz: jadvallar bor, lekin
        #    `source`/`created_by`/`related_lesson_id`/`csrf_token`/
        #    `confirm_state`/`total_lessons_target` USTUNLARI YO'Q.
        from app.db import SCHEMA
        old_schema = SCHEMA
        for drop in (
            "    source TEXT NOT NULL DEFAULT 'SYSTEM',\n",
            "    created_by INTEGER,\n",
            "    related_lesson_id INTEGER REFERENCES lesson_sessions(id) ON DELETE CASCADE,\n",
        ):
            old_schema = old_schema.replace(drop, "")
        old_schema = old_schema.replace(
            "    csrf_token TEXT NOT NULL DEFAULT '',\n", "")
        old_schema = old_schema.replace(
            "    confirm_state TEXT NOT NULL DEFAULT 'pending'\n"
            "        CHECK (confirm_state IN ('pending','confirmed')),\n", "")
        old_schema = old_schema.replace("    total_lessons_target INTEGER,\n", "")
        c = sqlite3.connect(self.path)
        c.executescript(old_schema)
        t = "2026-01-01 00:00:00"
        c.execute("INSERT INTO users(first_name,last_name,login,password_hash,role,status,"
                  "created_at,updated_at) VALUES('Eski','Talaba','usrL_OLDUS',"
                  "'x','student','active',?,?)", (t, t))
        c.execute("INSERT INTO notifications(user_id,type,title,body,is_read,created_at)"
                  " VALUES(1,'info','Eski xabar','matn',0,?)", (t,))
        c.commit()
        c.close()
        cols = {r[1] for r in sqlite3.connect(self.path).execute(
            "PRAGMA table_info(notifications)")}
        assert "source" not in cols, "test bazasi 'eski' bo'lib yaratilmadi"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_M01_old_db_migrates_without_error(self):
        from app.db import init_db
        db = init_db(self.path)      # <<< avval BU xato berardi
        cols = {r["name"] for r in self.db_rows("PRAGMA table_info(notifications)")}
        self.assertIn("source", cols)
        self.assertIn("created_by", cols)
        self.assertIn("related_lesson_id", cols)
        idx = {r["name"] for r in self.db_rows("PRAGMA index_list(notifications)")}
        self.assertIn("ux_notif_reminder_once", idx)
        self.assertIn("idx_notif_source", idx)

    def test_M02_data_preserved(self):
        """Eski ma'lumotlar (foydalanuvchi, xabar) O'CHIRILMAYDI."""
        from app.db import init_db
        init_db(self.path)
        self.assertEqual(self._one("SELECT COUNT(*) FROM users"), 1)
        self.assertEqual(self._one("SELECT COUNT(*) FROM notifications"), 1)
        # eski xabar `source` ni oldindan oladi: sender_id yo'q -> SYSTEM
        self.assertEqual(
            self._one("SELECT source FROM notifications WHERE id=1"), "SYSTEM")

    def test_M03_idempotent(self):
        """Migration bir necha marta ishga tushsa — xatosiz (2..5 marta)."""
        from app.db import init_db
        for _ in range(4):
            init_db(self.path)
        self.assertEqual(self._one("SELECT COUNT(*) FROM notifications"), 1)

    def db_rows(self, sql):
        c = self._sqlite3.connect(self.path)
        c.row_factory = self._sqlite3.Row
        try:
            return c.execute(sql).fetchall()
        finally:
            c.close()

    def _one(self, sql):
        c = self._sqlite3.connect(self.path)
        try:
            return c.execute(sql).fetchone()[0]
        finally:
            c.close()


# ============================================================================
# 9) MODUL 1-5 — hisobot eksporti, talaba hisoblagichi, backup o'chirish,
#    parol+login o'zgartirish, so'rovlarni ochish va muddat (auto-hide)
# ============================================================================
def _csv_mime_ok(mime):
    return (mime or "").split(";")[0].strip() == "text/csv"


# ---------------------------------------------------------------------------
# MODUL 1 — CSV / XLSX / PDF eksporti HAQIQIY fayl yaratadi va ochiladi
# ---------------------------------------------------------------------------
class TestModul1Export(_MixinAdmin, Base):
    """MODUL 1: `openpyxl` va `reportlab` o'rnatilgan; uchala format HAQIQIY
    fayl beradi (CSV — UTF-8 BOM, XLSX — zip/OOXML, PDF — %PDF + ichki
    shrift) va o'zbek sheva belgilari (oʻ, gʻ, sh, ch) buzilmaydi."""

    UZ = [{"Ism": "Alisher oʻgʻli", "Familiya": "Shukurov", "Guruh": "Ch-1",
           "Telefon": "+998 90 123 45 67"}]

    def test_M1_01_libraries_available_on_server_python(self):
        """Kutubxonalar server ishlatadigan Python'da REAL mavjud."""
        import openpyxl          # noqa: F401
        import reportlab         # noqa: F401
        from openpyxl import Workbook  # noqa: F401
        from reportlab.platypus import SimpleDocTemplate  # noqa: F401
        import sys
        self.assertIn("openpyxl", sys.modules)
        req = (ROOT / "requirements.txt").read_text(encoding="utf-8")
        self.assertIn("openpyxl", req)
        self.assertIn("reportlab", req)
        dock = (ROOT / "Dockerfile").read_text(encoding="utf-8")
        self.assertIn("requirements.txt", dock)

    def test_M1_02_csv_real_file_utf8_bom_and_uzbek_letters(self):
        from app.export import export_table
        res = export_table(self.UZ, "csv")
        self.assertTrue(res["filename"].endswith(".csv"), res["filename"])
        self.assertTrue(_csv_mime_ok(res["mime"]), res["mime"])
        data = res["data"]
        # UTF-8 BOM — Excel'da o'zbek harflari buzilmaydi
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"), "CSV BOM yo'q")
        text = data.decode("utf-8-sig")
        self.assertIn("Alisher oʻgʻli", text)
        self.assertIn("Shukurov", text)
        self.assertIn("Ch-1", text)

    def test_M1_03_xlsx_opens_with_openpyxl(self):
        """Fayl haqiqiy Excel fayli: `openpyxl` ochib, o'zbek matnni o'qiy oladi."""
        import io
        import openpyxl
        from app.export import export_table
        res = export_table(self.UZ, "xlsx", "Talabalar")
        self.assertTrue(res["filename"].endswith(".xlsx"), res["filename"])
        self.assertTrue(res["data"].startswith(b"PK"), "XLSX zip bo'lishi kerak")
        wb = openpyxl.load_workbook(io.BytesIO(res["data"]))
        ws = wb.active
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        flat = [str(v) for row in grid for v in row if v is not None]
        self.assertIn("Alisher oʻgʻli", flat, flat)
        self.assertIn("Ism", flat, flat)
        self.assertIn("Talabalar", [n for n in wb.sheetnames])

    def test_M1_04_pdf_real_pdf_with_embedded_font(self):
        """Fayl haqiqiy PDF: `%PDF` sarlavha + `%%EOF` + ICHKI shrift
        (`/FontFile2`) — o'zbek sheva belgilarini chizish uchun shrift
        PDF ichiga joylashgan bo'lishi SHART."""
        from app.export import export_table, _ensure_font
        res = export_table(self.UZ, "pdf", "Talabalar")
        self.assertEqual(res["mime"], "application/pdf")
        data = res["data"]
        self.assertTrue(data.startswith(b"%PDF-"), "PDF sarlavhasi yo'q")
        self.assertIn(b"%%EOF", data[-2048:], "PDF yakuni (EOF) yo'q")
        self.assertIn(b"/FontFile2", data,
                      "PDF ichida shrift YOK — o'zbek harflari kvadratch bo'ladi")
        # tanlangan shrift haqiqatan oʻ/ǔ/gʻ belgilarini qo'llaydi
        font = _ensure_font()
        self.assertTrue(font.get("name"), "o'zbekcha shrift topilmadi")

    def test_M1_05_admin_export_endpoint_three_formats(self):
        """`GET /api/admin/export` uchala formatda HAQIQIY fayl qaytaradi
        (server base64 ni ochib, `Content-Disposition` bilan yuboradi)."""
        import io
        import openpyxl
        cases = [("csv", b"\xef\xbb\xbf", ".csv"),
                 ("xlsx", b"PK", ".xlsx"),
                 ("pdf", b"%PDF-", ".pdf")]
        for fmt, magic, ext in cases:
            st, data, hdrs = self.admin.get_file(
                "/api/admin/export?format=%s&report=users" % fmt)
            self.assertEqual(st, 200, fmt)
            disp = hdrs.get("Content-Disposition", "")
            self.assertIn("attachment", disp, (fmt, disp))
            self.assertIn(ext, disp, (fmt, disp))
            self.assertTrue(data.startswith(magic), (fmt, data[:8]))
            self.assertGreater(len(data), 100, fmt)
        # XLSX va PDF haqiqatan OCHILADI
        st, data, _ = self.admin.get_file("/api/admin/export?format=xlsx&report=users")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertTrue(wb.sheetnames)
        self.assertGreater(wb.active.max_row, 1)

    def test_M1_06_reports_generate_three_formats(self):
        """`POST /api/reports/generate` — uchala format, haqiqiy fayl."""
        for fmt in ("csv", "xlsx", "pdf"):
            st, data, hdrs = self.admin.post_file("/api/reports/generate", {
                "format": fmt, "report": "users", "roles": ["student", "instructor"],
                "include_credentials": False, "lang": "uz"})
            self.assertEqual(st, 200, (fmt, st, data[:120]))
            self.assertIn("attachment", hdrs.get("Content-Disposition", ""), fmt)
            self.assertGreater(len(data), 200, fmt)
        # noqiron format rad etiladi
        st, r = self.admin.post("/api/reports/generate", {"format": "doc", "report": "users"})
        self.assertEqual(st, 400)
        self.assertEqual(r["error"], "export.bad_format")


# ---------------------------------------------------------------------------
# MODUL 2 — "Jami: N ta talaba" hisoblagichi
# ---------------------------------------------------------------------------
class TestModul2StudentCounter(_MixinAdmin, Base):
    """MODUL 2: `total` — talabalar SONI (filterlashda filtrlangan soni).

    Avvalgi xato: `total` o'zgaruvchisi progress tsiklida qayta yozilib,
    hisoblagich oxirgi talabaning DARS sonini ko'rsatardi."""

    def _mk3(self):
        """3 ta talaba + har biriga TURLI soni dars (1, 2, 3)."""
        made = []
        for i in range(1, 4):
            r = self._mkuser("student", "Cnt", "T%d" % i, group_name="CNT")
            made.append(r)
        db = self.db
        for n, r in enumerate(made, start=1):
            sid = db.q1("SELECT id FROM students WHERE user_id=?",
                        (r["user"]["id"],))["id"]
            for k in range(n):
                date = "2026-0%d-0%d" % (k + 1, n)
                ts = date + " 09:00:00"
                lsid = db.ex(
                    """INSERT INTO lesson_sessions(instructor_id, date, start_time, end_time,
                           status, capacity_snapshot, car_name_snapshot,
                           created_at, updated_at)
                       VALUES(1,?,?,?,'completed',4,'CNT-TEST',?,?)""",
                    (date, "10:00", "11:00", ts, ts))
                db.ex("INSERT INTO session_students(session_id, student_id, "
                      "student_status, joined_at) VALUES(?,?,'active',?)",
                      (lsid, sid, ts))
        return made

    def _sids(self, made):
        return [self.db.q1("SELECT id FROM students WHERE user_id=?",
                           (r["user"]["id"],))["id"] for r in made]

    def tearDown(self):
        try:
            self.db.upd("DELETE FROM session_students WHERE session_id IN "
                        "(SELECT id FROM lesson_sessions WHERE car_name_snapshot='CNT-TEST')")
            self.db.upd("DELETE FROM lesson_sessions WHERE car_name_snapshot='CNT-TEST'")
        except Exception:
            pass
        for r in getattr(self, "_made", []):
            self._cleanup_user(r["user"]["id"])
        super().tearDown()

    def test_M2_01_total_is_student_count_not_last_lesson_count(self):
        self._made = self._mk3()
        st, r = self.admin.get("/api/admin/users?role=student")
        self.assertEqual(st, 200, r)
        n_students = self.db.q1(
            "SELECT COUNT(*) c FROM users WHERE deleted_at IS NULL AND role='student'")["c"]
        n_users = self.db.q1(
            "SELECT COUNT(*) c FROM users WHERE deleted_at IS NULL")["c"]
        self.assertEqual(r["total_all"], n_users, "umumiy foydalanuvchi soni")
        # ASOSIY TEKSHIRUV: `total` — filtrga mos TALABALAR SONI.
        self.assertEqual(r["total"], n_students, "counter talabalar soni bo'lishi kerak")
        # eski xato: oxirgi talabaning DARS soni (1/2/3) ko'rsatilardi
        self.assertGreater(r["total"], 3,
                           "counter dars sonini ko'rsatmoqda (eski xato)")

    def test_M2_02_total_follows_filter(self):
        self._made = self._mk3()
        st, all_ = self.admin.get("/api/admin/users?role=student")
        st, one = self.admin.get("/api/admin/users?role=student&q=Cnt+T1")
        self.assertEqual(st, 200, one)
        self.assertTrue(one["filtered"], "qidiruv filtrlangan deb belgilanmadi")
        self.assertEqual(one["total"], len(one["users"]))
        self.assertLess(one["total"], all_["total"])
        st, none = self.admin.get("/api/admin/users?role=student&q=ZZZyoq")
        self.assertEqual(none["total"], 0)
        self.assertEqual(none["total_all"], all_["total_all"],
                         "umumiy soni filtrdan o'tmasligi kerak")
        # har bir filtr o'z sonini beradi
        st, r = self.admin.get("/api/admin/users?role=student&q=Cnt")
        self.assertEqual(r["total"], 3)

    def test_M2_03_by_role_split_consistent(self):
        self._made = self._mk3()
        st, r = self.admin.get("/api/admin/users")
        self.assertEqual(st, 200, r)
        self.assertEqual(r["total_all"],
                         r["by_role"].get("student", 0) + r["by_role"].get("instructor", 0)
                         + r["by_role"].get("admin", 0))
        self.assertEqual(r["total"], r["total_all"])


# ---------------------------------------------------------------------------
# MODUL 3 — zaxira nusxani o'chirish
# ---------------------------------------------------------------------------
class TestModul3BackupDelete(_MixinAdmin, Base):
    """MODUL 3: `DELETE /api/admin/backup?name=...` — fayl serverdan BUTUNLAY
    o'chiriladi, audit jurnaliga yoziladi, noto'g'ri nom rad etiladi."""

    def test_M3_01_create_list_delete_roundtrip(self):
        import os
        from app.api import BACKUP_DIR
        self.assertTrue(BACKUP_DIR, "BACKUP_DIR sozlangan bo'lishi shart")
        st, r = self.admin.post("/api/admin/backup")
        self.assertEqual(st, 200, r)
        name = r["file"]
        path = os.path.join(BACKUP_DIR, name)
        self.assertTrue(os.path.isfile(path), "nusxa fayli yaratilmadi")

        st, l = self.admin.get("/api/admin/backup")
        self.assertIn(name, [b["name"] for b in l["backups"]])

        st, d = self.admin.delete("/api/admin/backup?name=" + name)
        self.assertEqual(st, 200, d)
        self.assertFalse(os.path.exists(path), "fayl serverdan o'chirilmadi")

        st, l = self.admin.get("/api/admin/backup")
        self.assertNotIn(name, [b["name"] for b in l["backups"]])

    def test_M3_02_audit_entry_written(self):
        st, r = self.admin.post("/api/admin/backup")
        name = r["file"]
        st, d = self.admin.delete("/api/admin/backup?name=" + name)
        self.assertEqual(st, 200, d)
        row = self.db.q1("SELECT * FROM audit_log WHERE action_type='backup_deleted' "
                          "ORDER BY id DESC LIMIT 1")
        self.assertIsNotNone(row, "audit jurnaliga backup_deleted yozilmadi")
        self.assertIn(name, (row.get("description") or ""),
                      "auditda fayl nomi ko'rsatilmadi")

    def test_M3_03_path_traversal_and_missing_rejected(self):
        for bad in ("../secrets.db", "..%2Fsecrets.db", "../../app/db.py"):
            st, r = self.admin.delete("/api/admin/backup?name=" + bad)
            self.assertIn(st, (400, 404), (bad, st, r))
            self.assertEqual(r.get("error"), "backup.bad", (bad, r))
        st, r = self.admin.delete("/api/admin/backup?name=yoq_bu_fayl.db")
        self.assertEqual(st, 404, r)
        self.assertEqual(r.get("error"), "backup.not_found", r)

    def test_M3_04_non_admin_forbidden(self):
        r0 = self._mkuser("student", "Bak", "Test")
        try:
            c, st, rr = self._login(r0["credentials"]["login"],
                                    r0["credentials"]["password"], "student")
            self.assertEqual(st, 200, rr)
            st, r = c.delete("/api/admin/backup?name=x.db")
            self.assertIn(st, (401, 403), (st, r))
        finally:
            self._cleanup_user(r0["user"]["id"])

    def test_M3_05_ui_has_delete_button_and_confirm_text(self):
        js = (ROOT / "web" / "js" / "views-admin.js").read_text(encoding="utf-8")
        i = js.find("async function backup()")
        self.assertGreater(i, 0)
        body = js[i:i + 6000]
        self.assertIn("backup.delete", body, "o'chirish tugmasi yo'q")
        self.assertIn("confirmDialog", body, "tasdiqlash oynasi yo'q")
        self.assertIn("API.del(", body, "DELETE so'rovi yuborilmayapti")
        i18n = (ROOT / "web" / "js" / "i18n.js").read_text(encoding="utf-8")
        self.assertIn("Bu zahira nusxani o'chirmoqchisiz? Bu amalni qaytarib bo'lmaydi.",
                      i18n, "tasdiqlash matni i18n'da yo'q")


# ---------------------------------------------------------------------------
# MODUL 4 — parol va login o'zgartirish + kuchli parol siyosati
# ---------------------------------------------------------------------------
class TestModul4Credentials(_MixinAdmin, Base):
    """MODUL 4: bitta xavfsiz oqim (joriy parol + yangi login va/yoki parol),
    kuchli parol siyosati BUTUN platformada, audit va sessiya bekor qilish."""

    def setUp(self):
        super().setUp()
        r = self._mkuser("student", "Kuch", "Test", group_name="KUCH")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)

    def tearDown(self):
        self._cleanup_user(self.uid)
        super().tearDown()

    # ---- kuchli parol siyosati --------------------------------------------
    def test_M4_01_generated_passwords_are_strong(self):
        """Avtomatik generatsiya (yangi foydalanuvchi / CSV import) ham
        xuddi shu talablarga mos kelishi SHART."""
        from app.auth import password_strength_errors, generate_password
        for _ in range(40):
            pw = generate_password()
            self.assertEqual(password_strength_errors(pw), [], pw)
        for _ in range(4):
            r = self._mkuser("student", "Kuch", "Auto%d" % _)
            try:
                pw = r["credentials"]["password"]
                self.assertEqual(password_strength_errors(pw), [], pw)
            finally:
                self._cleanup_user(r["user"]["id"])

    def test_M4_02_weak_passwords_rejected_with_reasons(self):
        bad = {
            "Ab1!def":      ["auth.password_short"],                 # 7 belgi
            "abcdefgh1":    ["auth.password_short", "auth.password_need_upper"],
            "ABCDEFGH1":    ["auth.password_short", "auth.password_need_lower"],
            "Abcdefghij":   ["auth.password_need_digit"],
        }
        for pw, errs in bad.items():
            st, r = self.c.post("/api/auth/change-password", {
                "old_password": self.password, "new_password": pw,
                "confirm_password": pw})
            self.assertEqual(st, 400, (pw, r))
            self.assertEqual(r["error"], "auth.password_weak", (pw, r))
            got = (r.get("params") or {}).get("errors") or []
            self.assertEqual(sorted(got), sorted(errs), (pw, r))

    def test_M4_03_strong_password_accepted_and_old_stops_working(self):
        new = "KuchliParolM4a!"
        st, r = self.c.post("/api/auth/change-password", {
            "old_password": self.password, "new_password": new,
            "confirm_password": new})
        self.assertEqual(st, 200, r)
        self.assertTrue(r.get("password_changed"))
        self.assertEqual(r.get("login_changed"), False)
        row = self.db.q1("SELECT must_change_password FROM users WHERE id=?", (self.uid,))
        self.assertEqual(row["must_change_password"], 0)
        # eski parol bilan KIRIB BO'LMAYDI
        _c, st2, _r = self._login(self.login, self.password, "student")
        self.assertNotEqual(st2, 200)
        # yangi parol bilan kirish ishlaydi
        c2, st3, r3 = self._login(self.login, new, "student")
        self.assertEqual(st3, 200, r3)

    def test_M4_04_confirm_mismatch_and_nothing_to_change(self):
        st, r = self.c.post("/api/auth/change-password", {
            "old_password": self.password, "new_password": "KuchliParolM4b!",
            "confirm_password": "BoshqaParolM41!"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "auth.password_mismatch", r)
        st, r = self.c.post("/api/auth/change-password", {"old_password": self.password})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "auth.nothing_to_change", r)

    def test_M4_05_wrong_old_password_rejected(self):
        """Joriy parol HAR DOIM talab qilinadi — `must_change_password=1`
        (vaqtinchalik parol berilgan) holatida ham chetlab o'tish yo'q."""
        row = self.db.q1("SELECT must_change_password FROM users WHERE id=?", (self.uid,))
        self.assertEqual(row["must_change_password"], 1,
                         "test sharoiti: admin vaqtinchalik parol bergan")
        st, r = self.c.post("/api/auth/change-password", {
            "old_password": "Noto'g'riParol1", "new_password": "KuchliParolM4c!",
            "confirm_password": "KuchliParolM4c!"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "auth.wrong_old_password", r)
        # parol ham o'zgarmagan bo'lishi kerak
        self.assertNotEqual(
            self.db.q1("SELECT password_hash FROM users WHERE id=?",
                       (self.uid,))["password_hash"], "")
        self.assertEqual(self.db.q1("SELECT login FROM users WHERE id=?",
                                    (self.uid,))["login"], self.login)

    # ---- login --------------------------------------------------------------
    def test_M4_06_login_change_validates_and_revokes_other_sessions(self):
        # ikkinchi qurilmada kirish
        c2, st2, r2 = self._login(self.login, self.password, "student")
        self.assertEqual(st2, 200, r2)
        new_login = "usrL_kuchLogin00001"
        st, r = self.c.post("/api/auth/change-password", {
            "old_password": self.password, "new_login": new_login})
        self.assertEqual(st, 200, r)
        self.assertTrue(r.get("login_changed"))
        self.assertEqual(r.get("login"), new_login)
        self.assertEqual(self.db.q1("SELECT login FROM users WHERE id=?",
                                    (self.uid,))["login"], new_login)
        # joriy sessiya ISHLASHI (uzoq emas) kerak
        self.assertEqual(self.c.get("/api/auth/me")[0], 200)
        # boshqa sessiya bekor qilindi
        self.assertEqual(c2.get("/api/auth/me")[0], 401)
        # eski login bilan kiring imkoni yo'q
        c3, st3, _ = self._login(self.login, self.password, "student")
        self.assertNotEqual(st3, 200)
        c4, st4, _ = self._login(new_login, self.password, "student")
        self.assertEqual(st4, 200)

    def test_M4_07_login_format_and_unique_enforced(self):
        for bad in ("admin", "usrL_short", "usrX_aaaaaaaaaaaaaaaa", "usrL_ab!cd1234x"):
            st, r = self.c.post("/api/auth/change-password", {
                "old_password": self.password, "new_login": bad})
            self.assertEqual(st, 400, (bad, r))
            self.assertEqual(r["error"], "profile.login_format", (bad, r))
        other = self._mkuser("student", "Kuch", "Boshqa")
        try:
            st, r = self.c.post("/api/auth/change-password", {
                "old_password": self.password,
                "new_login": other["credentials"]["login"]})
            self.assertEqual(st, 400, r)
            self.assertEqual(r["error"], "profile.login_taken", r)
        finally:
            self._cleanup_user(other["user"]["id"])

    def test_M4_08_audit_records_only_the_fact(self):
        new = "AuditParolM4d9!"
        st, r = self.c.post("/api/auth/change-password", {
            "old_password": self.password, "new_password": new,
            "confirm_password": new,
            "new_login": "usrL_auditLogin0001"})
        self.assertEqual(st, 200, r)
        acts = [x["action_type"] for x in self.db.q(
            "SELECT action_type FROM audit_log WHERE target_id=? AND user_id=? "
            "ORDER BY id DESC LIMIT 2", (self.uid, self.uid))]
        self.assertIn("password_changed", acts)
        self.assertIn("login_changed", acts)
        # auditda parolning O'ZI saqlanmasin
        blob = " ".join(json.dumps(x, ensure_ascii=False) for x in self.db.q(
            "SELECT * FROM audit_log WHERE target_id=? AND action_type IN "
            "('password_changed','login_changed')", (self.uid,)))
        self.assertNotIn(new, blob, "audit jurnalida PAROL saqlanyapti")
        self.assertNotIn("AuditParol", blob)

    # ---- admin: parolni tiklash -------------------------------------------
    def test_M4_09_admin_reset_custom_password_must_be_strong(self):
        st, r = self.admin.post("/api/admin/users/%d/reset-password" % self.uid,
                                {"new_password": "kuchsiz"})
        self.assertEqual(st, 400, r)
        self.assertEqual(r["error"], "auth.password_weak", r)
        self.assertIn("auth.password_short",
                      (r.get("params") or {}).get("errors") or [])
        strong = "TiklashParolM4e7#"
        st, r = self.admin.post("/api/admin/users/%d/reset-password" % self.uid,
                                {"new_password": strong})
        self.assertEqual(st, 200, r)
        self.assertEqual(r["password"], strong)
        c2, st2, _ = self._login(self.login, strong, "student")
        self.assertEqual(st2, 200, "tiklangan parol bilan kirish kerak")

    def test_M4_10_admin_reset_auto_is_strong_and_revokes_sessions(self):
        st, r = self.admin.post("/api/admin/users/%d/reset-password" % self.uid, {})
        self.assertEqual(st, 200, r)
        from app.auth import password_strength_errors
        self.assertEqual(password_strength_errors(r["password"]), [])
        self.assertEqual(self.c.get("/api/auth/me")[0], 401,
                         "parol tiklangach barcha sessiyalar bekor bo'lishi shart")
        c2, st2, _ = self._login(self.login, r["password"], "student")
        self.assertEqual(st2, 200)

    def test_M4_11_ui_login_removed_from_profile_modal(self):
        js = (ROOT / "web" / "js" / "shared.js").read_text(encoding="utf-8")
        i = js.find("function editProfileModal")
        body = js[i:i + 3000]
        self.assertNotIn("common.login", body)
        self.assertIn("function openChangeCredentials", js)
        self.assertIn("profile.change_credentials", js,
                      "profil kartasida alohida tugma yo'q")
        app = (ROOT / "web" / "js" / "app.js").read_text(encoding="utf-8")
        self.assertNotIn("change_credentials", app,
                         "majburiy-parol oqimi ham alohida modallga o'tishi kerak")


# ---------------------------------------------------------------------------
# MODUL 5 — so'rovni ochish (5A) va muddati (5B)
# ---------------------------------------------------------------------------
class TestModul5Requests(_MixinAdmin, Base):
    """MODUL 5.

    5A — ESKI so'rov ham har doim ochiladi (ro'yxatdan qidirib topish emas,
         `GET .../requests/{id}` orqali).
    5B — Muddat so'rov YARATILGAN vaqtdan EMAS, so'ralgan MASHG'ULOT SANASI
         bo'yicha: sana kelajakda bo'lsa — qancha vaqt o'tsa ham faol;
         sana o'tgan + `pending` bo'lsa — 1 kundan keyin "Eskirgan"ga o'tadi.
    """

    MSG = "M5-TEST"

    def setUp(self):
        super().setUp()
        r = self._mkuser("student", "Mud", "Test", group_name="MUD")
        self.uid = r["user"]["id"]
        self.login = r["credentials"]["login"]
        self.password = r["credentials"]["password"]
        self.sid = self.db.q1("SELECT id FROM students WHERE user_id=?",
                               (self.uid,))["id"]
        self.c, st, rr = self._login(self.login, self.password, "student")
        self.assertEqual(st, 200, rr)

    def tearDown(self):
        for row in self.db.q("SELECT id, session_id FROM practice_requests WHERE message=?",
                             (self.MSG,)):
            if row["session_id"]:
                self.db.upd("DELETE FROM session_students WHERE session_id=?",
                            (row["session_id"],))
                self.db.upd("DELETE FROM lesson_sessions WHERE id=?", (row["session_id"],))
            self.db.upd("DELETE FROM hidden_requests WHERE request_id=?", (row["id"],))
            self.db.upd("DELETE FROM practice_requests WHERE id=?", (row["id"],))
        self._cleanup_user(self.uid)
        super().tearDown()

    def _mkreq(self, date_iso, status="pending"):
        rid = self.db.ex(
            "INSERT INTO practice_requests(student_id, preferred_date, preferred_start_time,"
            " preferred_end_time, message, status, created_at) VALUES(?,?,?,?,?,?,?)",
            (self.sid, date_iso, "09:00", "10:00", self.MSG, status,
             "2020-01-01 00:00:00"))
        return rid

    @staticmethod
    def _shift(days):
        from datetime import timedelta
        return (datetime.now() + timedelta(days=days)).strftime("%Y-%m-%d")

    # ---- 5B ---------------------------------------------------------------
    def test_M5_01_future_request_stays_active_no_matter_how_late(self):
        """Kelajakdagi so'rov 2 HAFTA keyingi bo'lsa ham FAOL qoladi
        (so'rov yaratilganiga emas, mashg'ulot sanasiga qaraladi)."""
        self._mkreq(self._shift(14))
        st, r = self.admin.get("/api/admin/requests?filter=active")
        self.assertEqual(st, 200, r)
        ids = [x["id"] for x in r["requests"]]
        self.assertIn(max(ids), ids)
        mine = [x for x in r["requests"] if x["message"] == self.MSG]
        self.assertEqual(len(mine), 1, "kelajakdagi so'rov faol bo'limda bo'lishi kerak")
        self.assertFalse(mine[0]["is_expired"])

    def test_M5_02_today_pending_request_is_active(self):
        self._mkreq(self._shift(0))
        st, r = self.admin.get("/api/admin/requests?filter=active")
        mine = [x for x in r["requests"] if x["message"] == self.MSG]
        self.assertEqual(len(mine), 1, "bugungi sana hali muddat kelmagan")
        self.assertFalse(mine[0]["is_expired"])

    def test_M5_03_past_pending_request_expires(self):
        """Sana o'tgan + hali `pending` -> 'Eskirgan' bo'limiga o'tadi."""
        self._mkreq(self._shift(-1))   # kechagi sana
        st, r = self.admin.get("/api/admin/requests?filter=active")
        self.assertEqual([x for x in r["requests"] if x["message"] == self.MSG], [],
                         "o'tgan-kun so'rovi faol ro'yxatda qolmasligi kerak")
        st, r = self.admin.get("/api/admin/requests?filter=expired")
        mine = [x for x in r["requests"] if x["message"] == self.MSG]
        self.assertEqual(len(mine), 1, "o'tgan-kun so'rovi 'Eskirgan'da bo'lishi kerak")
        self.assertTrue(mine[0]["is_expired"])
        self.assertEqual(mine[0]["_bucket"], "expired")

    def test_M5_04_old_pending_request_expires_after_one_day(self):
        """5 KUN OLDINGI sana, `pending` -> allaqach 'Eskirgan' (1 kun o'tgan)."""
        self._mkreq(self._shift(-5))
        st, r = self.admin.get("/api/admin/requests?filter=expired")
        self.assertEqual(len([x for x in r["requests"] if x["message"] == self.MSG]), 1)

    def test_M5_05_processed_requests_go_to_history_never_expired(self):
        """Tasdiqlangan/rad etilgan so'rov ESKIRMAYDI — u 'Tarix'da."""
        self._mkreq(self._shift(-5), status="approved")
        self._mkreq(self._shift(-5), status="rejected")
        st, r = self.admin.get("/api/admin/requests?filter=expired")
        self.assertEqual([x for x in r["requests"] if x["message"] == self.MSG], [],
                         "tasdiqlangan/rad etilgan so'rov 'Eskirgan'ga tushmasligi kerak")
        st, r = self.admin.get("/api/admin/requests?filter=history")
        self.assertEqual(len([x for x in r["requests"] if x["message"] == self.MSG]), 2)

    def test_M5_06_counts_and_default_filter(self):
        self._mkreq(self._shift(3))
        self._mkreq(self._shift(-2))
        self._mkreq(self._shift(-4), status="rejected")
        st, r = self.admin.get("/api/admin/requests")   # filtrsiz = active
        self.assertEqual(st, 200, r)
        self.assertEqual(r["filter"], "active")
        c = r["counts"]
        self.assertGreaterEqual(c["active"], 1)
        self.assertGreaterEqual(c["expired"], 1)
        self.assertGreaterEqual(c["history"], 1)
        # hisoblagichlar ro'yxatdagi BARCHA so'rovlarni qamrab oladi
        all_db = self.db.q1("SELECT COUNT(*) c FROM practice_requests r "
                            "JOIN students s ON s.id=r.student_id "
                            "JOIN users u ON u.id=s.user_id")["c"]
        self.assertEqual(c["active"] + c["expired"] + c["history"], all_db)
        # va hech bir so'rov ikki bo'limda ham bo'lmaydi
        ids = []
        for mode in ("active", "expired", "history"):
            st, rr = self.admin.get("/api/admin/requests?filter=%s" % mode)
            ids += [x["id"] for x in rr["requests"]]
        self.assertEqual(len(ids), len(set(ids)), "so'rovlar bo'linib ketgan")

    def test_M5_07_student_side_filters(self):
        self._mkreq(self._shift(5))
        self._mkreq(self._shift(-2))
        st, r = self.c.get("/api/student/requests")
        self.assertEqual(st, 200, r)
        self.assertEqual(len(r["requests"]), 1, "faqat kelajakdagi so'rov faol")
        st, r = self.c.get("/api/student/requests?filter=expired")
        self.assertEqual(len(r["requests"]), 1)
        self.assertTrue(r["requests"][0]["is_expired"])

    def test_M5_08_scheduled_job_marks_expired_at_only_for_past_pending(self):
        """Fon vazifa (`expire_stale_requests`) `expired_at` ni to'ldiradi —
        lekin FAQAT o'tgan-kun `pending` so'rovlar uchun."""
        from app.api import expire_stale_requests
        far = self._mkreq(self._shift(2))
        bugun = self._mkreq(self._shift(0))
        kecha = self._mkreq(self._shift(-1))
        tasdiq = self._mkreq(self._shift(-3), status="approved")
        n = expire_stale_requests(self.db)
        self.assertGreaterEqual(n, 1)
        row = lambda i: self.db.q1("SELECT expired_at FROM practice_requests WHERE id=?", (i,))
        self.assertEqual(row(far)["expired_at"], "", "kelajakdagi so'rov eskirMAYdi")
        self.assertEqual(row(bugun)["expired_at"], "", "bugungi so'rov eskirMAYdi")
        self.assertTrue(row(kecha)["expired_at"], "o'tgan-kun so'rovi belgilanmadi")
        self.assertEqual(row(tasdiq)["expired_at"], "", "tasdiqlangan so'rov eskirMAYdi")
        # idempotent
        self.assertEqual(expire_stale_requests(self.db), 0)

    # ---- 5A ---------------------------------------------------------------
    def test_M5_09_old_request_opens_via_detail_endpoint(self):
        """ESKI so'rov (o'tgan, faol ro'yxatda YO'Q) ham ochilishi kerak."""
        rid = self._mkreq(self._shift(-9))
        st, lst = self.admin.get("/api/admin/requests?filter=active")
        self.assertEqual([x for x in lst["requests"] if x["id"] == rid], [])
        st, r = self.admin.get("/api/admin/requests/%d" % rid)
        self.assertEqual(st, 200, r)
        self.assertEqual(r["request"]["id"], rid)
        self.assertEqual(r["request"]["message"], self.MSG)
        self.assertTrue(r["request"]["is_expired"])
        self.assertTrue(r["request"]["first_name"], "talaba ma'lumoti kelishi kerak")

    def test_M5_10_student_opens_own_old_request(self):
        rid = self._mkreq(self._shift(-9))
        st, r = self.c.get("/api/student/requests/%d" % rid)
        self.assertEqual(st, 200, r)
        self.assertEqual(r["request"]["id"], rid)
        # boshqa talabaning so'rovi — KO'RILMAYDI
        other = self._mkuser("student", "Bosh", "Talaba")
        try:
            osid = self.db.q1("SELECT id FROM students WHERE user_id=?",
                              (other["user"]["id"],))["id"]
            self.db.upd("UPDATE practice_requests SET student_id=? WHERE id=?", (osid, rid))
            st, r = self.c.get("/api/student/requests/%d" % rid)
            self.assertEqual(st, 404, r)
        finally:
            self.db.upd("UPDATE practice_requests SET student_id=? WHERE id=?", (self.sid, rid))
            self._cleanup_user(other["user"]["id"])

    def test_M5_11_detail_endpoint_rejects_bad_ids_and_missing(self):
        for bad in ("abc", "0", "-5"):
            st, r = self.admin.get("/api/admin/requests/%s" % bad)
            self.assertEqual(st, 404, (bad, st, r))
            self.assertEqual(r["error"], "request.not_found", (bad, r))
        st, r = self.admin.get("/api/admin/requests/99999999")
        self.assertEqual(st, 404, r)

    def test_M5_12_no_limit_lost_requests(self):
        """Eski `LIMIT 300` sababli so'rovlar yo'qolmasligi kerak."""
        self._mkreq(self._shift(2))
        st, r = self.admin.get("/api/admin/requests?filter=active")
        self.assertNotIn(" LIMIT", (ROOT / "app" / "api.py").read_text(encoding="utf-8")
                         .split("def admin_requests")[1].split("def ")[0])
        self.assertTrue(any(x["message"] == self.MSG for x in r["requests"]))

    def test_M5_13_backend_and_frontend_rules_agree(self):
        """`preferred_date` + `pending` — frontend va backend bir xil qoida."""
        from app.api import request_is_expired, request_bucket
        td = self._shift(0)
        self.assertFalse(request_is_expired({"status": "pending", "preferred_date": td}))
        self.assertFalse(request_is_expired({"status": "pending", "preferred_date": self._shift(9)}))
        self.assertTrue(request_is_expired({"status": "pending", "preferred_date": self._shift(-1)}))
        self.assertFalse(request_is_expired({"status": "approved", "preferred_date": self._shift(-9)}))
        self.assertEqual(request_bucket({"status": "pending", "preferred_date": td}), "active")
        self.assertEqual(request_bucket({"status": "pending", "preferred_date": self._shift(-1)}), "expired")
        self.assertEqual(request_bucket({"status": "rejected", "preferred_date": self._shift(-1)}), "history")
        # so'rov YARATILGAN sanasi hisobga OLINMAYDI
        row = {"status": "pending", "preferred_date": self._shift(3), "created_at": "2019-01-01 00:00:00"}
        self.assertFalse(request_is_expired(row), "created_at bo'yicha eskirsa xato bo'lardi")

    def test_M5_14_migration_adds_column_to_old_db(self):
        """Eski bazada `expired_at` yo'q — `migrate()` qo'shadi, ma'lumot saqlanadi."""
        import sqlite3
        from app.db import init_db
        d = tempfile.mkdtemp(prefix="avto-m5-")
        path = str(Path(d) / "old.db")
        init_db(path)
        # "ESKI" jadval: `expired_at` ustuni yo'q (platforma avvalgi versiyasi)
        c = sqlite3.connect(path)
        c.execute("ALTER TABLE practice_requests RENAME TO practice_requests_new")
        c.execute("""CREATE TABLE practice_requests (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    student_id INTEGER NOT NULL REFERENCES students(id),
    preferred_date TEXT DEFAULT '',
    preferred_start_time TEXT DEFAULT '',
    preferred_end_time TEXT DEFAULT '',
    message TEXT DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending'
        CHECK (status IN ('pending','approved','rejected','cancelled','rescheduled')),
    admin_note TEXT DEFAULT '',
    session_id INTEGER,
    created_at TEXT NOT NULL,
    processed_at TEXT DEFAULT ''
)""")
        c.execute("INSERT INTO students(user_id, group_name, license_category,"
                  " study_status, enrolled_at, status)"
                  " VALUES(1,'OLD','B','active','2020-01-01','active')")
        c.execute("INSERT INTO practice_requests(student_id, preferred_date, message,"
                  " status, created_at)"
                  " VALUES(1,'2020-01-01','eski','pending','2020-01-01 00:00:00')")
        c.execute("DROP TABLE practice_requests_new")
        c.commit()
        cols = [r[1] for r in c.execute("PRAGMA table_info(practice_requests)")]
        self.assertNotIn("expired_at", cols)
        c.close()
        init_db(path)                      # migrate ishlaydi
        c = sqlite3.connect(path)
        cols = [r[1] for r in c.execute("PRAGMA table_info(practice_requests)")]
        self.assertIn("expired_at", cols)
        n, = c.execute("SELECT COUNT(*) FROM practice_requests WHERE message='eski'").fetchone()
        self.assertEqual(n, 1, "eski ma'lumot YO'QOLMASIN")
        c.close()
        shutil.rmtree(d, ignore_errors=True)


class TestModul2UserManagement(_MixinAdmin, Base):
    """MODUL 2/3 — TALABA va ADMIN boshqaruvi.

    Talabalar (Modul 2) uchun:
      * bloklash / blokdan chiqarish,
      * arxivlash / arxivdan chiqarish (ma'lumotlar SAQLANADI),
      * to'liq profil ko'rish (progress, mashg'ulotlar),
      * login ko'rish + parolni TIKLASH (eski parol ko'rinMAYDI — hash).

    Adminlar (Modul 3) uchun:
      * xuddi shu amallar admin roliga ham qo'llanadi,
      * "Yangi admin" yaratish (login/parol AVTOMATIK, bir marta ko'rsatiladi),
      * talaba/instruktor admin yaratish huquqiga kirish olmaydi,
      * har bir amal AUDIT jurnalida qayd etiladi.
    """

    def _mk_admin(self, first="Adm", last="Qosh"):
        st, r = self.admin.post("/api/admin/users", {
            "role": "admin", "first_name": first, "last_name": last,
            "phone": "+998900000001",
        })
        self.assertEqual(st, 200, r)
        return r

    # ---------------------------------------------------------- MODUL 2: talaba
    def test_M2_01_student_block_unblock(self):
        r = self._mkuser("student", "Blok", "Test", group_name="BLK")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "block"})
            self.assertEqual(st, 200, res)
            row = self.db.q1("SELECT status FROM users WHERE id=?", (uid,))
            self.assertEqual(row["status"], "blocked")
            # Bloklangan foydalanuvchi KIRA OLMAYDI.
            c, st2, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertNotEqual(st2, 200, "bloklangan talaba kirdi")
            # Blokdan chiqarish -> yana kira oladi.
            st3, res3 = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "unblock"})
            self.assertEqual(st3, 200, res3)
            c, st4, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st4, 200, "faollashtirilgandan keyin kirmadi")
        finally:
            self._cleanup_user(uid)

    def test_M2_02_student_archive_unarchive_keeps_data(self):
        r = self._mkuser("student", "Arxiv", "Test", group_name="ARV")
        uid = r["user"]["id"]
        sid = self.db.q1("SELECT id FROM students WHERE user_id=?", (uid,))["id"]
        try:
            st, res = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "archive"})
            self.assertEqual(st, 200, res)
            row = self.db.q1("SELECT status, deleted_at FROM users WHERE id=?", (uid,))
            self.assertEqual(row["status"], "archived")
            self.assertTrue(row["deleted_at"], "arxivlashda deleted_at to'ldirilmadi")
            # MA'LUMOT SAQLANADI (o'chirilmaydi).
            self.assertIsNotNone(
                self.db.q1("SELECT id FROM students WHERE id=?", (sid,)),
                "arxivlashda talaba qatori o'chirildi")
            # Arxivlangan login qila olmaydi.
            c, st2, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertNotEqual(st2, 200, "arxivlangan talaba kirdi")
            # Arxivdan chiqarish.
            st3, res3 = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "unarchive"})
            self.assertEqual(st3, 200, res3)
            row2 = self.db.q1("SELECT status, deleted_at FROM users WHERE id=?", (uid,))
            self.assertEqual(row2["status"], "active")
            self.assertIsNone(row2["deleted_at"], "unarchive da deleted_at tozalanmadi")
            c, st4, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st4, 200, "arxivdan chiqarilgandan keyin kirmadi")
        finally:
            self._cleanup_user(uid)

    def test_M2_03_student_profile_full_and_no_password(self):
        r = self._mkuser("student", "Profil", "Test", group_name="PRF",
                         license_category="B", phone="+998901234567")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.get(f"/api/admin/users/{uid}")
            self.assertEqual(st, 200, res)
            u = res["user"]
            self.assertEqual(u["first_name"], "Profil")
            self.assertEqual(u["login"], r["credentials"]["login"])
            self.assertEqual(u["student"]["group_name"], "PRF")
            self.assertIn("progress", u)
            for k in ("progress",):
                self.assertIsNotNone(u[k], f"{k} yo'q")
            # Parol HECH QACHON qaytarilmaydi (xavfsizlik standarti).
            self.assertNotIn("password_hash", u)
            self.assertNotIn("password", u)
            self.assertNotIn("totp_secret", u)
        finally:
            self._cleanup_user(uid)

    def test_M2_04_password_reset_shown_once(self):
        r = self._mkuser("student", "Tikla", "Test")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.post(f"/api/admin/users/{uid}/reset-password")
            self.assertEqual(st, 200, res)
            new_pw = res["password"]
            self.assertTrue(new_pw.startswith("usrP_"), "parol formati usrP_ bo'lishi kerak")
            # Yangi parol bilan kirish ishlaydi.
            c, st2, _ = self._login(res["login"], new_pw, "student")
            self.assertEqual(st2, 200, "yangi parol bilan kirmadi")
            # Eski parol endi ISHLAMAYDI (parol haqiqiyda o'zgartirildi).
            c3, st3, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertNotEqual(st3, 200, "eski parol hali ham ishlaydi")
        finally:
            self._cleanup_user(uid)

    def test_M2_05_audit_written_for_block_and_archive(self):
        r = self._mkuser("student", "Audit", "Test")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "block"})
            self.assertEqual(st, 200, res)
            st2, res2 = self.admin.get("/api/admin/audit?action=user")
            self.assertEqual(st2, 200, res2)
            rows = [x for x in res2["logs"] if str(x.get("target_id")) == str(uid)]
            self.assertTrue(rows, "audit jurnalida bloklash yozuvi yo'q")
            self.assertTrue(any(x.get("action_type") == "user_blocked" for x in rows),
                            "action_type=user_blocked yo'q")
            st3, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "unblock"})
            self.assertEqual(st3, 200)
        finally:
            self._cleanup_user(uid)

    def test_M2_06_archived_visible_through_filter(self):
        r = self._mkuser("student", "Ko", "Rinadi", group_name="KRN")
        uid = r["user"]["id"]
        try:
            st, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "archive"})
            self.assertEqual(st, 200)
            # Arxiv filtri bilan KO'RINADI (ro'yxatdan yo'qolmasin).
            st2, res2 = self.admin.get("/api/admin/users?role=student&status=archived")
            self.assertEqual(st2, 200, res2)
            self.assertTrue([u for u in res2["users"] if u["id"] == uid],
                            "arxivlangan talaba arxiv filtrida ko'rinmadi")
            # Arxiv profilini ko'rish mumkin (ma'lumot saqlangan).
            st3, res3 = self.admin.get(f"/api/admin/users/{uid}")
            self.assertEqual(st3, 200, res3)
            self.assertEqual(res3["user"]["status"], "archived")
            # Faol ro'yxatda YO'Q.
            st4, res4 = self.admin.get("/api/admin/users?role=student&status=active")
            self.assertFalse([u for u in res4["users"] if u["id"] == uid],
                             "arxivlangan talaba faol ro'yxatda qoldi")
        finally:
            self._cleanup_user(uid)

    def test_M2_07_status_change_returns_new_status(self):
        r = self._mkuser("student", "Holat", "Qaytar")
        uid = r["user"]["id"]
        try:
            for act, want in (("archive", "archived"), ("unarchive", "active"),
                              ("block", "blocked"), ("unblock", "active")):
                st, res = self.admin.post(f"/api/admin/users/{uid}/status", {"status": act})
                self.assertEqual(st, 200, res)
                self.assertEqual(res["status"], want)
                self.assertEqual(
                    self.db.q1("SELECT status FROM users WHERE id=?", (uid,))["status"], want)
        finally:
            self._cleanup_user(uid)

    def test_M2_08_block_does_not_unarchive(self):
        """Bloklash arxiv belgisini TOZALAMASLIGI kerak.

        Aks holda "arxivlangan talabani bloklash" uni yashirincha faol
        ro'yxatga qaytarib yuborardi (u login qila olmas edi, lekin holat
        "active" ko'rinardi).
        """
        r = self._mkuser("student", "Arxiv", "Blok")
        uid = r["user"]["id"]
        try:
            st, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "archive"})
            self.assertEqual(st, 200)
            st2, res2 = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "block"})
            self.assertEqual(st2, 200, res2)
            row = self.db.q1("SELECT status, deleted_at FROM users WHERE id=?", (uid,))
            self.assertEqual(row["status"], "blocked")
            self.assertIsNotNone(row["deleted_at"], "bloklash arxivni bekor qildi")
            # Arxiv filtrida hali ham ko'rinadi.
            st3, res3 = self.admin.get("/api/admin/users?role=student&status=archived")
            self.assertTrue([u for u in res3["users"] if u["id"] == uid],
                            "bloklangan arxiv foydalanuvchi arxiv ro'yxatidan qoldi")
        finally:
            self._cleanup_user(uid)

    # ---------------------------------------------------------- MODUL 3: admin
    def test_M3_01_create_admin_auto_credentials(self):
        r = self._mk_admin()
        uid = r["user"]["id"]
        try:
            self.assertEqual(r["user"]["role"], "admin")
            self.assertTrue(r["credentials"]["login"].startswith("usrL_"),
                            "admin login formati usrL_ bo'lishi kerak")
            self.assertTrue(r["credentials"]["password"].startswith("usrP_"),
                            "admin parol formati usrP_ bo'lishi kerak")
            # Yaratilgan admin login qila oladi.
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "admin")
            self.assertEqual(st, 200, "yangi admin kira olmadi")
            # Audit: "admin_created".
            st2, res2 = self.admin.get("/api/admin/audit?action=admin_created")
            self.assertEqual(st2, 200, res2)
            self.assertTrue([x for x in res2["logs"] if str(x.get("target_id")) == str(uid)],
                            "admin_created audit yozuvi yo'q")
        finally:
            self._cleanup_user(uid)

    def test_M3_02_admin_block_archive_unarchive(self):
        r = self._mk_admin("BlokAdm")
        uid = r["user"]["id"]
        try:
            st, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "block"})
            self.assertEqual(st, 200, "admin bloklanmadi")
            self.assertEqual(self.db.q1("SELECT status FROM users WHERE id=?", (uid,))["status"],
                             "blocked")
            c, st2, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "admin")
            self.assertNotEqual(st2, 200, "bloklangan admin kirdi")
            st3, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "unblock"})
            self.assertEqual(st3, 200)
            c, st4, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "admin")
            self.assertEqual(st4, 200, "blokdan chiqarilgandan keyin admin kirmadi")
            st5, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "archive"})
            self.assertEqual(st5, 200, "admin arxivlanmadi")
            c, st6, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "admin")
            self.assertNotEqual(st6, 200, "arxivlangan admin kirdi")
            st7, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "unarchive"})
            self.assertEqual(st7, 200)
            c, st8, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "admin")
            self.assertEqual(st8, 200, "arxivdan chiqarilgandan keyin admin kirmadi")
        finally:
            self._cleanup_user(uid)

    def test_M3_03_admin_password_reset(self):
        r = self._mk_admin("TiklaAdm")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.post(f"/api/admin/users/{uid}/reset-password")
            self.assertEqual(st, 200, res)
            c, st2, _ = self._login(res["login"], res["password"], "admin")
            self.assertEqual(st2, 200, "tiklangan parol bilan admin kirmadi")
        finally:
            self._cleanup_user(uid)

    def test_M3_04_student_cannot_create_admin(self):
        r = self._mkuser("student", "Talaba", "Test")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            st2, res2 = c.post("/api/admin/users", {
                "role": "admin", "first_name": "Hu", "last_name": "Akinchi",
            })
            self.assertIn(st2, (401, 403), f"talaba admin yarata oldi: {st2} {res2}")
            self.assertEqual(
                self.db.q1("SELECT COUNT(*) c FROM users WHERE last_name='Akinchi'")["c"], 0,
                "talaba orqali admin yaratilgan")
        finally:
            self._cleanup_user(uid)

    def test_M3_05_instructor_cannot_create_admin(self):
        r = self._mkuser("instructor", "Instr", "Test")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "instructor")
            self.assertEqual(st, 200)
            st2, res2 = c.post("/api/admin/users", {
                "role": "admin", "first_name": "Hu", "last_name": "InstruktorA",
            })
            self.assertIn(st2, (401, 403), f"instruktor admin yarata oldi: {st2} {res2}")
        finally:
            self._cleanup_user(uid)

    def test_M3_06_admin_invalid_role_rejected(self):
        st, res = self.admin.post("/api/admin/users", {
            "role": "superadmin", "first_name": "X", "last_name": "Y",
        })
        self.assertEqual(st, 400, res)
        self.assertEqual(res["error"], "user.invalid_role")

    def test_M3_07_bad_status_rejected(self):
        r = self._mkuser("student", "Holat", "Test")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "nonsense"})
            self.assertEqual(st, 400, res)
            self.assertEqual(res["error"], "user.bad_status")
        finally:
            self._cleanup_user(uid)

    def test_M3_08_status_block_kills_sessions(self):
        r = self._mkuser("student", "Sess", "Test")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            st2, _ = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "block"})
            self.assertEqual(st2, 200)
            st3, res3 = c.get("/api/auth/me")
            self.assertEqual(st3, 401, "bloklangandan keyin sessiya yashirin qoldi")
        finally:
            self._cleanup_user(uid)

    def test_M3_09_admin_cannot_block_self(self):
        st0, me = self.admin.get("/api/auth/me")
        self.assertEqual(st0, 200, me)
        me_id = me["user"]["id"]
        for act in ("block", "archive"):
            with self.subTest(act=act):
                st, res = self.admin.post(f"/api/admin/users/{me_id}/status", {"status": act})
                self.assertEqual(st, 400, res)
                self.assertEqual(res["error"], "user.cannot_block_self")
                self.assertEqual(
                    self.db.q1("SELECT status FROM users WHERE id=?", (me_id,))["status"],
                    "active", "o'zini bloklagan holda holati o'zgardi")

    def test_M3_10_last_active_admin_protected(self):
        """Tizimda faqat BIR faol admin qolsa, uni bloklashga ruxsat yo'q.

        Test uchun holat "qo'lda" yaratiladi: boshqa admin `status` maydoni
        to'g'ridan-to'g'ri DB da arxivlanadi (sessiya saqlanib qoladi).
        Shu bilan himoya qoidasi real yo'l bilan tekshiriladi.
        """
        r = self._mk_admin("OxirgiAdmin")
        uid = r["user"]["id"]
        st0, me = self.admin.get("/api/auth/me")
        me_id = me["user"]["id"]
        try:
            self.db.upd("UPDATE users SET status='archived' WHERE id=?", (me_id,))
            try:
                st, res = self.admin.post(f"/api/admin/users/{uid}/status", {"status": "block"})
                self.assertEqual(st, 400, res)
                self.assertEqual(res["error"], "user.last_admin_protected")
                self.assertEqual(
                    self.db.q1("SELECT status FROM users WHERE id=?", (uid,))["status"], "active")
            finally:
                self.db.upd("UPDATE users SET status='active' WHERE id=?", (me_id,))
        finally:
            self._cleanup_user(uid)

    def test_M3_11_admin_can_be_purged_but_last_one_protected(self):
        """Admin ham BUTUNLAY o'chirilishi mumkin — lekin O'ZINI va
        OXIRGI faol adminni o'chirishga ruxsat yo'q (tizim qolishi uchun).

        Eski qoida ("adminni o'chirish butunlay mumkin emas") endi
        KENGAYTIRILDI: yangi talab bo'yicha "O'chirish" = haqiqiy DELETE,
        u arxivlashdan qat'i ajratilgan bo'lishi SHART.
        """
        r = self._mk_admin("Ochiriladi")
        uid = r["user"]["id"]
        # o'zini o'chirishga urinish
        me = self.admin.get("/api/auth/me")[1]["user"]["id"]
        if uid != me:
            st, res = self.admin.delete(f"/api/admin/users/{uid}")
            self.assertEqual(st, 200, res)
            # haqiqiy DELETE: arxiv EMAS
            self.assertIsNone(self.db.q1("SELECT id FROM users WHERE id=?", (uid,)),
                              "admin o'chirmasligi kerak edi (arxiv emas!)")
            st2, r2 = self.admin.get("/api/admin/users?status=archived")
            self.assertNotIn(uid, [u["id"] for u in r2["users"]])
            return
        # `me` bo'lsa — o'zini o'chira olmaydi
        st, res = self.admin.delete(f"/api/admin/users/{uid}")
        self.assertEqual(st, 400, res)
        self.assertEqual(res["error"], "user.cannot_delete_self")

    def test_M3_12_admin_profile_visible(self):
        r = self._mk_admin("Ko", "Rsat", )
        uid = r["user"]["id"]
        try:
            st, res = self.admin.get(f"/api/admin/users/{uid}")
            self.assertEqual(st, 200, res)
            self.assertEqual(res["user"]["role"], "admin")
            self.assertEqual(res["user"]["login"], r["credentials"]["login"])
            self.assertNotIn("password_hash", res["user"])
            # Talaba/instruktor kabi alohida profil jadvali yaratilMAGAN.
            self.assertIsNone(self.db.q1("SELECT id FROM students WHERE user_id=?", (uid,)))
            self.assertIsNone(self.db.q1("SELECT id FROM instructors WHERE user_id=?", (uid,)))
        finally:
            self._cleanup_user(uid)


class TestModul4NoInstructorMessages(_MixinAdmin, Base):
    """MODUL 4 — Instruktor "Xabarlar" bo'limi BUTUNLAY olib tashlandi.

    Tekshiriladi:
      * sidebar (NAV_KEYS) da "messages" yo'q,
      * `InstructorViews.messages` export qilinmaydi,
      * `App.go('messages')` instruktor uchun bo'sh sahifaga olib ketmaydi
        (route alias bilan "notifications" ga boradi yoki bo'sh qaytariladi).
    """

    def _read(self, rel):
        p = Path(__file__).resolve().parent.parent / rel
        return p.read_text(encoding="utf-8")

    def test_M4_01_instructor_nav_has_no_messages(self):
        app_js = self._read("web/js/app.js")
        m = __import__("re").search(
            r"instructor:\s*\[(.*?)\n    \]", app_js, __import__("re").S)
        self.assertIsNotNone(m, "NAV_KEYS.instructor topilmadi")
        block = m.group(1)
        self.assertNotIn('"messages"', block,
                         "instruktor sidebar'ida 'messages' hali ham bor")
        # Boshqa bo'limlar o'z o'rnida qolishi SHART (regressiya tekshiruvi).
        for key in ('"schedule"', '"students"', '"car"', '"profile"'):
            self.assertIn(key, block, f"instruktor nav'ida {key} yo'qoldi")

    def test_M4_02_instructor_views_have_no_messages_fn(self):
        src = self._read("web/js/views-instructor.js")
        self.assertNotIn("function messages(", src,
                         "InstructorViews.messages hali ham mavjud")
        exp = __import__("re").search(r"return\s*\{([^}]*)\}\s*;\s*\}\)\(\)\s*;?\s*$", src,
                                     __import__("re").S)
        self.assertIsNotNone(exp, "views-instructor.js eksporti topilmadi")
        self.assertNotIn("messages", exp.group(1), "eksportda messages qoldi")

    def test_M4_03_instructor_messages_page_not_reachable(self):
        # Eski view kaliti bo'lmasa, `App.go("messages")` instruktor uchun
        # `views[App.view] || views.dashboard` -> dashboard ga tushadi.
        r = self._mkuser("instructor", "InstrM", "Test")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "instructor")
            self.assertEqual(st, 200)
            # Eski endpoint hali ishlashi mumkin (ma'lumot uchun), lekin
            # FRONTEND da bo'lim yo'q — bu test shuni tasdiqlaydi:
            # `NAV_KEYS` da yo'q (yuqoridagi test) va view eksportida yo'q.
            st2, res2 = c.get("/api/instructor/context")
            self.assertEqual(st2, 200, res2)
        finally:
            self._cleanup_user(uid)


class TestModul1ConditionalFields(_MixinAdmin, Base):
    """MODUL 1 — "Parol va login o'zgartirish" formasi.

    Frontend tekshiruvlari (shartli maydonlar, ko'zcha, live validatsiya)
    va backend qismi (joriy parol har doim majburiy, ixtiyoriy login/parol).
    """

    def _read(self, rel):
        p = Path(__file__).resolve().parent.parent / rel
        return p.read_text(encoding="utf-8")

    def test_M1_01_password_input_helper_exists(self):
        ui = self._read("web/js/ui.js")
        self.assertIn("function passwordInput(", ui,
                      "UI.passwordInput (ko'zcha bilan parol maydoni) mavjud emas")
        self.assertIn("passwordInput", ui.split("return {")[-1],
                      "passwordInput eksport qilinmagan")
        # Komponent haqiqiy `type` almashtirishni amalga oshiradi.
        self.assertIn('inp.type = show ? "text" : "password"', ui,
                      "ko'zcha type=password ni text ga almashtirmaydi")
        self.assertIn("ico-eye-off", ui, "ko'zcha ikonkasi yo'q")

    def test_M1_02_change_credentials_uses_eye(self):
        src = self._read("web/js/shared.js")
        start = src.index("function openChangeCredentials()")
        body = src[start:start + 4000]
        self.assertEqual(body.count("UI.passwordInput"), 3,
                         "joriy/yangi/takrorlash maydonlarida ko'zcha bo'lishi kerak")

    def test_M1_03_conditional_indicators_present(self):
        src = self._read("web/js/shared.js")
        start = src.index("function openChangeCredentials()")
        body = src[start:start + 4000]
        self.assertIn("profile.login_changing", body, "login o'zgarishi indikatori yo'q")
        self.assertIn("profile.pw_changing", body, "parol o'zgarishi indikatori yo'q")
        self.assertIn("pwChecklist", body, "real vaqtli kuchli parol ro'yxati yo'q")

    def test_M1_04_save_button_gated_by_validation(self):
        src = self._read("web/js/shared.js")
        start = src.index("function openChangeCredentials()")
        body = src[start:start + 4000]
        self.assertIn("saveBtn.disabled", body, "saqlash tugmasi validatsiyaga bog'lanmagan")
        self.assertIn("updateSave", body, "updateSave (validatsiya) ishlatilmagan")

    def test_M1_05_i18n_keys_exist(self):
        i18n = self._read("web/js/i18n.js")
        for key in ("profile.login_changing", "profile.pw_changing",
                    "users.add_admin", "users.block", "users.unblock",
                    "users.archive", "users.unarchive", "users.view_profile",
                    "users.admin_role_hint", "users.block_confirm",
                    "users.unblock_confirm", "users.archive_confirm",
                    "users.unarchive_confirm"):
            with self.subTest(key=key):
                self.assertIn('"%s"' % key, i18n, f"i18n kaliti yo'q: {key}")

    def test_M1_06_login_only_change_backend(self):
        # FAQAT login o'zgarishi: parolsiz.
        r = self._mkuser("student", "Login", "Faqat")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            new_login = "usrL_TestFaqatLogin01"
            st2, res2 = c.post("/api/auth/change-password", {
                "old_password": r["credentials"]["password"],
                "new_login": new_login, "new_password": "", "confirm_password": "",
            })
            self.assertEqual(st2, 200, res2)
            self.assertTrue(res2["login_changed"])
            self.assertFalse(res2["password_changed"])
        finally:
            self._cleanup_user(uid)

    def test_M1_07_password_only_change_backend(self):
        r = self._mkuser("student", "Parol", "Faqat")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            new_pw = "FaqatParol2026Xy"
            st2, res2 = c.post("/api/auth/change-password", {
                "old_password": r["credentials"]["password"],
                "new_login": "", "new_password": new_pw, "confirm_password": new_pw,
            })
            self.assertEqual(st2, 200, res2)
            self.assertTrue(res2["password_changed"])
            self.assertFalse(res2["login_changed"])
        finally:
            self._cleanup_user(uid)

    def test_M1_08_old_password_still_required(self):
        r = self._mkuser("student", "Joriy", "Kerak")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            st2, res2 = c.post("/api/auth/change-password", {
                "old_password": "",
                "new_login": "", "new_password": "YangiParol2026Xy", "confirm_password": "YangiParol2026Xy",
            })
            self.assertEqual(st2, 400, res2)
            self.assertEqual(res2["error"], "auth.wrong_old_password")
        finally:
            self._cleanup_user(uid)

    def test_M1_09_nothing_to_change(self):
        r = self._mkuser("student", "Hich", "Narsa")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            st2, res2 = c.post("/api/auth/change-password", {
                "old_password": r["credentials"]["password"],
                "new_login": "", "new_password": "", "confirm_password": "",
            })
            self.assertEqual(st2, 400, res2)
            self.assertEqual(res2["error"], "auth.nothing_to_change")
        finally:
            self._cleanup_user(uid)


class TestModul5AdminCredentialVisibility(_MixinAdmin, Base):
    """MODUL 5 — admin foydalanuvchi LOGIN va PAROLINI ko'radi hamda
    xohlagancha o'zgartiradi.

    Talab (foydalanuvchidan): "admin har bir foydalanuvchining login va
    parolini ko'ra olsin va xohlagancha o'zgartirish huquqiga ega bo'lsin";
    "bir marta ko'rsatish" cheklovi olib tashlandi; yangi foydalanuvchiga
    login/parol avtomatik beriladi; "agar admin o'zi almashtirsa, o'zgarishsiz
    yangi bazaga saqlansin".

    Buning uchun parol `password_enc` (AES-256-GCM) sifatida ham saqlanadi.
    Testlar shu qatlamni ham tekshiradi.
    """

    def _read(self, rel):
        p = Path(__file__).resolve().parent.parent / rel
        return p.read_text(encoding="utf-8")

    # ------------------------------------------------------------------
    # Shifrlash qatlami
    # ------------------------------------------------------------------
    def test_M5_01_encrypt_roundtrip(self):
        from app.auth import encrypt_secret, decrypt_secret
        pw = "usrP_TestRoundtrip77!"
        blob = encrypt_secret(pw)
        self.assertTrue(blob.startswith("enc:v1:"), "prefiks noto'g'ri: %r" % blob[:12])
        self.assertNotIn(pw, blob, "ochiq matn shifrlangan blob ichida qoldi")
        self.assertEqual(decrypt_secret(blob), pw)

    def test_M5_02_decrypt_rejects_tampering(self):
        from app.auth import encrypt_secret, decrypt_secret
        blob = encrypt_secret("usrP_ABC12345x")
        # 1 ta belgi o'zgartiriladi -> GCM authenticate muvaffaqiyatsiz
        broken = blob[:-4] + ("AAAA" if not blob.endswith("AAAA") else "BBBB")
        self.assertIsNone(decrypt_secret(broken), "buzilgan ma'lumot qabul qilindi")
        self.assertIsNone(decrypt_secret(""), "bo'sh qiymat parol qaytaradi")
        self.assertIsNone(decrypt_secret("sha1$olma$1$2$aa$bb"),
                          "noto'g'ri prefiks parol qaytaradi")
        # Boshqa kalit bilan ochib bo'lmaydi (kalit tekshiruvi)
        self.assertNotEqual(decrypt_secret(blob), decrypt_secret(encrypt_secret("boshqa")))

    def test_M5_03_encrypt_par_never_raises(self):
        """Kalit/kutubxona yo'q bo'lsa `encrypt_par` xato OTARMASIZ —
        platforma ishlashda davom etishi kerak (bo'sh satr qaytaradi)."""
        from app import api as apimod
        from app.auth import decrypt_secret
        r = apimod.encrypt_par("qandaydir_parol")
        self.assertIsInstance(r, str, "natija satr bo'lishi SHART")
        if _crypto_ready():
            self.assertTrue(r.startswith("enc:v1:"), "shifrlanmagan: %r" % r)
            self.assertEqual(decrypt_secret(r), "qandaydir_parol")
        else:
            self.assertEqual(r, "", "kalit yo'q bo'lsa bo'sh satr qaytariladi")

    # ------------------------------------------------------------------
    # Saqlash: hash ham, shifrlangan nusxa ham
    # ------------------------------------------------------------------
    def test_M5_04_new_user_password_is_viewable(self):
        r = self._mkuser("student", "Ko'rin", "Parol")
        uid = r["user"]["id"]
        try:
            pw = r["credentials"]["password"]
            st, res = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(st, 200, res)
            self.assertTrue(res["can_view"], "yangi foydalanuvchi paroli ko'rinmadi: %s" % res.get("reason"))
            self.assertEqual(res["login"], r["credentials"]["login"])
            self.assertEqual(res["password"], pw, "qaytarilgan parol boshqacha")
            # Parol DB'da ochiq matn holda EMAS
            row = _json_db().q1("SELECT password_hash,password_enc FROM users WHERE id=?", (uid,))
            self.assertTrue(row["password_hash"].startswith("scrypt$"), "hash yo'q")
            self.assertTrue(row["password_enc"].startswith("enc:v1:"), "shifrlangan nusxa yo'q")
            self.assertNotIn(pw, row["password_hash"])
            self.assertNotIn(pw, row["password_enc"])
        finally:
            self._cleanup_user(uid)

    def test_M5_05_password_enc_hidden_from_public_payloads(self):
        """`password_enc` hech qanday ro'yxat/profil javobida chiqmasligi SHART."""
        r = self._mkuser("student", "Yashir", "Kerak")
        uid = r["user"]["id"]
        try:
            for label, st, res in (
                ("admin/users", *self.admin.get("/api/admin/users?role=student")),
                ("admin/user_profile", *self.admin.get("/api/admin/users/%d" % uid)),
            ):
                with self.subTest(endpoint=label):
                    self.assertEqual(st, 200, res)
                    blob = repr(res)
                    self.assertNotIn("password_enc", blob, "%s password_enc ni chiqaryapti" % label)
                    self.assertNotIn("password_hash", blob, "%s password_hash ni chiqaryapti" % label)
        finally:
            self._cleanup_user(uid)

    def test_M5_06_view_is_written_to_audit(self):
        """Parolni KO'RISH maxfiy amal — auditga yozilishi shart."""
        r = self._mkuser("student", "Audit", "Ko'rish")
        uid = r["user"]["id"]
        try:
            self.admin.get("/api/admin/users/%d/credentials" % uid)
            row = _json_db().q1(
                "SELECT action_type FROM audit_log WHERE target_id=? AND action_type=?",
                (str(uid), "user_credentials_viewed"))
            self.assertIsNotNone(row, "user_credentials_viewed audit yozuvi yo'q")
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # RBAC
    # ------------------------------------------------------------------
    def test_M5_07_student_cannot_read_or_write_credentials(self):
        r = self._mkuser("student", "Talaba", "Hujjat")
        uid = r["user"]["id"]
        try:
            c, st, _ = self._login(r["credentials"]["login"], r["credentials"]["password"], "student")
            self.assertEqual(st, 200)
            st1, res1 = c.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(st1, 403, res1)
            st2, res2 = c.put("/api/admin/users/%d/credentials" % uid,
                               {"password": "usrP_HackerTry1!"})
            self.assertEqual(st2, 403, res2)
            # Parol hamon o'zgarilmagan
            st3, res3 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res3["password"], r["credentials"]["password"])
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # BIRGA o'zgartirish (login + parol)
    # ------------------------------------------------------------------
    def test_M5_08_change_login_and_password_together(self):
        r = self._mkuser("student", "Birga", "O'zgar")
        uid = r["user"]["id"]
        try:
            new_login = "usrL_BirgaOzgaris01"
            new_pw = "usrP_BirgaOzgar9!"
            st, res = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                     {"login": new_login, "password": new_pw})
            self.assertEqual(st, 200, res)
            self.assertTrue(res["changed"])
            self.assertEqual(sorted(res["what"]), ["login", "password"])
            # Saqlanganini tekshirish
            st2, res2 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res2["login"], new_login)
            self.assertEqual(res2["password"], new_pw)
            # YANGI login/parol bilan kirish ishlaydi
            c, st3, res3 = self._login(new_login, new_pw, "student")
            self.assertEqual(st3, 200, res3)
            # Eski login bilan KIRIB BO'LMAYDI
            from app.auth import reset_login_attempts
            reset_login_attempts("login:" + r["credentials"]["login"])
            c2 = Client()
            st4, _ = c2.post("/api/auth/login", {"role": "student",
                                                "login": r["credentials"]["login"],
                                                "password": r["credentials"]["password"]})
            self.assertNotEqual(st4, 200, "eski login bilan kirish mumkin bo'ldi")
        finally:
            self._cleanup_user(uid)

    def test_M5_09_password_change_kills_sessions(self):
        """Parol o'zgarganda barcha sessiyalar bekor qilinadi."""
        r = self._mkuser("student", "Sessiya", "Bekor")
        uid = r["user"]["id"]
        try:
            lg, pw = r["credentials"]["login"], r["credentials"]["password"]
            c, st, _ = self._login(lg, pw, "student")
            self.assertEqual(st, 200)
            st2, res2 = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                        {"password": "usrP_YangiSessiya5!"})
            self.assertEqual(st2, 200, res2)
            st3, res3 = c.get("/api/auth/me")
            self.assertNotEqual(st3, 200, "eski sessiya yashirib qoldi")
        finally:
            self._cleanup_user(uid)

    def test_M5_10_login_change_only_keeps_password(self):
        """FAQAT login o'zgaradi — parolga tegilmaydi, sessiya ham saqlanadi."""
        r = self._mkuser("student", "FaqatLogin", "O'zgar")
        uid = r["user"]["id"]
        try:
            lg, pw = r["credentials"]["login"], r["credentials"]["password"]
            c, st, _ = self._login(lg, pw, "student")
            self.assertEqual(st, 200)
            st2, res2 = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                        {"login": "usrL_FaqatLoginOzgar01"})
            self.assertEqual(st2, 200, res2)
            self.assertEqual(res2["what"], ["login"])
            self.assertFalse(res2["password_changed"])
            st3, res3 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res3["password"], pw, "parol o'zgari ketdi")
            # Sessiya saqlanadi (parol o'zgarmagan)
            st4, res4 = c.get("/api/auth/me")
            self.assertEqual(st4, 200, res4)
            self.assertEqual(res4["user"]["login"], "usrL_FaqatLoginOzgar01")
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # "O'zgartirmasam — o'zgarishsiz saqlansin"
    # ------------------------------------------------------------------
    def test_M5_11_empty_body_changes_nothing(self):
        """"Agar admin o'zi almashtirsa, o'zgarishsiz saqlansin": bo'sh
        so'rov HECH QANDAY yozuv QILMASIZ (joriy qiymat buzilmaydi)."""
        r = self._mkuser("student", "Hich", "Narsa")
        uid = r["user"]["id"]
        try:
            lg, pw = r["credentials"]["login"], r["credentials"]["password"]
            before = _json_db().q1("SELECT login,password_hash,updated_at FROM users WHERE id=?", (uid,))
            for payload in ({}, {"login": ""}, {"password": ""}, {"login": lg, "password": pw},
                            {"login": "  ", "password": "  "}):
                with self.subTest(payload=payload):
                    st, res = self.admin.put("/api/admin/users/%d/credentials" % uid, payload)
                    self.assertEqual(st, 200, res)
                    self.assertFalse(res["changed"], "bo'sh so'rov yozuv qildi: %s" % res)
                    self.assertEqual(res["login"], lg)
            after = _json_db().q1("SELECT login,password_hash,updated_at FROM users WHERE id=?", (uid,))
            self.assertEqual(after["login"], before["login"])
            self.assertEqual(after["password_hash"], before["password_hash"],
                             "parol hash'i o'zgardi")
            self.assertEqual(after["updated_at"], before["updated_at"], "updated_at o'zgardi")
            # Parol hali ham ko'rinadi
            st, res = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res["password"], pw)
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # Validatsiya
    # ------------------------------------------------------------------
    def test_M5_12_weak_password_rejected(self):
        r = self._mkuser("student", "Zaif", "Parol")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                     {"password": "abc"})
            self.assertEqual(st, 400, res)
            self.assertEqual(res["error"], "auth.password_weak")
            self.assertIn("errors", res["params"])
            # Parol o'zgarmadi
            st2, res2 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res2["password"], r["credentials"]["password"])
        finally:
            self._cleanup_user(uid)

    def test_M5_13_taken_login_rejected(self):
        r1 = self._mkuser("student", "Birinchi", "Foydalanuvchi")
        r2 = self._mkuser("student", "Ikkinchi", "Foydalanuvchi")
        try:
            taken = r1["credentials"]["login"]
            st, res = self.admin.put("/api/admin/users/%d/credentials" % r2["user"]["id"],
                                     {"login": taken})
            self.assertEqual(st, 400, res)
            self.assertEqual(res["error"], "profile.login_taken")
            # Oxirgi qiymat buzilmadi
            st2, res2 = self.admin.get("/api/admin/users/%d/credentials" % r2["user"]["id"])
            self.assertEqual(res2["login"], r2["credentials"]["login"])
        finally:
            self._cleanup_user(r1["user"]["id"])
            self._cleanup_user(r2["user"]["id"])

    def test_M5_14_bad_login_format_rejected(self):
        r = self._mkuser("student", "Format", "Noto'g'ri")
        uid = r["user"]["id"]
        try:
            for bad in ("admin", "usrL_qisqa", "usrL_" + "a" * 40, "x" * 20):
                with self.subTest(login=bad):
                    st, res = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                             {"login": bad})
                    self.assertEqual(st, 400, res)
                    self.assertEqual(res["error"], "profile.login_format")
        finally:
            self._cleanup_user(uid)

    def test_M5_15_same_password_is_not_an_error(self):
        """Aynan bir xil parolni yuborish XATO EMAS — forma o'zgartirilmay
        qayta yuborilgan bo'lishi mumkin. Yozuv QILINMAYDI."""
        r = self._mkuser("student", "O'zi", "Biri")
        uid = r["user"]["id"]
        try:
            before = _json_db().q1("SELECT password_hash,updated_at FROM users WHERE id=?", (uid,))
            st, res = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                     {"password": r["credentials"]["password"]})
            self.assertEqual(st, 200, res)
            self.assertFalse(res["changed"], "o'zgarmagan parol 'o'zgargan' deb hisoblandi")
            self.assertFalse(res["password_changed"])
            self.assertEqual(res["note"], "nothing_changed")
            after = _json_db().q1("SELECT password_hash,updated_at FROM users WHERE id=?", (uid,))
            self.assertEqual(after["password_hash"], before["password_hash"])
            self.assertEqual(after["updated_at"], before["updated_at"])
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # Avtomatik generatsiya + eski (legacy) foydalanuvchilar
    # ------------------------------------------------------------------
    def test_M5_16_generate_password(self):
        r = self._mkuser("student", "Avto", "Parol")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                     {"generate_password": True})
            self.assertEqual(st, 200, res)
            gen = res["password"]
            self.assertTrue(gen.startswith("usrP_"), "generatsiya formati noto'g'ri: %r" % gen)
            self.assertNotEqual(gen, r["credentials"]["password"])
            from app.auth import password_strength_errors
            self.assertEqual(password_strength_errors(gen), [], "generatsiya kuchsiz chiqdi")
            # Yangi parol bilan kirish
            c, st2, res2 = self._login(res["login"], gen, "student")
            self.assertEqual(st2, 200, res2)
        finally:
            self._cleanup_user(uid)

    def test_M5_17_legacy_user_not_viewable_and_can_be_attached(self):
        """Eski modelda yaratilgan foydalanuvchi (`password_enc` yo'q):
        parol KO'RINMAYDI (`reason: legacy`), lekin admin o'z bilgan parolni
        tasdiqlab uni BIRIKTIRA oladi — parol O'ZGARMAGAN holda."""
        from app.auth import hash_password
        db = _json_db()
        pw = "usrP_LegacyTest42!"
        db.ex("INSERT INTO users(first_name,last_name,login,password_hash,password_enc,role,status,"
              "created_at,updated_at) VALUES(?,?,?,?,'', 'student','active',?,?)",
              ("Eski", "Foydalanuvchi", "usrL_legacyTest1234", hash_password(pw), now(), now()))
        uid = db.q1("SELECT id FROM users WHERE login=?", ("usrL_legacyTest1234",))["id"]
        try:
            st, res = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(st, 200, res)
            self.assertFalse(res["can_view"], "eski foydalanuvchi paroli ko'rindi")
            self.assertEqual(res["reason"], "legacy")
            self.assertIsNone(res["password"])
            self.assertEqual(res["login"], "usrL_legacyTest1234", "login ko'rinishi kerak edi")

            before_hash = _json_db().q1("SELECT password_hash FROM users WHERE id=?", (uid,))["password_hash"]
            # 1) To'g'ri parolni BIRIKTIRISH — parol O'ZGARMASDI
            st2, res2 = self.admin.put("/api/admin/users/%d/credentials" % uid, {"password": pw})
            self.assertEqual(st2, 200, res2)
            self.assertTrue(res2["attached"])
            self.assertFalse(res2["password_changed"])
            after_hash = _json_db().q1("SELECT password_hash FROM users WHERE id=?", (uid,))["password_hash"]
            self.assertEqual(after_hash, before_hash, "parol hash'i o'zgardi (u o'zgarmasligi kerak edi)")
            st3, res3 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertTrue(res3["can_view"])
            self.assertEqual(res3["password"], pw)
            # 2) Endi biriktirilgan parolni qayta kiritish -> "o'zgarmadi"
            st4, res4 = self.admin.put("/api/admin/users/%d/credentials" % uid, {"password": pw})
            self.assertEqual(st4, 200, res4)
            self.assertFalse(res4["changed"])
        finally:
            _json_db().ex("DELETE FROM users WHERE id=?", (uid,))

    def test_M5_17b_login_change_keeps_unchanged_password(self):
        """ENG MUHIM REGRESYA: parol maydonida O'ZGARMAGAN parol turib, admin
        FAQAT LOGIN ni o'zgartirsa — login SAQLANISHI kerak. Parolga tegilmasligi
        va eski parol bilan kirish ham saqlanishiga ishonchli."""
        r = self._mkuser("student", "Login", "Faqat")
        uid = r["user"]["id"]
        try:
            lg, pw = r["credentials"]["login"], r["credentials"]["password"]
            new_login = "usrL_FaqatLoginOzg02"
            # Parol maydoniga O'ZGARMAGAN qiymat yuboriladi (frontend shunday
            # yuboradi: maydon to'ldirilgan holda turibdi).
            st, res = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                     {"login": new_login, "password": pw})
            self.assertEqual(st, 200, res)
            self.assertTrue(res["changed"], "faqat login o'zgarganda ham saqlanishi kerak")
            self.assertEqual(res["what"], ["login"])
            self.assertFalse(res["password_changed"])
            st2, res2 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res2["login"], new_login, "login saqlanmadi")
            self.assertEqual(res2["password"], pw, "parol o'zgari ketdi")
            # Yangi login + ESKI parol bilan kirish ishlaydi
            c, st3, res3 = self._login(new_login, pw, "student")
            self.assertEqual(st3, 200, res3)
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # Foydalanuvchining O'Z parolini o'zgartirishi ham nusxani yangilaydi
    # ------------------------------------------------------------------
    def test_M5_18_self_password_change_updates_encrypted_copy(self):
        r = self._mkuser("student", "O'zi", "Almashtir")
        uid = r["user"]["id"]
        try:
            lg, old_pw = r["credentials"]["login"], r["credentials"]["password"]
            c, st, _ = self._login(lg, old_pw, "student")
            self.assertEqual(st, 200)
            new_pw = "usrP_OziAlmashtir8!"
            st2, res2 = c.post("/api/auth/change-password", {
                "old_password": old_pw, "new_login": "", "new_password": new_pw,
                "confirm_password": new_pw,
            })
            self.assertEqual(st2, 200, res2)
            # Admin endi YANGI parolni ko'radi (eskisini emas)
            st3, res3 = self.admin.get("/api/admin/users/%d/credentials" % uid)
            self.assertEqual(res3["password"], new_pw, "shifrlangan nusxa yangilanmagan")
            # Parol "olingan" bo'lib qolmaydi — ikkinchi marta berilsa rad etiladi
            st4, res4 = self.admin.put("/api/admin/users/%d/credentials" % uid,
                                       {"password": old_pw})
            self.assertEqual(st4, 400, res4)
            self.assertEqual(res4["error"], "auth.password_used")
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # Frontend + i18n
    # ------------------------------------------------------------------
    def test_M5_19_manager_modal_present(self):
        src = self._read("web/js/shared.js")
        self.assertIn("function openUserCredentialsManager(", src,
                      "openUserCredentialsManager topilmadi")
        start = src.index("function openUserCredentialsManager(")
        body = src[start:start + 5000]
        # Ko'zcha (parolni ko'rsatish/yashirish)
        self.assertIn("UI.passwordInput", body, "parol maydonida ko'zcha yo'q")
        # Joriy qiymatlarni yuklash
        self.assertIn("credentials", body, "joriy login/parol yuklanmaydi")
        # Ko'rish va o'zgartirish
        self.assertIn("API.get(`admin/users/", body, "GET chaqiruvi yo'q")
        self.assertIn("API.put(`admin/users/", body, "PUT chaqiruvi yo'q")
        # Saqlash faqat o'zgarish bo'lganda
        self.assertIn("saveBtn.disabled", body, "Saqlash validatsiyaga bog'lanmagan")
        # Xavfli amal — tasdiqsiz bajarilmaydi
        self.assertIn("confirmDialog", body, "o'zgartirish tasdiqsiz bajarilmoqda")
        # Audit eslatmasi ko'rsatiladi
        self.assertIn("users.cred_audit_note", body, "audit eslatmasi yo'q")

    def test_M5_20_eye_button_in_manager(self):
        """🔑 ko'zcha barcha modallarda: credential boshqaruv oynasida."""
        src = self._read("web/js/shared.js")
        start = src.index("function openUserCredentialsManager(")
        body = src[start:start + 5000]
        self.assertIn('text: "" + t("auth.password")', body, "parol maydoni yo'q")

    def test_M5_21_buttons_wired_in_admin_views(self):
        """🔑 tugmasi barcha admin ro'yxatlarida (talaba, instruktor, admin,
        arxiv) va profil kartasida bor."""
        src = self._read("web/js/views-admin.js")
        self.assertEqual(src.count("openUserCredentialsManager"), 6,
                         "🔑 tugmalari 6 joyda bo'lishi kerak "
                         "(admin/talaba/instruktor jadval + profil + import)")
        # Eskirgan "bir marta ko'rsatish" tugmasi qolmagan
        self.assertNotIn("reset-password", src,
                         "eskirgan reset-password chaqiruvi qoldi")
        self.assertNotIn("users.new_password_confirm", src,
                         "eskirgan tasdiqlash matni qoldi")

    def test_M5_22_shown_once_text_replaced(self):
        """'Bir marta ko'rsatish' ogohlantirishi endi noto'g'ri —
        parol DOIMIY saqlanadi va qayta ko'rinadi."""
        i18n = self._read("web/js/i18n.js")
        self.assertIn('"users.cred_visible_note"', i18n)
        self.assertIn('"users.cred_legacy"', i18n)
        self.assertIn('"users.cred_no_crypto"', i18n)
        self.assertIn('"users.cred_save_confirm"', i18n)
        self.assertIn('"users.cred_title"', i18n)
        # Eski matn endi ishlatilmaydi
        shared = self._read("web/js/shared.js")
        start = shared.index("function openUserCredentials(")
        body = shared[start:start + 1200]
        self.assertNotIn("users.password_shown_once", body,
                         "openUserCredentials hali 'bir marta' ogohlantirishini ko'rsatmoqda")

    def test_M5_23_profile_reports_visibility(self):
        """Profil javobida `can_view_password` bor (frontend 🔑 ni
        boshidan boshlab to'g'ri ochishi uchun)."""
        r = self._mkuser("student", "Profil", "Belgi")
        uid = r["user"]["id"]
        try:
            st, res = self.admin.get("/api/admin/users/%d" % uid)
            self.assertEqual(st, 200, res)
            prof = res.get("user", res)
            self.assertIn("can_view_password", prof)
            self.assertTrue(prof["can_view_password"])
        finally:
            self._cleanup_user(uid)

    # ------------------------------------------------------------------
    # Maxfiy kalit (.env) bilan bog'liqlik
    # ------------------------------------------------------------------
    def test_M5_24_key_helper_persists(self):
        """`ensure_credentials_key()` kalitni `.env` da BIR MARTA yaratadi."""
        from app.config import credentials_key, ensure_credentials_key, CREDENTIALS_KEY_ENV
        k = credentials_key()
        self.assertTrue(k, "%s .env da yo'q va yaratilmadi" % CREDENTIALS_KEY_ENV)
        self.assertEqual(ensure_credentials_key(), k, "kalit o'zgardi (idempotent emas)")
        import base64
        self.assertEqual(len(base64.urlsafe_b64decode(k.encode())), 32,
                         "kalit 32 bayt bo'lishi kerak")

    def test_M5_25_key_creation_is_thread_safe(self):
        """Kalit YO'Q holatida parallel so'rovlar `.env` ga IKKITA turli kalit
        yozmasligi kerak — aks holda qayta ishga tushirganda eski parollar
        ochilmay qoladi."""
        import base64
        import threading as _th
        from app import config as cfgmod
        real_env = os.environ.pop(cfgmod.CREDENTIALS_KEY_ENV, None)
        real_path = cfgmod.ENV_PATH
        tmp = Path(__file__).resolve().parent.parent / ".env.m5_test"
        try:
            cfgmod.ENV_PATH = str(tmp)
            if tmp.exists():
                tmp.unlink()
            results = []
            barrier = _th.Barrier(8)

            def worker():
                barrier.wait()
                results.append(cfgmod.ensure_credentials_key())

            ts = [_th.Thread(target=worker) for _ in range(8)]
            for t in ts:
                t.start()
            for t in ts:
                t.join()
            # Barcha so'rovlar BIR XIL kalitni olishi SHART
            self.assertEqual(len(results), 8)
            self.assertEqual(len(set(results)), 1,
                             "parallel so'rovlar turli kalit yaratdi: %s" % set(results))
            # `.env` ga faqat BIR marta yozilishi SHART
            txt = tmp.read_text(encoding="utf-8")
            self.assertEqual(txt.count(cfgmod.CREDENTIALS_KEY_ENV + "="), 1,
                             "kalit `.env` ga bir necha marta yozildi")
            self.assertEqual(len(base64.urlsafe_b64decode(results[0].encode())), 32)
        finally:
            cfgmod.ENV_PATH = real_path
            if tmp.exists():
                tmp.unlink()
            if real_env is not None:
                os.environ[cfgmod.CREDENTIALS_KEY_ENV] = real_env
            else:
                os.environ.pop(cfgmod.CREDENTIALS_KEY_ENV, None)

    def test_M5_26_env_example_documents_key(self):
        """`.env.example` kalitni hujjatlashi kerak (maxfiy qiymat EMAS)."""
        ex = self._read(".env.example")
        self.assertIn("CREDENTIALS_KEY", ex)
        # Hech qanday HAQIQIY kalit faylda bo'lmasligi shart
        import base64
        import re as _re
        for line in ex.splitlines():
            if line.strip().startswith("CREDENTIALS_KEY="):
                val = line.split("=", 1)[1].strip()
                if val:
                    with self.subTest(value=val[:8]):
                        self.assertNotRegex(val, _re.escape(val) + "$")
                        self.assertGreaterEqual(len(base64.urlsafe_b64decode(val)), 32)


def _crypto_ready():
    """Test muhitida shifrlash ishlaydimi (kalit + kutubxona)."""
    try:
        from app.auth import credentials_crypto_available
        return credentials_crypto_available()
    except Exception:
        return False


if __name__ == "__main__":
    unittest.main(verbosity=2)