# -*- coding: utf-8 -*-
"""CSV / XLSX / PDF eksporti.

Uchta format HAQIQIY fayl yaratadi:
  * CSV  — standart `csv` moduli + UTF-8 BOM (Excel/O'zbekcha harflar buzilmaydi),
           `;` ajratgich (o'zbek Excel shuni talab qiladi).
  * XLSX — `openpyxl` bilan haqiqiy Excel fayli.
  * PDF  — `reportlab` bilan HAQIQIY PDF (avvalgi versiya HTML fayl edi).

O'ZBEKCHA HARFLAR: PDF uchun shrift avtomatik tanlanadi (Arial/DejaVu/Noto/
Times). `ǒ`, `ǔ`, `gʻ`, `oʻ` kabi sheva belgilari `Helvetica` da YO'Q, shuning
uchun oddiy shrift ishlatilsa, PDF'da kvadratchalar chiqadi. Shuning uchun
TTF shrift majburiy izlanadi.

KUTUBXONALAR YO'Q BO'LSA: xato (500) emas — aniq `unavailable` xabari bilan
rad etiladi, "bo'sh/buzilgan fayl" qaytarilMAYDI.
"""
from __future__ import annotations

import csv
import io
import os

from .db import now

# --------------------------------------------------------------------------
# i18n: eksport sarlavhalari DB ustun nomlari sifatida keladi, shuning uchun
# ular uch tilga bog'liq bo'lishi kerak. Sarlavha -> {uz, ru, en}.
# --------------------------------------------------------------------------
H = {
    "Ism":               {"uz": "Ism",        "ru": "Имя",        "en": "First name"},
    "Familiya":          {"uz": "Familiya",   "ru": "Фамилия",    "en": "Last name"},
    "Rol":               {"uz": "Rol",        "ru": "Роль",       "en": "Role"},
    "Guruh":             {"uz": "Guruh",      "ru": "Группа",     "en": "Group"},
    "Toifa":             {"uz": "Toifa",      "ru": "Категория",  "en": "License category"},
    "Telefon":           {"uz": "Telefon",    "ru": "Телефон",    "en": "Phone"},
    "Holat":             {"uz": "Holat",      "ru": "Статус",     "en": "Status"},
    "Tug'ilgan_sana":    {"uz": "Tug'ilgan sana", "ru": "Дата рождения", "en": "Date of birth"},
    "Qabul_sanasi":      {"uz": "Qabul sanasi",    "ru": "Дата зачисления", "en": "Enrolled"},
    "Otilgan_darslar":   {"uz": "O'tilgan mashg'ulotlar", "ru": "Проведённые занятия", "en": "Lessons taken"},
    "Jami_belgilangan":  {"uz": "Jami belgilangan", "ru": "Всего назначено", "en": "Total assigned"},
    "Qolgan":            {"uz": "Qolgan",     "ru": "Осталось",   "en": "Remaining"},
    "Progress":          {"uz": "Progress",   "ru": "Прогресс",   "en": "Progress"},
    "Royxatga_olingan":  {"uz": "Ro'yxatga olingan", "ru": "Зачислен", "en": "Enrolled"},
    "Oxirgi_kirish":     {"uz": "Oxirgi kirish", "ru": "Последний вход", "en": "Last login"},
    "Kirish":            {"uz": "Kirish",     "ru": "Вход",       "en": "Login"},
    "Parol":             {"uz": "Parol",      "ru": "Пароль",     "en": "Password"},
    "Izoh":              {"uz": "Izoh",       "ru": "Примечание", "en": "Notes"},
    "Manzil":            {"uz": "Manzil",     "ru": "Адрес",      "en": "Address"},
    "Talabalar_soni":    {"uz": "Talabalar soni", "ru": "Кол-во учеников", "en": "Students count"},
    "Avtomobil":         {"uz": "Avtomobil",  "ru": "Автомобиль", "en": "Car"},
    "Davlat_raqami":     {"uz": "Davlat raqami", "ru": "Госномер", "en": "Plate"},
    "Ish_vaqti":         {"uz": "Ish vaqti",  "ru": "Рабочее время", "en": "Working hours"},
    "Huquqlar":          {"uz": "Huquqlar",   "ru": "Права",      "en": "Permissions"},
    # --- mashg'ulotlar (eski eksport) ---
    "Sana":              {"uz": "Sana",       "ru": "Дата",       "en": "Date"},
    "Boshlanish":        {"uz": "Boshlanish", "ru": "Начало",     "en": "Start"},
    "Tugash":            {"uz": "Tugash",     "ru": "Конец",      "en": "End"},
    "Instruktor":        {"uz": "Instruktor", "ru": "Инструктор", "en": "Instructor"},
    "Talabalar":         {"uz": "Talabalar",  "ru": "Ученики",    "en": "Students"},
    "Sigim":             {"uz": "Sigim",      "ru": "Вместимость", "en": "Capacity"},
    "Bekor_sababi":      {"uz": "Bekor sababi", "ru": "Причина отмены", "en": "Cancel reason"},
    "Vaqt":              {"uz": "Vaqt",       "ru": "Время",      "en": "Time"},
    "Qatnashish":        {"uz": "Qatnashish", "ru": "Посещение",  "en": "Attendance"},
    "Belgilangan":       {"uz": "Belgilangan", "ru": "Отмечено",  "en": "Marked"},
    "Yaratilgan":        {"uz": "Yaratilgan", "ru": "Создан",     "en": "Created"},
}

# Rol tarjimalari (qiymat ham DB'dan o'zbekcha/inglizcha kelishi mumkin)
ROLE_TXT = {
    "student":    {"uz": "Talaba",    "ru": "Ученик",   "en": "Student"},
    "instructor": {"uz": "Instruktor", "ru": "Инструктор", "en": "Instructor"},
    "admin":      {"uz": "Admin",     "ru": "Админ",    "en": "Admin"},
}
STATUS_TXT = {
    "active":   {"uz": "Faol",       "ru": "Активен",  "en": "Active"},
    "blocked":  {"uz": "Bloklangan", "ru": "Заблокирован", "en": "Blocked"},
    "archived": {"uz": "Arxivlangan", "ru": "В архиве", "en": "Archived"},
}

CSV_MIME = "text/csv; charset=utf-8"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"


class ExportUnavailable(Exception):
    """Kutubxona yoki shrift topilmadi — aniq xabar bilan rad etish."""

    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


# ==========================================================================
# 1) CSV — UTF-8 BOM bilan (Excel o'zbekcha harflarni to'g'ri ko'rsatadi)
# ==========================================================================
def to_csv(rows: list, filename: str = "hisobot.csv") -> dict:
    """UTF-8 BOM + `;` ajratgich. O'zbek Excel `;` ni kutadi (`,` esa qatorni
    aralashitiradi). BOMsiz faylda Excel `O'g'` ni `O?g'` qilib ko'rsatadi."""
    headers = list(rows[0].keys()) if rows else []
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL,
                        lineterminator="\r\n")
    if headers:
        writer.writerow(headers)
        for r in rows:
            writer.writerow(["" if r.get(h) is None else str(r.get(h)) for h in headers])
    # utf-8-sig = BOM + UTF-8
    data = buf.getvalue().encode("utf-8-sig")
    return {"filename": filename, "mime": CSV_MIME, "data": data}


# ==========================================================================
# 2) XLSX — openpyxl bilan haqiqiy Excel fayli
# ==========================================================================
def to_xlsx(rows: list, sheet: str = "Hisobot", lang: str = "uz") -> dict:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ExportUnavailable("export.xlsx_unavailable")

    headers = list(rows[0].keys()) if rows else []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet or "Hisobot")[:31]

    # --- sarlavha qatori ---
    if headers:
        ws.append([_label(h, lang) for h in headers])
        fill = PatternFill("solid", fgColor="0F2B5B")
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=ci)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- ma'lumot qatorlari (qiymatlar ham tarjima qilinadi) ---
    for r in rows:
        ws.append([_cell(r.get(h), lang) for h in headers])

    # --- ustun kengligi ---
    for ci, h in enumerate(headers, start=1):
        vals = [_label(h, lang)] + [str(r.get(h) or "") for r in rows[:200]]
        width = min(46, max(10, max((len(v) for v in vals), default=10) + 2))
        ws.column_dimensions[get_column_letter(ci)].width = width
    if headers:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return {"filename": "hisobot.xlsx", "mime": XLSX_MIME, "data": buf.getvalue()}


def _label(col: str, lang: str) -> str:
    d = H.get(col)
    if not d:
        return col
    return d.get(lang) or d["uz"]


def _cell(v, lang: str):
    """Qiymatni tarjima qiladi (rol/holat) yoki o'zgarishsiz qaytaradi."""
    if v is None:
        return ""
    s = str(v)
    if s in ROLE_TXT:
        return ROLE_TXT[s].get(lang) or ROLE_TXT[s]["uz"]
    if s in STATUS_TXT:
        return STATUS_TXT[s].get(lang) or STATUS_TXT[s]["uz"]
    return s


# ==========================================================================
# 3) PDF — reportlab bilan HAQIQIY PDF
# ==========================================================================
# O'zbekcha sheva belgilari (ǒ, ǔ, gʻ, oʻ) faqat TTF shriftda bor.
_FONT = {"name": None, "path": None, "family": "Helvetica"}

# Tartib: o'zbekcha to'liq qo'llab-quvvatlanadigan shriftlar birinchi.
_FONT_CANDIDATES = (
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ("DejaVuSans", "/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    ("DejaVuSans", "/usr/share/fonts/TTF/DejaVuSans.ttf"),
    ("NotoSans", "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
    ("Arial", "C:/Windows/Fonts/arial.ttf"),
    ("Arial", "C:/Windows/Fonts/ARIALUNI.TTF"),
    ("Arial", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    ("LiberationSans", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    ("Times", "C:/Windows/Fonts/times.ttf"),
    ("Times", "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"),
    ("DejaVuSans", str(os.path.join(os.path.dirname(__file__), "..", "assets", "DejaVuSans.ttf"))),
)

# Tekshiriladigan maxsus belgilar — bular bo'lmasa shrift foydali emas.
_PROBE = "ǒǔgʻoʻ’‘"


def _ensure_font():
    """Ilk chaqiruvda o'zbekcha belgilarni qo'llab-quvvatlaydigan TTF topadi."""
    if _FONT["name"]:
        return _FONT
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return _FONT  # reportlab yo'q — keyin import check buni ushlaydi
    for name, path in _FONT_CANDIDATES:
        if not path or not os.path.isfile(path):
            continue
        try:
            pdfmetrics.registerFont(TTFont(name, path))
            face = pdfmetrics.getFont(name).face
            cmap = getattr(face, "charToGlyph", {})
            if all(ord(ch) in cmap for ch in _PROBE):
                _FONT.update(name=name, path=path, family=name)
                return _FONT
        except Exception:
            continue
    return _FONT  # topilmadi -> Helvetica (lotin harflar ishlaydi, ǒ/ǔ yo'q)


def to_pdf(rows: list, title: str = "Hisobot", lang: str = "uz") -> dict:
    """`reportlab` bilan HAQIQIY PDF (A4 landskap, jadval ko'rinishida).

    Eslatma: avvalgi versiya `text/html` + `window.print()` edi — bu PDF
    EMAS edi. Endi haqiqiy PDF fayl yaratiladi.
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    except ImportError:
        raise ExportUnavailable("export.pdf_unavailable")

    font = _ensure_font()
    f_name = font["name"] or "Helvetica"

    headers = list(rows[0].keys()) if rows else []
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=title, author="AVTO MAKTAB",
    )
    styles = getSampleStyleSheet()
    ttl = ParagraphStyle("T", parent=styles["Title"], fontName=f_name,
                         fontSize=15, leading=19, spaceAfter=4)
    sub = ParagraphStyle("S", parent=styles["Normal"], fontName=f_name,
                         fontSize=9, textColor=colors.HexColor("#555555"))

    story = [Paragraph(_esc_xml(title), ttl),
             Paragraph(_esc_xml("%s  ·  %s" % (now(), f"{len(rows)} ta yozuv")), sub),
             Spacer(1, 5 * mm)]

    if not rows:
        story.append(Paragraph(_esc_xml(_label("Holat", lang)) + ": —", sub))
    else:
        data = [[_esc_xml(_label(h, lang)) for h in headers]]
        for r in rows:
            data.append([_esc_xml(_cell(r.get(h), lang)) for h in headers])
        ncol = max(1, len(headers))
        avail = landscape(A4)[0] - 24 * mm
        # Juda ko'p ustun bo'lsa shriftni kichraytiramiz
        fs = 8 if ncol <= 8 else (7 if ncol <= 12 else 6)
        t = Table(data, repeatRows=1, colWidths=[avail / ncol] * ncol)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), f_name),
            ("FONTNAME", (0, 1), (-1, -1), f_name),
            ("FONTSIZE", (0, 0), (-1, -1), fs),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F2B5B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D4E8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F4F7FC")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        story.append(t)

    doc.build(story)
    data_bytes = buf.getvalue()
    # Ishlashini tekshiramiz — bo'sh/buzilgan fayl qaytarMAYmiz
    if not data_bytes.startswith(b"%PDF"):
        raise ExportUnavailable("export.pdf_broken")
    return {"filename": "hisobot.pdf", "mime": PDF_MIME, "data": data_bytes}


def _esc_xml(s: str) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


# ==========================================================================
# Umumiy kirish nuqtasi
# ==========================================================================
def export_table(rows: list, fmt: str, sheet: str = "Hisobot",
                 lang: str = "uz", filename: str = None) -> dict:
    """`rows` — dict ro'yxati (bir xil kalitlar). `fmt` — csv|xlsx|pdf."""
    fmt = (fmt or "csv").lower()
    if fmt == "csv":
        return to_csv(rows, filename or ((sheet or "hisobot").replace(" ", "_") + ".csv"))
    if fmt == "xlsx":
        return to_xlsx(rows, sheet, lang)
    if fmt == "pdf":
        return to_pdf(rows, sheet, lang)
    raise ExportUnavailable("export.bad_format")
